"""Create the Block 7 editorial revision from the verified Block 6 draft."""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "docs"
    / "revision_octubre_2026"
    / "word"
    / "TRABAJO_FINAL_TPScouting_CORREGIDO_OCTUBRE_2026_BORRADOR.docx"
)
OUTPUT = (
    ROOT
    / "docs"
    / "revision_octubre_2026"
    / "word"
    / "TRABAJO_FINAL_TPScouting_CORREGIDO_BLOQUE7_2026-10-06.docx"
)
EXPECTED_SOURCE_SHA256 = (
    "c6c0d18fc7f6b4cbec1c05bd6c55a4f385ece04040ec1a411bcb7af3f174a1e4"
)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def body_paragraphs(document: Document) -> list[Paragraph]:
    return [p for p in document.paragraphs if not p.style.name.lower().startswith("toc")]


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in body_paragraphs(document) if p.text.strip().startswith(prefix)]
    if len(matches) > 1:
        exact = [p for p in matches if p.text.strip() == prefix]
        if len(exact) == 1:
            return exact[0]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one body paragraph starting {prefix!r}; found {len(matches)}")
    return matches[0]


def replace(document: Document, prefix: str, text: str, style: str | None = None) -> None:
    paragraph = find_paragraph(document, prefix)
    paragraph.text = text
    if style:
        paragraph.style = style


def insert_before(anchor: Paragraph, text: str, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def delete_body_range(document: Document, first_prefix: str, stop_prefix: str) -> None:
    paragraphs = list(document.paragraphs)
    first_element = find_paragraph(document, first_prefix)._p
    stop_element = find_paragraph(document, stop_prefix)._p
    first = next(i for i, p in enumerate(paragraphs) if p._p is first_element)
    stop = next(i for i, p in enumerate(paragraphs) if p._p is stop_element)
    if first >= stop:
        raise RuntimeError(f"Invalid paragraph range {first_prefix!r} to {stop_prefix!r}")
    for paragraph in paragraphs[first:stop]:
        if paragraph._p.xpath("./w:pPr/w:sectPr"):
            paragraph.text = ""
        else:
            delete_paragraph(paragraph)


def delete_table(document: Document, header: str) -> None:
    matches = [t for t in document.tables if t.cell(0, 0).text.strip() == header]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one table headed {header!r}; found {len(matches)}")
    table = matches[0]
    table._element.getparent().remove(table._element)


def delete_table_containing(document: Document, text: str) -> None:
    matches = [
        table
        for table in document.tables
        if any(text in cell.text for row in table.rows for cell in row.cells)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one table containing {text!r}; found {len(matches)}")
    table = matches[0]
    table._element.getparent().remove(table._element)


def replace_in_tables(document: Document, old: str, new: str) -> None:
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if old in cell.text:
                    cell.text = cell.text.replace(old, new)


def remove_glossary_rows(document: Document, terms: set[str]) -> None:
    glossary = next(
        table
        for table in document.tables
        if table.cell(0, 0).text.strip() == "Término"
        and table.cell(0, 1).text.strip() == "Definición"
    )
    for row in list(glossary.rows)[1:]:
        if row.cells[0].text.strip() in terms:
            glossary._tbl.remove(row._tr)


def main() -> None:
    if digest(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("The Block 6 source differs from its verified SHA-256")

    document = Document(SOURCE)

    # Summary and introductory prose: punctuation, impersonal voice and scope.
    replace(
        document,
        "El seguimiento de futbolistas juveniles",
        "El seguimiento de futbolistas juveniles en clubes con recursos limitados suele depender de "
        "registros dispersos y criterios difíciles de comparar a lo largo del tiempo. Ante este problema, "
        "se diseñó y desarrolló TPScouting, un producto mínimo viable orientado a organizar y analizar "
        "información deportiva de jugadores de 12 a 18 años. La aplicación permite gestionar perfiles, "
        "atributos técnicos, estadísticas, evaluaciones, reportes, disponibilidad e historiales. También "
        "incluye un panel general, herramientas de comparación y una estimación de potencial producida por "
        "un modelo implementado en PyTorch. La solución se evaluó mediante datos sintéticos, pruebas "
        "automatizadas, métricas de clasificación, modelos de referencia y evidencias fechadas de la "
        "aplicación, el despliegue y la integración continua. Los resultados demuestran la factibilidad "
        "técnica del flujo comprendido entre el registro y la persistencia de datos y la presentación de "
        "una estimación en la interfaz. El sistema se plantea como apoyo para organizar información y hacer "
        "trazables los criterios de evaluación, sin sustituir el análisis de entrenadores y scouts.",
    )
    replace(
        document,
        "El software se inspira",
        "El trabajo se sitúa en el uso de datos y aprendizaje automático como apoyo al análisis deportivo. "
        "La literatura muestra aplicaciones para valorar acciones, comparar perfiles y estudiar la evolución "
        "de jugadores, aunque la utilidad de cada enfoque depende del origen de los datos y de su contexto "
        "de validación (Pappalardo et al., 2019; Decroos et al., 2019; Lacan, 2024).",
    )
    replace(
        document,
        "Los clubes pequeños y medianos",
        "En instituciones con recursos limitados, los registros pueden permanecer distribuidos entre "
        "planillas y observaciones no normalizadas. En esas condiciones resulta difícil conservar la "
        "trazabilidad y comparar evaluaciones; el trabajo no cuantifica la frecuencia ni el efecto deportivo "
        "de esa situación.",
    )
    replace(
        document,
        "Este software busca reducir",
        "TPScouting evalúa la factibilidad de centralizar esos registros mediante componentes de software de "
        "uso abierto. La reducción de una brecha tecnológica o la mejora de decisiones no fueron medidas y se "
        "mantienen como posibles líneas de validación futura.",
    )
    replace(
        document,
        "La propuesta de este sistema",
        "La justificación técnica consiste en integrar, dentro de un único MVP, registros estructurados, "
        "visualizaciones, comparadores y una estimación reproducible. La evaluación comprueba el funcionamiento "
        "del flujo, pero no permite afirmar mejoras en la precisión del scouting ni igualdad entre clubes.",
    )
    replace(
        document,
        "Innovación tecnológica:",
        "Integración técnica: analítica de datos y aprendizaje automático aplicados a un flujo de scouting juvenil.",
    )
    replace(
        document,
        "Aplicación práctica:",
        "Aplicación funcional: registro, consulta y comparación de información deportiva en un MVP web.",
    )
    replace(
        document,
        "Contribución a la igualdad:",
        "Accesibilidad potencial: uso de componentes abiertos y reproducibles; la adopción y el impacto social no fueron evaluados.",
    )
    replace(
        document,
        "Con la expansión de la IA",
        "La incorporación de inteligencia artificial al análisis deportivo ofrece herramientas para organizar "
        "y comparar información. Plataformas como Wyscout centralizan video y datos para apoyar la observación, "
        "la comparación y el reclutamiento (Hudl, s. f.). Sin embargo, disponer de más datos no garantiza "
        "decisiones más precisas: los resultados dependen de la calidad del registro, del objetivo definido y "
        "de la revisión humana.",
    )
    replace(
        document,
        "Una de las principales limitaciones",
        "Una limitación relevante para clubes con recursos reducidos es la gestión sistemática de la información "
        "de sus jugadores. El proyecto aborda el problema de registro y consulta, pero no evalúa inversiones ni "
        "resultados institucionales.",
    )
    replace(
        document,
        "Otra dificultad relevante",
        "Los clubes pueden generar información de entrenamientos, partidos y evaluaciones sin disponer de un "
        "medio uniforme para procesarla. Los registros manuales pueden dificultar la recuperación, comparación "
        "y trazabilidad de esa información.",
    )
    replace(document, "• Dashboard o panel general", "• Panel general con métricas adaptadas al rol del usuario.")
    replace(
        document,
        "• Acceder al deploy público",
        "• Acceder al despliegue documentado, iniciar sesión y navegar el panel general, el listado, la ficha, la predicción y los comparadores.",
    )
    replace(
        document,
        "El alcance del proyecto es",
        "El alcance del proyecto comprende un MVP académico de scouting juvenil que permite registrar jugadores, "
        "cargar atributos y estadísticas, consultar fichas, comparar perfiles y calcular una estimación de "
        "potencial mediante un modelo de aprendizaje automático integrado al backend Flask.",
    )

    # Remove pseudo-headings and generic examples that were not needed by the implemented scope.
    for prefix in [
        "Definición y Principios Básicos",
        "Subcampos del Machine Learning",
        "Aprendizaje Supervisado:",
        "Aprendizaje No Supervisado:",
        "Aprendizaje por Refuerzo:",
        "Cómo las técnicas de Machine Learning",
        "Predicciones Basadas en Datos:",
    ]:
        delete_paragraph(find_paragraph(document, prefix))

    # Condense theory sections that contained unsourced, promotional enumerations.
    replace(
        document,
        "Importancia del Análisis de Datos",
        "El análisis de datos deportivos permite resumir eventos, comparar perfiles y conservar evidencia de "
        "las evaluaciones. Los enfoques de PlayeRank y VAEP muestran que la utilidad de una métrica depende de "
        "la función asignada al jugador y del contexto de cada acción (Pappalardo et al., 2019; Decroos et al., 2019).",
    )
    for prefix in [
        "Transformación del deporte moderno:",
        "Toma de decisiones basada en evidencia:",
        "Técnicas y Herramientas de Análisis de Datos",
        "Técnicas Estadísticas:",
        "Herramientas analíticas:",
    ]:
        delete_paragraph(find_paragraph(document, prefix))
    replace(
        document,
        "Transformación del entrenamiento y la estrategia",
        "En este trabajo, la inteligencia artificial se utiliza en una tarea acotada de clasificación sobre "
        "datos sintéticos. No se evaluaron recomendaciones tácticas, programas de entrenamiento ni efectos sobre "
        "el rendimiento real. Los antecedentes de scouting predictivo se consideran comparaciones metodológicas "
        "y no pruebas de validez para TPScouting (Lacan, 2024; van Arem et al., 2025).",
    )
    for prefix in [
        "Cambios en la forma de entrenamiento:",
        "Mejoramiento de las tácticas de juego:",
        "Mejora en la detección y desarrollo de talentos",
        "Identificación temprana de perfiles:",
        "Desarrollo Personalizado de Jugadores:",
    ]:
        delete_paragraph(find_paragraph(document, prefix))
    replace(
        document,
        "Desafíos y consideraciones éticas",
        "El uso de información deportiva de menores plantea riesgos de privacidad, seguridad y sesgo. La edad "
        "relativa, la maduración y el contexto pueden afectar tanto los registros como su interpretación; por "
        "ello se requieren transparencia y revisión humana (Cobley et al., 2009). Las obligaciones para un uso "
        "real se desarrollan en la Sección 6.5.",
    )
    for prefix in [
        "Privacidad y Seguridad de los Datos:",
        "Ética en el uso de datos:",
        "Sesgos en algoritmos de IA:",
    ]:
        delete_paragraph(find_paragraph(document, prefix))
    replace(
        document,
        "Data Science",
        "La ciencia de datos combina procedimientos de obtención, preparación, análisis y comunicación de "
        "datos. El aprendizaje automático aporta modelos para tareas predictivas, cuya validez depende de la "
        "separación entre entrenamiento y evaluación (Hastie et al., 2009; Goodfellow et al., 2016).",
    )
    replace(
        document,
        "La Ciencia de Datos es",
        "Aunque Big Data suele describir problemas de volumen, velocidad y variedad, TPScouting no procesa "
        "datos masivos. El proyecto utiliza un conjunto sintético de 20.000 jugadores para comprobar un flujo "
        "técnico; por ello se emplean los términos ciencia de datos y analítica sin atribuir al MVP una "
        "arquitectura de Big Data.",
    )
    for prefix in [
        "Big Data",
        "Data Science y Big Data en el deporte:",
        "Transformación del Scouting",
        "Predicción de Potencial",
        "Influencia en la Estrategia Deportiva",
        "Optimización de Estrategias",
        "Análisis en tiempo real.",
        "Monitoreo de la condición física.",
        "Decisiones Basadas en Datos",
    ]:
        try:
            delete_paragraph(find_paragraph(document, prefix))
        except RuntimeError:
            pass
    try:
        delete_paragraph(find_paragraph(document, "Big Data se refiere"))
    except RuntimeError:
        pass
    replace(
        document,
        "Descripción: Los analistas",
        "La analítica de datos comprende la preparación, exploración, síntesis y visualización de información "
        "para responder preguntas definidas. En TPScouting se aplica al historial de atributos, las estadísticas, "
        "los reportes y las comparaciones. El componente predictivo se trata por separado porque requiere "
        "entrenamiento y evaluación de un modelo.",
    )
    for prefix in [
        "Responsabilidades:",
        "Recopilar, organizar",
        "Analizar patrones",
        "Crear visualizaciones",
        "Trabajar con herramientas",
        "La analítica de datos y Big Data",
        "Data Analytics se centra",
    ]:
        delete_paragraph(find_paragraph(document, prefix))
    replace(
        document,
        "La discriminación y la calibración",
        "La discriminación y la calibración responden a preguntas diferentes: ROC-AUC evalúa el orden de "
        "los casos, mientras que una probabilidad calibrada busca coherencia entre frecuencias observadas y "
        "probabilidades estimadas. La calibración isotónica debe ajustarse fuera del conjunto de test para "
        "conservar una evaluación final independiente (Niculescu-Mizil & Caruana, 2005). Con una clase "
        "positiva minoritaria, PR-AUC complementa ROC-AUC porque expone la relación entre precisión y recall "
        "para la clase de interés (Saito & Rehmsmeier, 2015). Cualquier estadístico calculado con todo el "
        "conjunto antes del split puede transferir información de validación o test y debe declararse como "
        "riesgo de fuga (Kaufman et al., 2012).",
    )

    # Methodology: describe the practice actually followed and rename the undated table.
    replace(
        document,
        "Se aplicó un enfoque incremental",
        "Se aplicó un enfoque incremental organizado en bloques pequeños y verificables. Cada cambio relevante "
        "se acompañó con pruebas focales, una ejecución de la suite y documentación técnica. Se tomaron como "
        "referencia conceptos de iteración e inspección descritos en la Guía Scrum, pero no se aplicó Scrum de "
        "forma completa porque no se documentaron sus responsabilidades, eventos y artefactos (Schwaber & "
        "Sutherland, 2020).",
    )
    lifecycle = find_paragraph(document, "3.3 Ciclo de vida")
    insert_before(
        find_paragraph(document, "1. Análisis de requisitos"),
        "El ciclo se presenta como una secuencia de ingeniería de software adaptada al proyecto y no como "
        "evidencia de sprints formales (Pressman, 2010; Sommerville, 2011).",
    )
    lifecycle.text = "3.3 Ciclo de vida del software"
    replace(
        document,
        "3.4 Desglose de tareas y cronograma consolidado",
        "3.4 Secuencia de actividades realizadas",
        "Heading 2",
    )
    replace(
        document,
        "Tabla 3-1. Desglose de tareas y cronograma consolidado",
        "Tabla 3-1. Secuencia de actividades realizadas.",
    )

    # Remove the unused-technology essay and its three comparison tables.
    delete_body_range(document, "Análisis de componentes específicos:", "La arquitectura actual es")
    architecture = find_paragraph(document, "La arquitectura actual es")
    insert_before(architecture, "4.1.1.1 Decisiones de arquitectura", "Heading 4")
    insert_before(
        architecture,
        "Se seleccionó Flask porque la aplicación y el pipeline de aprendizaje automático utilizan Python y "
        "porque el renderizado server-side cubre el alcance del MVP (Pallets Projects, s. f.). SQLAlchemy "
        "mantiene un mismo modelo de dominio sobre SQLite local y PostgreSQL en el despliegue (SQLAlchemy, "
        "s. f.). .NET Minimal APIs, Peewee, MySQL y SQL Server no fueron implementados ni comparados mediante "
        "benchmarks; por ese motivo no se les atribuyen ventajas o desventajas experimentales.",
    )
    delete_table_containing(document, ".NET Minimal APIs")
    delete_table_containing(document, "Peewee")
    delete_table_containing(document, "SQL Server")
    for old, new in [
        ("Tabla 4-3. Capas", "Tabla 4-1. Capas"),
        ("Tabla 4-5. Dimensiones", "Tabla 4-2. Dimensiones"),
        ("Tabla 4-6. Requisitos", "Tabla 4-3. Requisitos"),
        ("Tabla 4-7. Entidades", "Tabla 4-4. Entidades"),
        ("Tabla 4-8. Endpoints", "Tabla 4-5. Endpoints"),
        ("Tabla 4-9. Decisiones", "Tabla 4-6. Decisiones"),
        ("Tabla 4-10. Casos", "Tabla 4-7. Casos"),
    ]:
        replace(document, old, new)

    # Align the database narrative with the eleven entities already listed in its table.
    replace(
        document,
        "Tablas principales:",
        "El modelo implementa once entidades: Player, Match, Coach, Director, User, PlayerStat, "
        "PlayerAttributeHistory, PlayerMatchParticipation, ScoutReport, PhysicalAssessment y "
        "PlayerAvailability. Sus responsabilidades se resumen en la Tabla 4-4.",
    )
    for prefix in [
        "players (Player):",
        "player_stats (PlayerStat):",
        "player_attribute_history",
        "Relaciones: Player se vincula",
        "Análisis de componente específico:",
        "Para una producción real",
    ]:
        delete_paragraph(find_paragraph(document, prefix))
    replace(
        document,
        "La interfaz fue desarrollada",
        "La interfaz se implementó con Bootstrap, plantillas Jinja2, CSS propio y Chart.js. Se priorizaron "
        "formularios, tablas, modales y gráficos compatibles con el flujo server-side del MVP.",
    )
    replace(
        document,
        "PyTorch se emplea",
        "PyTorch se emplea para implementar PlayerNet (Paszke et al., 2019; PyTorch Foundation, s. f.). "
        "Scikit-learn aporta preprocesamiento, partición, métricas, regresión logística y calibración "
        "isotónica (Pedregosa et al., 2011). La inferencia web utiliza PlayerNet como modelo operativo; la "
        "regresión logística se conserva como referencia experimental.",
    )
    replace(
        document,
        "El MVP implementa autenticación",
        "La autenticación se basa en sesiones, contraseñas hasheadas, protección CSRF y autorización por rol. "
        "Las medidas se contrastaron con riesgos de control de acceso, fallos criptográficos e inyección "
        "descritos por OWASP (2021), sin presentar esa referencia como una certificación de seguridad.",
    )
    replace(
        document,
        "El login incorpora rate limiting",
        "El login aplica rate limiting local mediante una clave de cliente definida por la política de proxies "
        "confiables. Esta mitigación fue verificada para una instancia, pero no sustituye un limitador distribuido "
        "en un despliegue multi-instancia.",
    )

    # Terminology, voice, punctuation and consistent captions.
    replacements = {
        "• Panel general dinámico por rol:": "• Panel general dinámico por rol: prioridades de scouting para scout/administrador y estado del plantel para director técnico.",
        "La Figura 6-3 presenta el panel general": "La Figura 6-3 presenta el panel general, que concentra indicadores, prioridades y visualizaciones para apoyar el seguimiento.",
        "Figura 6-3. Panel general o mesa de scouting": "Figura 6-3. Panel general con métricas por rol.",
        "A partir del desarrollo realizado": "El desarrollo permitió construir un producto mínimo viable que integra el registro de jugadores, el seguimiento de atributos e historiales, la consulta de estadísticas y evaluaciones, la visualización de indicadores, la comparación de perfiles y un módulo de aprendizaje automático. La solución recorre el flujo comprendido entre la carga y persistencia de datos y la presentación de una estimación en la interfaz. Las pruebas automatizadas y las evidencias técnicas demuestran la factibilidad de esa integración dentro del alcance definido.",
        "El aporte principal del proyecto": "El aporte comprobado del proyecto reside en centralizar y estructurar información que puede permanecer distribuida entre planillas, observaciones y registros aislados. El sistema permite consultar la evolución de cada jugador y aplicar criterios consistentes de registro. No se midieron reducciones de subjetividad, tiempo de evaluación ni costos operativos; esos beneficios continúan como posibilidades sujetas a validación externa.",
        "En relación con el componente": "El componente de aprendizaje automático integra preprocesamiento, artefactos, calibración e inferencia. Las métricas obtenidas con datos sintéticos permiten describir la corrida y compararla con referencias. La regresión logística obtuvo resultados levemente superiores en algunas métricas, mientras que PlayerNet crudo alcanzó el mayor PR-AUC. El bootstrap no aporta evidencia suficiente de superioridad de ninguno de los dos modelos y tampoco demuestra equivalencia. La selección futura debe considerar evidencia con varias semillas, costo de errores y condiciones reales de uso.",
        "• Persistencia y operación suficientes": "• Persistencia y operación suficientes para una demostración, pero no para uso multiusuario intensivo.",
        "• Incorporar auditoría de cambios": "• Incorporar auditoría de cambios, registros estructurados, caché distribuida y rate limiting centralizado.",
        "La Figura 10-1 conserva": "La Figura 10-1 conserva, con fines de trazabilidad histórica, una vista ampliada de la CI #73 del repositorio de desarrollo. Esa ejecución corresponde al commit bc5ddd35d0fa3bf6d85faab637772d4e9025fc98 y no certifica el repositorio de entrega ni las correcciones locales.",
        "Figura 10-1. Historial público": "Figura 10-1. Evidencia histórica de la CI #73 del repositorio de desarrollo.",
        "Esta sección incorpora los diagramas": "Esta sección reproduce ocho diagramas del cuerpo en páginas independientes. Se conservan porque el tamaño del cuerpo dificulta leer etiquetas y relaciones; los anexos sirven como ampliación visual y no como evidencia adicional ni como diagramas diferentes.",
        "La Figura 10-2 presenta": "Las Figuras 10-2 a 10-6 amplían, respectivamente, los componentes, el modelo de dominio, la inferencia, el panel general y el despliegue.",
        "La Figura 10-7 presenta": "Las Figuras 10-7 a 10-9 amplían los casos de uso generales, la gestión de jugadores y el análisis para decisión scout.",
    }
    for prefix, text in replacements.items():
        replace(document, prefix, text)
    replace(
        document,
        "En una línea más próxima",
        "En una línea próxima al scouting predictivo, Lacan (2024) utiliza una arquitectura de aprendizaje "
        "profundo por stacking para detectar jugadores de alto potencial sobre una base abierta. Van Arem et "
        "al. (2025) comparan modelos explicables para pronosticar la calidad y el valor futuro de futbolistas "
        "profesionales. TPScouting se diferencia por integrar captura, seguimiento histórico, visualización, "
        "reglas operativas e inferencia sobre datos sintéticos; no se plantea como competidor de plataformas "
        "profesionales.",
    )
    replace(
        document,
        "Por último, la integración entre desarrollo",
        "La integración entre desarrollo, pruebas, despliegue y documentación permitió conservar evidencia "
        "trazable. El registro de métricas, hashes, cobertura, limitaciones y fechas mejora la reproducibilidad "
        "del trabajo.",
    )
    for caption in [
        "Figura 4-1. Diagrama de componentes de TPScouting",
        "Figura 5-2. Secuencia de carga del panel general",
        "Figura 10-2. Diagrama de componentes de TPScouting ampliado",
        "Figura 10-6. Diagrama de despliegue en Render ampliado",
    ]:
        paragraph = find_paragraph(document, caption)
        if not paragraph.text.rstrip().endswith("."):
            paragraph.text = paragraph.text.rstrip() + "."
    stray = [p for p in body_paragraphs(document) if p.text.strip() == "."]
    for paragraph in stray:
        delete_paragraph(paragraph)

    # Introduce a proper Discussion section before ethics and use plural chapter naming.
    ethics = find_paragraph(document, "6.4 Consideraciones éticas")
    ethics.text = "6.5 Consideraciones éticas y protección de datos"
    insert_before(ethics, "6.4 Discusión", "Heading 2")
    insert_before(
        ethics,
        "Los resultados muestran que el flujo técnico es reproducible en el entorno evaluado, pero no permiten "
        "inferir utilidad deportiva. PlayerNet y la regresión logística presentan desempeños próximos en un único "
        "split sintético; el intervalo bootstrap de sus diferencias incluye cero. En consecuencia, la complejidad "
        "de PlayerNet no queda justificada por una superioridad estadística en esta corrida. El score combinado "
        "que observa el usuario tampoco cuenta con una evaluación independiente en test.",
    )
    insert_before(
        ethics,
        "La validez está limitada por una prevalencia fijada por diseño, cuantiles calculados antes del split, "
        "una sola semilla y ausencia de jugadores reales. El MVP sí aporta evidencia de integración, persistencia, "
        "control de acceso y presentación de resultados. Una conclusión sobre identificación de talento, ahorro "
        "de recursos o igualdad de oportunidades requiere datos longitudinales, varios splits y participación de "
        "clubes y profesionales.",
    )
    replace(
        document,
        "La aplicación fue verificada con datos sintéticos",
        "La aplicación fue verificada con datos sintéticos y no se utilizó con menores reales. Un uso real "
        "debería definir una base jurídica y una finalidad específica, informar a los titulares y a sus "
        "representantes, obtener el consentimiento que corresponda, recolectar sólo datos necesarios, aplicar "
        "plazos de retención y mecanismos de acceso, rectificación y supresión, y limitar el acceso por rol. "
        "La Ley 25.326 regula en Argentina los principios de protección de datos, los derechos de los titulares "
        "y las obligaciones de responsables de archivos y bancos de datos (Honorable Congreso de la Nación "
        "Argentina, 2000). Estas medidas corresponden a una etapa posterior; el MVP no implementa un flujo "
        "completo de consentimiento de tutores ni una política automatizada de retención y baja.",
    )
    replace(document, "7. CONCLUSIÓN", "7. CONCLUSIONES", "Heading 1")
    replace(document, "7.1 Conclusiones", "7.1 Conclusiones principales", "Heading 2")

    # Reproduction details that remained incomplete.
    replace(
        document,
        "Repositorio de entrega revisado:",
        "Repositorio de entrega revisado: https://github.com/indio21/TPScouting-entrega, commit "
        "ffdefdf8035c994ae285a270de0a4ff4e9f336a8. El trabajo de corrección se realiza en el repositorio "
        "principal a partir del checkpoint 6e29b45a85396b5cbe82e8b0ece2b0eb394a7bfd y todavía no fue "
        "sincronizado ni publicado en la entrega.",
    )
    replace(
        document,
        ".\\.venv\\Scripts\\python.exe .\\scouting_app\\create_admin.py",
        "Windows PowerShell:\n$env:ADMIN_USERNAME='<usuario-admin>'\n$env:ADMIN_PASSWORD='<contraseña-segura>'\n.\\.venv\\Scripts\\python.exe .\\scouting_app\\create_admin.py\nLinux/macOS:\nADMIN_USERNAME='<usuario-admin>' ADMIN_PASSWORD='<contraseña-segura>' .venv/bin/python scouting_app/create_admin.py",
    )
    replace(document, "Volver a la raíz del proyecto:", "Volver a la raíz del proyecto:")
    insert_before(find_paragraph(document, "Smoke del despliegue:"), "Set-Location ..", "Código técnico")

    # Remove glossary terms that imply formal Scrum usage and standardize panel terminology.
    remove_glossary_rows(document, {"SCRUM", "Sprint"})
    replace_in_tables(document, "dashboard, comparadores", "panel general, comparadores")
    replace_in_tables(document, "ficha y dashboard", "ficha y panel general")
    replace_in_tables(document, "comparadores, dashboard", "comparadores, panel general")
    replace_in_tables(document, "Dashboard / panel general", "Panel general")
    replace_in_tables(
        document,
        "49aa51c0167fdb24c5f2f3a6ab6e3f397830b462 más cambios locales no publicados",
        "Checkpoint 6e29b45a85396b5cbe82e8b0ece2b0eb394a7bfd más cambios documentales locales del Bloque 7",
    )
    replace_in_tables(document, "Status", "Estado")

    # Replace the bibliography with an audited, alphabetical APA-oriented list.
    bibliography_heading = find_paragraph(document, "8. BIBLIOGRAFÍA")
    annex_heading = find_paragraph(document, "9. ANEXOS TÉCNICOS")
    paragraphs = list(document.paragraphs)
    start = next(i for i, p in enumerate(paragraphs) if p._p is bibliography_heading._p) + 1
    stop = next(i for i, p in enumerate(paragraphs) if p._p is annex_heading._p)
    for paragraph in paragraphs[start:stop]:
        if paragraph._p.xpath("./w:pPr/w:sectPr"):
            paragraph.text = ""
        else:
            delete_paragraph(paragraph)
    references = [
        "Cobley, S., Baker, J., Wattie, N., & McKenna, J. (2009). Annual age-grouping and athlete development: A meta-analytical review of relative age effects in sport. Sports Medicine, 39(3), 235–256. https://doi.org/10.2165/00007256-200939030-00005",
        "Decroos, T., Bransen, L., Van Haaren, J., & Davis, J. (2019). Actions speak louder than goals: Valuing player actions in soccer. Proceedings of the 25th ACM SIGKDD International Conference on Knowledge Discovery & Data Mining, 1851–1861. https://doi.org/10.1145/3292500.3330758",
        "Goodfellow, I., Bengio, Y., & Courville, A. (2016). Deep learning. MIT Press.",
        "Hastie, T., Tibshirani, R., & Friedman, J. (2009). The elements of statistical learning (2nd ed.). Springer. https://doi.org/10.1007/978-0-387-84858-7",
        "Honorable Congreso de la Nación Argentina. (2000). Ley 25.326 de Protección de los Datos Personales. https://www.argentina.gob.ar/normativa/nacional/ley-25326-2000-64790",
        "Hudl. (s. f.). Wyscout. https://www.hudl.com/products/wyscout",
        "Kaufman, S., Rosset, S., Perlich, C., & Stitelman, O. (2012). Leakage in data mining: Formulation, detection, and avoidance. ACM Transactions on Knowledge Discovery from Data, 6(4), Article 15. https://doi.org/10.1145/2382577.2382579",
        "Lacan, S. (2024). Stacking-based deep neural network for player scouting in football. arXiv. https://doi.org/10.48550/arXiv.2403.08835",
        "Niculescu-Mizil, A., & Caruana, R. (2005). Predicting good probabilities with supervised learning. Proceedings of the 22nd International Conference on Machine Learning, 625–632. https://doi.org/10.1145/1102351.1102430",
        "OWASP Foundation. (2021). OWASP Top 10: The ten most critical web application security risks. https://owasp.org/Top10/",
        "Pallets Projects. (s. f.). Flask documentation. https://flask.palletsprojects.com/",
        "Pappalardo, L., Cintia, P., Ferragina, P., Massucco, E., Pedreschi, D., & Giannotti, F. (2019). PlayeRank: Data-driven performance evaluation and player ranking in soccer via a machine learning approach. ACM Transactions on Intelligent Systems and Technology, 10(5), Article 59. https://doi.org/10.1145/3343172",
        "Paszke, A., Gross, S., Massa, F., Lerer, A., Bradbury, J., Chanan, G., Killeen, T., Lin, Z., Gimelshein, N., Antiga, L., Desmaison, A., Köpf, A., Yang, E., DeVito, Z., Raison, M., Tejani, A., Chilamkurthy, S., Steiner, B., Fang, L., Bai, J., & Chintala, S. (2019). PyTorch: An imperative style, high-performance deep learning library. Advances in Neural Information Processing Systems, 32. https://proceedings.neurips.cc/paper/2019/hash/bdbca288fee7f92f2bfa9f7012727740-Abstract.html",
        "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, É. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830. https://www.jmlr.org/papers/v12/pedregosa11a.html",
        "Pressman, R. S. (2010). Ingeniería del software: Un enfoque práctico (7.ª ed.). McGraw-Hill.",
        "PyTorch Foundation. (s. f.). PyTorch documentation. https://docs.pytorch.org/docs/",
        "Saito, T., & Rehmsmeier, M. (2015). The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets. PLOS ONE, 10(3), e0118432. https://doi.org/10.1371/journal.pone.0118432",
        "Schwaber, K., & Sutherland, J. (2020). La Guía Scrum: La guía definitiva de Scrum: Las reglas del juego. https://scrumguides.org/docs/scrumguide/v2020/2020-Scrum-Guide-Spanish-Latin-South-American.pdf",
        "Sommerville, I. (2011). Ingeniería de software (9.ª ed.). Pearson Educación.",
        "SQLAlchemy. (s. f.). SQLAlchemy documentation. https://docs.sqlalchemy.org/",
        "van Arem, K. W., Goes-Smit, F., & Söhl, J. (2025). Forecasting the future development in quality and value of professional football players. Applied Sciences, 15(16), 8916. https://doi.org/10.3390/app15168916",
    ]
    for reference in references:
        paragraph = insert_before(annex_heading, reference)
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        paragraph.paragraph_format.space_after = Inches(0.08)

    document.core_properties.title = "TPScouting — borrador corregido hasta Bloque 7"
    document.core_properties.comments = (
        "Edición de redacción, fuentes y estructura derivada del borrador verificado del Bloque 6. "
        "Índices, campos y revisión visual del PDF permanecen para el Bloque 8."
    )
    document.save(OUTPUT)
    print(f"source={SOURCE}")
    print(f"source_sha256={digest(SOURCE)}")
    print(f"output={OUTPUT}")
    print(f"output_sha256={digest(OUTPUT)}")


if __name__ == "__main__":
    main()
