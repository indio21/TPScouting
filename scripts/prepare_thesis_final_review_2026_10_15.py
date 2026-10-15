"""Prepare a new thesis copy for the final document review."""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "docs/revision_octubre_2026/word/"
    "TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_COMPARACION_ML_2026-10-15.docx"
)
OUTPUT = (
    ROOT
    / "docs/revision_octubre_2026/word/"
    "TRABAJO_FINAL_TPScouting_PRE_REVISION_FINAL_2026-10-15.docx"
)
EXPECTED_SOURCE_SHA256 = "021ad5624f1ae5e6796904ab5afc3ad2682887fd0338846e24c3ff8b3917ed7e"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_runs(paragraph, text: str) -> None:
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def paragraph_starting(document: Document, prefix: str):
    return next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix))


def set_cell(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.style = "Normal"


def main() -> None:
    if digest(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("La fuente no coincide con la copia comparativa verificada.")

    document = Document(SOURCE)

    replace_runs(
        paragraph_starting(document, "La verificación local del 10/10/2026"),
        "La verificación local del 15/10/2026 obtuvo 123 pruebas aprobadas, 1 omitida y "
        "4 advertencias conocidas, con una cobertura de 83,29 %. El smoke visual Playwright "
        "opt-in se ejecutó por separado. Las advertencias corresponden a RuntimeWarning de "
        "scikit-learn por columnas completamente NaN en dos pruebas de preprocesamiento. Las "
        "ejecuciones de CI y del contenedor se vinculan con commits exactos de los repositorios "
        "principal y de entrega.",
    )

    evidence_table = next(table for table in document.tables if table.cell(0, 0).text == "Evidencia")
    evidence_rows = {
        1: ("pytest local (15/10/2026)", "123 passed, 1 skipped, 4 warnings"),
        2: ("pytest-cov local", "83,29 %"),
        5: (
            "Commit funcional del repositorio principal de respaldo",
            "754ccb3601a489ed1249a078ad700d1eb546a494; distribución reproducible verificada",
        ),
        6: (
            "Commit funcional publicado en el repositorio de entrega",
            "8093fe616c9c50b72078ff4f241d9ab3427e763d",
        ),
        7: (
            "CI de las correcciones",
            "Entrega: CI 38656130065 aprobada. Respaldo: CI 38655727560 aprobada.",
        ),
        8: (
            "Alcance",
            (
                "La CI 38656130065 certifica el commit de entrega 8093fe6. Los smoke tests "
                "de contenedor 38656129892 (entrega) y 38655727631 (respaldo) fueron aprobados."
            ),
        ),
    }
    for row_index, (label, result) in evidence_rows.items():
        set_cell(evidence_table.cell(row_index, 0), label)
        set_cell(evidence_table.cell(row_index, 1), result)

    objectives_table = next(
        table for table in document.tables if table.cell(0, 0).text == "Objetivo declarado en 1.2"
    )
    set_cell(
        objectives_table.cell(1, 2),
        "Registro, historiales, paneles, comparadores y PlayerNet; 123 pruebas aprobadas, "
        "1 omitida y 83,29 % de cobertura; CI de entrega 38656130065 aprobada.",
    )

    replace_runs(
        paragraph_starting(document, "La CI #73 se conserva"),
        "La CI #73 se conserva como evidencia histórica del commit "
        "bc5ddd35d0fa3bf6d85faab637772d4e9025fc98 en el repositorio de desarrollo. La "
        "entrega ejecutable más reciente se publicó en el commit "
        "8093fe616c9c50b72078ff4f241d9ab3427e763d y fue certificada por la ejecución CI "
        "38656130065 y el smoke de contenedor 38656129892. La figura siguiente documenta "
        "la ejecución histórica y no sustituye esa evidencia posterior.",
    )

    repository_paragraph = paragraph_starting(document, "Repositorio de entrega:")
    replace_runs(
        repository_paragraph,
        "Repositorio de entrega: https://github.com/indio21/TPScouting-entrega, commit "
        "publicado 8093fe616c9c50b72078ff4f241d9ab3427e763d, validado por la ejecución CI "
        "38656130065 y el smoke de contenedor 38656129892. El repositorio principal conserva "
        "el historial de trabajo y las evidencias internas que no forman parte de la entrega "
        "académica.",
    )

    commands_paragraph = paragraph_starting(document, "Los comandos siguientes se ejecutan")
    rights = commands_paragraph.insert_paragraph_before(
        "La distribución incluye Dockerfile, .dockerignore y .env.example para reproducir el "
        "entorno sin publicar secretos. El repositorio no incluye un archivo LICENSE. Por "
        "decisión de su titular, todos los derechos se encuentran reservados y la entrega no "
        "concede una autorización general de copia, modificación o redistribución.",
        style="Normal",
    )
    rights.paragraph_format.keep_together = True
    replace_runs(
        commands_paragraph,
        "Los comandos siguientes se ejecutan desde la raíz del repositorio clonado. En Windows "
        "y Linux CPU, el runtime se instala desde requirements-lock.txt y PyTorch desde "
        "requirements-torch-cpu.txt; requirements-dev.txt agrega pruebas, auditoría y lint. En "
        "macOS se usa requirements.txt para obtener el wheel estándar de PyTorch y luego "
        "requirements-dev.txt. requirements-docs.txt se reserva para trabajar con el documento. "
        "Se verificó Windows con Python 3.11, Linux con Python 3.11 y 3.12, y macOS Apple Silicon "
        "con Python 3.11. macOS Intel no se declara compatible porque PyTorch 2.14.1 no dispone "
        "de un wheel x86_64 para esa plataforma.",
    )

    # The approved document stores its lists as static text. The expanded
    # evidence table moves its caption by one visible page.
    list_pages = {
        "Tabla 6-2. Resultados de validación automatizada y CI.": 54,
        "Tabla 7-1. Cumplimiento de objetivos.": 73,
    }
    for caption, page in list_pages.items():
        list_entry = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith(f"{caption}\t")
        )
        list_entry.runs[-1].text = f"\t{page}"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    checked = Document(OUTPUT)
    if (len(checked.inline_shapes), len(checked.tables), len(checked.sections)) != (28, 19, 9):
        raise RuntimeError("La copia alteró imágenes, tablas o secciones.")
    full_text = "\n".join(paragraph.text for paragraph in checked.paragraphs)
    required = (
        "123 pruebas aprobadas",
        "8093fe616c9c50b72078ff4f241d9ab3427e763d",
        "38656130065",
        "todos los derechos se encuentran reservados",
        "macOS Apple Silicon con Python 3.11",
    )
    missing = [text for text in required if text not in full_text]
    if missing:
        raise RuntimeError(f"Faltan actualizaciones requeridas: {missing}")

    print(f"OUTPUT={OUTPUT}")
    print(f"SHA256={digest(OUTPUT).upper()}")
    print(
        f"PARAGRAPHS={len(checked.paragraphs)} "
        f"TABLES={len(checked.tables)} IMAGES={len(checked.inline_shapes)}"
    )


if __name__ == "__main__":
    main()
