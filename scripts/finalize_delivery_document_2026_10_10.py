"""Create the delivery copy tied to the verified delivery commit and CI run."""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/revision_octubre_2026/word/TRABAJO_FINAL_TPScouting_FINAL_POSTAUDITORIA_2026-10-10.docx"
OUTPUT = ROOT / "docs/revision_octubre_2026/word/TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_2026-10-10.docx"
EXPECTED_SOURCE_SHA256 = "1326866abba0e04672072ad9078af3b65d9110beafb0acc577c116723a6f8c8a"
DELIVERY_SHA = "72314c2072729c524c0eb4ca57e5e241feb433b6"
DELIVERY_RUN = "38068749412"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_runs(paragraph, text: str) -> None:
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def main() -> None:
    if digest(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("La fuente no coincide con el documento postauditoría verificado.")

    document = Document(SOURCE)
    replacements = {
        452: (
            "La CI #73 se conserva como evidencia histórica del commit "
            "bc5ddd35d0fa3bf6d85faab637772d4e9025fc98 en el repositorio de desarrollo. "
            f"La entrega final se publicó en el commit {DELIVERY_SHA} y fue certificada "
            f"por la ejecución CI {DELIVERY_RUN}, aprobada en Linux con Python 3.11 y "
            "3.12 junto con la auditoría del entorno instalado y del manifest directo. "
            "Esta ejecución incorpora el ajuste de PyTorch 2.14.1 verificado previamente "
            "en el repositorio principal."
        ),
        525: (
            "Repositorio de entrega: https://github.com/indio21/TPScouting-entrega, "
            f"commit publicado {DELIVERY_SHA}, validado por la ejecución CI "
            f"{DELIVERY_RUN}. El repositorio principal conserva el historial de trabajo "
            "y las evidencias internas que no forman parte de la entrega académica."
        ),
    }
    for index, text in replacements.items():
        replace_runs(document.paragraphs[index], text)

    table = document.tables[11]
    table_values = {
        (5, 0): "Commit funcional del repositorio principal de respaldo",
        (5, 1): "6f83e4066db7ccb9900554568a0389cdf1d12d20; PyTorch 2.14.1 y auditoría directa del manifest",
        (6, 0): "Commit funcional publicado en el repositorio de entrega",
        (6, 1): DELIVERY_SHA,
        (7, 1): f"Entrega: CI {DELIVERY_RUN} aprobada. Respaldo funcional: CI 37709526994 aprobada.",
        (8, 1): f"La CI {DELIVERY_RUN} certifica el commit de entrega {DELIVERY_SHA[:7]} con PyTorch 2.14.1.",
    }
    for (row, column), text in table_values.items():
        replace_runs(table.cell(row, column).paragraphs[0], text)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    checked = Document(OUTPUT)
    if (len(checked.inline_shapes), len(checked.tables), len(checked.sections)) != (28, 19, 9):
        raise RuntimeError("La copia alteró la estructura del documento.")
    all_text = "\n".join(p.text for p in checked.paragraphs)
    all_text += "\n" + "\n".join(cell.text for table in checked.tables for row in table.rows for cell in row.cells)
    for obsolete in ("aún no fue sincronizado", "pendiente de una futura sincronización", "7a47bd3f164e865677f2c70075cbedc9fa63427d"):
        if obsolete in all_text:
            raise RuntimeError(f"Texto obsoleto conservado: {obsolete}")
    if all_text.count(DELIVERY_SHA) < 2 or all_text.count(DELIVERY_RUN) < 2:
        raise RuntimeError("La trazabilidad de entrega no quedó asentada.")
    print(OUTPUT)
    print(digest(OUTPUT).upper())


if __name__ == "__main__":
    main()
