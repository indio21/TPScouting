"""Apply the final academic-style edits before Word updates fields."""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/revision_octubre_2026/word/TRABAJO_FINAL_TPScouting_CORREGIDO_BLOQUE7_2026-10-06.docx"
OUTPUT = ROOT / "docs/revision_octubre_2026/word/TRABAJO_FINAL_TPScouting_BLOQUE8_PREFIELDS_2026-10-06.docx"
EXPECTED_SHA256 = "52ab6c6c6201cf783952631904a1a5cb07da088f0903c53d960d39dd077b7dcb"

REPLACEMENTS = {
    "Esta situación plantea una paradoja: mientras el talento juvenil puede convertirse en una fuente de valor deportivo y económico en el mediano plazo, la falta de inversión y de procesos de registro y análisis de información reduce la capacidad de los clubes para identificar oportunidades, planificar la formación y sostener decisiones basadas en evidencia.":
        "Esta situación plantea una paradoja. El talento juvenil puede convertirse en una fuente de valor deportivo y económico en el mediano plazo. Sin embargo, la falta de inversión y de procesos de registro y análisis reduce la capacidad de los clubes para identificar oportunidades, planificar la formación y sostener decisiones basadas en evidencia.",
    "La literatura reciente también explora modelos de aprendizaje automático para detectar perfiles de alto potencial y pronosticar la evolución futura de los jugadores, aunque sus resultados dependen de la calidad de los datos, la definición del objetivo y el contexto de aplicación (Lacan, 2024; van Arem et al., 2025).":
        "La literatura reciente también explora modelos de aprendizaje automático para detectar perfiles de alto potencial y pronosticar la evolución futura de los jugadores. Sus resultados dependen de la calidad de los datos, la definición del objetivo y el contexto de aplicación (Lacan, 2024; van Arem et al., 2025).",
    "Se utilizan datos sintéticos para validar el flujo completo de captura, persistencia, entrenamiento e inferencia, ya que esta decisión favorece la reproducibilidad y evita depender de bases privadas de clubes o proveedores externos durante la evaluación académica, pero limita la validez externa de los resultados.":
        "Se utilizan datos sintéticos para validar el flujo completo de captura, persistencia, entrenamiento e inferencia. Esta decisión favorece la reproducibilidad y evita depender de bases privadas de clubes o proveedores externos durante la evaluación académica, pero limita la validez externa de los resultados.",
    "La aplicación fue verificada con datos sintéticos y no se utilizó con menores reales. Un uso real debería definir una base jurídica y una finalidad específica, informar a los titulares y a sus representantes, obtener el consentimiento que corresponda, recolectar sólo datos necesarios, aplicar plazos de retención y mecanismos de acceso, rectificación y supresión, y limitar el acceso por rol. La Ley 25.326 regula en Argentina los principios de protección de datos, los derechos de los titulares y las obligaciones de responsables de archivos y bancos de datos (Honorable Congreso de la Nación Argentina, 2000). Estas medidas corresponden a una etapa posterior; el MVP no implementa un flujo completo de consentimiento de tutores ni una política automatizada de retención y baja.":
        "La aplicación fue verificada con datos sintéticos y no se utilizó con menores reales. Un uso real debería definir una base jurídica y una finalidad específica. También debería informar a los titulares y a sus representantes, obtener el consentimiento que corresponda, recolectar sólo los datos necesarios, aplicar plazos de retención y mecanismos de acceso, rectificación y supresión, y limitar el acceso según el rol. La Ley 25.326 regula en Argentina los principios de protección de datos, los derechos de los titulares y las obligaciones de responsables de archivos y bancos de datos (Honorable Congreso de la Nación Argentina, 2000). Estas medidas corresponden a una etapa posterior; el MVP no implementa un flujo completo de consentimiento de tutores ni una política automatizada de retención y baja.",
    "•   README.md: instalación, pruebas, deploy y limitaciones.":
        "• README.md: instalación, pruebas, despliegue y limitaciones.",
    "Tabla 4-1. Capas": "Tabla 4-1. Capas y responsabilidades del MVP.",
    "Tabla 4-4. Comparación entre motores relacionales": "Tabla 4-2. Comparación entre motores relacionales.",
    "Tabla 4-2. Dimensiones": "Tabla 4-3. Dimensiones de evaluación deportiva.",
    "Tabla 4-3. Requisitos": "Tabla 4-4. Requisitos funcionales del sistema.",
    "Tabla 4-4. Entidades": "Tabla 4-5. Entidades principales del modelo de datos.",
    "Tabla 4-5. Endpoints": "Tabla 4-6. Endpoints principales del MVP.",
    "Tabla 4-6. Decisiones": "Tabla 4-7. Decisiones tecnológicas.",
    "Tabla 4-7. Casos": "Tabla 4-8. Casos de uso principales.",
    "El modelo implementa once entidades: Player, Match, Coach, Director, User, PlayerStat, PlayerAttributeHistory, PlayerMatchParticipation, ScoutReport, PhysicalAssessment y PlayerAvailability. Sus responsabilidades se resumen en la Tabla 4-4.":
        "El modelo implementa once entidades: Player, Match, Coach, Director, User, PlayerStat, PlayerAttributeHistory, PlayerMatchParticipation, ScoutReport, PhysicalAssessment y PlayerAvailability. Sus responsabilidades se resumen en la Tabla 4-5.",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_static_list(document: Document, heading: str, next_heading: str, captions: list[str]) -> None:
    paragraphs = list(document.paragraphs)
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == heading)
    stop = next(i for i, p in enumerate(paragraphs) if i > start and p.text.strip() == next_heading)
    anchor = paragraphs[start]._p
    for paragraph in paragraphs[start + 1 : stop]:
        paragraph._p.getparent().remove(paragraph._p)
    for caption in captions:
        element = OxmlElement("w:p")
        anchor.addnext(element)
        paragraph = Paragraph(element, document._body)
        paragraph.style = "toc 1"
        paragraph.add_run(f"{caption}\t0")
        anchor = element


def main() -> None:
    if sha256(SOURCE) != EXPECTED_SHA256:
        raise RuntimeError("El documento del Bloque 7 no coincide con el hash verificado.")
    document = Document(SOURCE)
    changed = set()
    for paragraph in document.paragraphs:
        replacement = REPLACEMENTS.get(paragraph.text)
        if replacement is not None:
            paragraph.text = replacement
            changed.add(paragraph.text)
    if len(changed) != len(REPLACEMENTS):
        raise RuntimeError(f"Se aplicaron {len(changed)} de {len(REPLACEMENTS)} reemplazos esperados.")

    document.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"SHA256={sha256(OUTPUT)}")


if __name__ == "__main__":
    main()
