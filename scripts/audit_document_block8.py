"""Audit the Block 8 thesis candidate without changing it."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DOCX = ROOT / "docs/revision_octubre_2026/word/TRABAJO_FINAL_TPScouting_REVISION_TEXTO_SIN_IMAGENES_2026-10-06.docx"
DEFAULT_REPORT = ROOT / "docs/revision_octubre_2026/auditoria_redaccion_bloque8.md"

SUSPICIOUS = {
    "doble espacio": re.compile(r"(?<! ) {2,}"),
    "espacio antes de puntuacion": re.compile(r"\s+[,:;.](?:\s|$)"),
    "primera persona plural": re.compile(r"\b(nosotros|nuestro|nuestra|hemos|realizamos|desarrollamos)\b", re.IGNORECASE),
    "segunda persona": re.compile(r"\b(tu|tus|usted|ustedes|podras|puedes)\b", re.IGNORECASE),
    "lenguaje promocional": re.compile(r"\b(revolucionari[oa]s?|innovador(?:a|as|es)?|garantiza|excelente|potente|de vanguardia)\b", re.IGNORECASE),
    "pendiente editorial": re.compile(r"\b(FIXME|PENDIENTE DE INFORMAR)\b", re.IGNORECASE),
}


def iter_text(document: Document):
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text and not paragraph.style.name.lower().startswith("toc"):
            yield f"párrafo {index}", text
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                if text:
                    yield f"tabla {table_index}, fila {row_index}, celda {cell_index}", text


def sentence_candidates(location: str, text: str):
    for sentence in re.split(r"(?<=[.!?])\s+", text):
        words = re.findall(r"\b[\wÁÉÍÓÚÜÑáéíóúüñ]+\b", sentence)
        if len(words) >= 45:
            yield location, len(words), sentence


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()

    docx_path = args.docx.resolve()
    report_path = args.report.resolve()
    document = Document(docx_path)
    entries = list(iter_text(document))
    findings: dict[str, list[tuple[str, str]]] = {name: [] for name in SUSPICIOUS}
    long_sentences = []
    duplicate_words = []

    for location, text in entries:
        for name, pattern in SUSPICIOUS.items():
            if pattern.search(text):
                findings[name].append((location, text))
        duplicate = re.search(r"\b([A-Za-zÁÉÍÓÚÜÑáéíóúüñ]{3,})\s+\1\b", text, re.IGNORECASE)
        if duplicate:
            duplicate_words.append((location, duplicate.group(0), text))
        long_sentences.extend(sentence_candidates(location, text))

    body = [
        "# Auditoría reproducible de redacción — Bloque 8",
        "",
        f"Documento: `{docx_path.relative_to(ROOT)}`",
        f"Párrafos y celdas examinados: {len(entries)}.",
        "",
        "Este control detecta candidatos para revisión humana; una coincidencia no equivale por sí sola a un error.",
        "",
        "## Resumen",
        "",
        "| Control | Coincidencias |",
        "|---|---:|",
    ]
    for name, items in findings.items():
        body.append(f"| {name} | {len(items)} |")
    body.append(f"| palabra consecutiva repetida | {len(duplicate_words)} |")
    body.append(f"| oraciones de 45 palabras o más | {len(long_sentences)} |")

    for name, items in findings.items():
        body.extend(["", f"## {name}", ""])
        if not items:
            body.append("Sin coincidencias.")
        else:
            for location, text in items:
                body.append(f"- **{location}:** {text}")

    body.extend(["", "## Palabras consecutivas repetidas", ""])
    if not duplicate_words:
        body.append("Sin coincidencias.")
    else:
        for location, match, text in duplicate_words:
            body.append(f"- **{location} ({match}):** {text}")

    body.extend(["", "## Oraciones extensas", ""])
    if not long_sentences:
        body.append("Sin coincidencias.")
    else:
        for location, count, sentence in long_sentences:
            body.append(f"- **{location}, {count} palabras:** {sentence}")

    report_path.write_text("\n".join(body) + "\n", encoding="utf-8")
    print(f"REPORT={report_path}")
    print("COUNTS=" + repr({name: len(items) for name, items in findings.items()}))
    print(f"DUPLICATES={len(duplicate_words)} LONG_SENTENCES={len(long_sentences)}")


if __name__ == "__main__":
    main()
