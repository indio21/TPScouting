"""Assemble the reviewed text with one final set of regenerated images."""

from __future__ import annotations

import hashlib
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/revision_octubre_2026/word/TRABAJO_FINAL_TPScouting_TEXTO_REVISADO_SIN_IMAGENES_2026-10-06.docx"
SOURCE_WITH_LOGO = ROOT / "docs/revision_octubre_2026/word/TRABAJO_FINAL_TPScouting_CANDIDATO_FINAL_BLOQUE8_v4_2026-10-06.docx"
OUTPUT = ROOT / "docs/revision_octubre_2026/word/TRABAJO_FINAL_TPScouting_CONTENIDO_FINAL_CON_IMAGENES_2026-10-06.docx"
ASSET_DIR = ROOT / "docs/revision_octubre_2026/recursos_finales_bloque8"
DIAGRAM_DIR = ROOT / "docs/diagramas/export"
CAPTURE_DIR = ROOT / "docs/evidencia_word_render/bloque8_2026-10-06"
EVIDENCE_DIR = ROOT / "docs/evidencia_word_render"
EXPECTED_SHA256 = "d889a26c5ee4a0cd26070d5f232afbb5934d7ebc20830f84f843b3acb34ec072"

OLD_CAPTURE_TEXT = (
    "Las Figuras 6-2 a 6-5 y 6-7 fueron capturadas desde el despliegue público en Render mediante una sesión "
    "autenticada. La Figura 6-6 fue reproducida el 21 de agosto de 2026 en un navegador real contra una base "
    "SQLite temporal aislada, después de corregir el mapeo de current_age en la vista de predicción. Cada captura "
    "se interpreta como evidencia fechada, no como garantía de disponibilidad permanente."
)
NEW_CAPTURE_TEXT = (
    "Las Figuras 6-2 a 6-7 fueron reproducidas el 6 de octubre de 2026 en un navegador real contra una base "
    "SQLite temporal con datos sintéticos y semilla 42. La sesión autenticada verificó el acceso, el panel general, "
    "el listado, la ficha, la predicción y el comparador múltiple. La base temporal se eliminó al finalizar; estas "
    "capturas no demuestran disponibilidad continua del despliegue público."
)
OLD_PREDICTION_TEXT = (
    "La Figura 6-6 presenta una reproducción local auditada de la vista de predicción, con la edad derivada de "
    "birth_date y el ajuste combinado respecto de PlayerNet crudo. Ese ajuste incorpora tanto el historial disponible "
    "como el ajuste posicional; la base temporal utilizada para la captura fue eliminada al finalizar."
)
NEW_PREDICTION_TEXT = (
    "La Figura 6-6 presenta la vista de predicción reproducida en la verificación local del 6 de octubre de 2026, "
    "con la edad derivada de birth_date y el ajuste combinado respecto de PlayerNet crudo. Ese ajuste incorpora el "
    "historial disponible y la adecuación posicional."
)

IMAGES = [
    ("Figura 4-1. Diagrama de componentes de TPScouting.", DIAGRAM_DIR / "03_componentes.png", 6.0),
    ("Figura 4-2. Diagrama de clases del modelo de dominio.", DIAGRAM_DIR / "05_clases.png", 5.9),
    ("Figura 4-3. Diagrama de casos de uso general", DIAGRAM_DIR / "06_casos_uso_general.png", 5.9),
    ("Figura 4-4. Diagrama de casos de uso de gestión de jugadores.", DIAGRAM_DIR / "07_casos_uso_gestion_jugadores.png", 4.25),
    ("Figura 4-5. Diagrama de casos de uso de análisis y decisión scout.", DIAGRAM_DIR / "08_casos_uso_analisis_decision.png", 6.0),
    ("Figura 5-1. Secuencia real de inferencia: score combinado principal y calibración secundaria.", DIAGRAM_DIR / "01_secuencia_prediccion.png", 6.0),
    ("Figura 5-2. Secuencia de carga del panel general.", DIAGRAM_DIR / "02_secuencia_dashboard.png", 6.0),
    ("Figura 5-3. Arquitectura de despliegue en Render y mecanismos operativos del MVP.", DIAGRAM_DIR / "04_despliegue.png", 6.0),
    ("Figura 6-1. Curva real de entrenamiento generada desde training_metadata.json.", EVIDENCE_DIR / "07_training_curve_real.png", 5.75),
    ("Figura 6-2. Login del MVP en la verificación local.", CAPTURE_DIR / "01_login.png", 6.0),
    ("Figura 6-3. Panel general con métricas por rol.", CAPTURE_DIR / "02_dashboard.png", 6.0),
    ("Figura 6-4. Listado paginado de jugadores con edad, categoría y potencial.", CAPTURE_DIR / "03_players.png", 6.0),
    ("Figura 6-5. Ficha de jugador con historial, atributos y acciones CRUD en modales.", CAPTURE_DIR / "04_player_detail.png", 6.0),
    ("Figura 6-6. Vista de predicción con edad y ajuste combinado de historial y posición.", CAPTURE_DIR / "05_prediction.png", 6.0),
    ("Figura 6-7. Comparador múltiple de jugadores.", CAPTURE_DIR / "06_compare_multi.png", 6.0),
    ("Figura 6-8. Resumen público de la ejecución CI #73 en GitHub Actions.", EVIDENCE_DIR / "08_ci_run_73_resumen.png", 5.75),
    ("Figura 10-1. Evidencia histórica de la CI #73 del repositorio de desarrollo.", EVIDENCE_DIR / "09_ci_run_73_ampliada.png", 3.8),
    ("Figura 10-2. Diagrama de componentes de TPScouting ampliado.", DIAGRAM_DIR / "03_componentes.png", 6.0),
    ("Figura 10-3. Diagrama de clases del modelo de dominio ampliado.", DIAGRAM_DIR / "05_clases.png", 6.0),
    ("Figura 10-4. Diagrama de secuencia de inferencia ampliado.", DIAGRAM_DIR / "01_secuencia_prediccion.png", 6.0),
    ("Figura 10-5. Diagrama de secuencia del panel general ampliado.", DIAGRAM_DIR / "02_secuencia_dashboard.png", 6.0),
    ("Figura 10-6. Diagrama de despliegue en Render ampliado.", DIAGRAM_DIR / "04_despliegue.png", 6.0),
    ("Figura 10-7. Diagrama de casos de uso general ampliado.", DIAGRAM_DIR / "06_casos_uso_general.png", 6.0),
    ("Figura 10-8. Diagrama de casos de uso de gestión de jugadores ampliado.", DIAGRAM_DIR / "07_casos_uso_gestion_jugadores.png", 4.65),
    ("Figura 10-9. Diagrama de casos de uso de análisis y decisión scout ampliado.", DIAGRAM_DIR / "08_casos_uso_analisis_decision.png", 6.0),
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(text: str) -> str:
    return " ".join((text or "").split())


def exact(document: Document, text: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if norm(paragraph.text) == norm(text)]
    if len(matches) != 1:
        raise RuntimeError(f"Se esperaba una coincidencia para {text!r}; se hallaron {len(matches)}.")
    return matches[0]


def insert_before(document: Document, caption: Paragraph, image: Path, width: float) -> None:
    if not image.exists():
        raise FileNotFoundError(image)
    previous = caption._p.getprevious()
    paragraph = None
    if previous is not None and previous.tag.endswith("}p"):
        candidate = Paragraph(previous, document._body)
        if not norm(candidate.text) and not candidate._p.xpath(".//w:drawing | .//w:pict"):
            paragraph = candidate
    if paragraph is None:
        element = OxmlElement("w:p")
        caption._p.addprevious(element)
        paragraph = Paragraph(element, document._body)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image), width=Inches(width))


def compact_glossary_table(document: Document) -> None:
    """Recover enough space to keep the glossary caption with its table."""
    table = document.tables[-1]
    if "Glosario" not in table.cell(0, 0).text and "Término" not in table.cell(0, 0).text:
        raise RuntimeError("No se pudo identificar la tabla final del glosario.")
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 0.95
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)


def main() -> None:
    if sha256(SOURCE) != EXPECTED_SHA256:
        raise RuntimeError("El texto revisado no coincide con el hash esperado.")
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    logo = ASSET_DIR / "logo_ucse.png"
    with ZipFile(SOURCE_WITH_LOGO) as archive:
        logo.write_bytes(archive.read("word/media/image1.png"))

    document = Document(SOURCE)
    exact(document, OLD_CAPTURE_TEXT).text = NEW_CAPTURE_TEXT
    exact(document, OLD_PREDICTION_TEXT).text = NEW_PREDICTION_TEXT
    exact(document, "Figura 6-2. Login del MVP desplegado en Render.").text = "Figura 6-2. Login del MVP en la verificación local."
    compact_glossary_table(document)

    cover = document.paragraphs[0]
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_logo = cover.add_run().add_picture(str(logo), width=Inches(6.135))
    cover_logo.height = Inches(2.1844)

    for caption_text, image, width in IMAGES:
        insert_before(document, exact(document, caption_text), image, width)

    document.save(OUTPUT)
    if len(Document(OUTPUT).inline_shapes) != len(IMAGES) + 1:
        raise RuntimeError("La cantidad final de imágenes no coincide con el inventario de ensamblado.")
    print(f"OUTPUT={OUTPUT}")
    print(f"IMAGES={len(IMAGES) + 1}")
    print(f"SHA256={sha256(OUTPUT)}")


if __name__ == "__main__":
    main()
