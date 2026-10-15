"""Add a bounded historical-versus-candidate ML comparison to the thesis."""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "docs/revision_octubre_2026/word/"
    "TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_ACTUALIZADA_2026-10-15.docx"
)
OUTPUT = (
    ROOT
    / "docs/revision_octubre_2026/word/"
    "TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_COMPARACION_ML_2026-10-15.docx"
)
EXPECTED_SOURCE_SHA256 = "b4a80877ffe0e8688bd53140eb7a6e6f82aac753e5db60553b7e2c4aaa07a0f9"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_runs(paragraph, text: str) -> None:
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def main() -> None:
    if digest(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("La fuente no coincide con la copia documental verificada.")

    document = Document(SOURCE)
    anchor = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("La validez está limitada por una prevalencia")
    )

    comparison = (
        "La comparación entre la corrida histórica y la candidata se limita a la salida cruda de "
        "PlayerNet sobre sus respectivos conjuntos de test. La corrida histórica obtuvo ROC-AUC "
        "0,9203, PR-AUC 0,5461, Brier 0,1049, log loss 0,3294 y F1 0,5291. La corrida candidata "
        "20261010T_local_seed42_leakage_safe_v2 obtuvo "
        "ROC-AUC 0,9056, PR-AUC 0,4888, Brier 0,1175, log loss 0,3673 y F1 0,4931. En términos "
        "descriptivos, la corrida histórica presenta mejores valores en esas cinco métricas."
    )
    interpretation = (
        "La diferencia no permite atribuir una pérdida de desempeño únicamente al modelo, porque las "
        "corridas no comparten exactamente la misma definición del target. La histórica calculó "
        "cuantiles y cuotas antes del split; la candidata ajustó la política solo con train y la aplicó "
        "sin recalcularla en validation y test. La candidata aporta un procedimiento metodológicamente "
        "más seguro y evidencia reproducible, pero no demuestra una mejora predictiva ni justifica "
        "reemplazar los artefactos utilizados por la aplicación. Por ello se conserva como comparación "
        "experimental y no como nuevo resultado oficial del MVP."
    )
    anchor.insert_paragraph_before(comparison, style="Normal").paragraph_format.keep_together = True
    anchor.insert_paragraph_before(interpretation, style="Normal").paragraph_format.keep_together = True

    replace_runs(
        anchor,
        "La validez de la corrida histórica está limitada por una prevalencia fijada por diseño, "
        "cuantiles calculados antes del split, una sola semilla y ausencia de jugadores reales. La "
        "corrida candidata corrige el ajuste de cuantiles y registra los metadatos faltantes, aunque "
        "conserva las limitaciones de una sola semilla y datos sintéticos. El MVP sí aporta evidencia "
        "de integración, persistencia, control de acceso y presentación de resultados. Una validación "
        "externa requeriría datos longitudinales reales, consentimiento y criterios definidos antes "
        "del experimento.",
    )

    # The approved document stores figure and table lists as static text. Keep
    # the entries aligned with the one-page shift introduced by this analysis.
    list_pages = {
        "Figura 10-1. Evidencia histórica de la CI #73 del repositorio de desarrollo.": 85,
        "Figura 10-2. Diagrama de componentes de TPScouting ampliado.": 87,
        "Figura 10-3. Diagrama de clases del modelo de dominio ampliado.": 88,
        "Figura 10-4. Diagrama de secuencia de inferencia ampliado.": 89,
        "Figura 10-5. Diagrama de secuencia del panel general ampliado.": 90,
        "Figura 10-6. Diagrama de despliegue en Render ampliado.": 91,
        "Figura 10-7. Diagrama de casos de uso general ampliado.": 93,
        "Figura 10-8. Diagrama de casos de uso de gestión de jugadores ampliado.": 94,
        "Figura 10-9. Diagrama de casos de uso de análisis y decisión scout ampliado.": 95,
        "Tabla 7-1. Cumplimiento de objetivos.": 72,
        "Tabla 9-1. Rutas principales.": 80,
        "Tabla 9-2. Variables de entorno relevantes.": 80,
        "Tabla 9-3. Glosario técnico.": 81,
    }
    updated_lists = set()
    for paragraph in document.paragraphs:
        if "\t" not in paragraph.text:
            continue
        caption, _, _ = paragraph.text.rpartition("\t")
        if caption not in list_pages:
            continue
        paragraph.runs[-1].text = f"\t{list_pages[caption]}"
        updated_lists.add(caption)
    if updated_lists != set(list_pages):
        missing = sorted(set(list_pages) - updated_lists)
        raise RuntimeError(f"No se pudieron actualizar estas entradas de listas: {missing}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    checked = Document(OUTPUT)
    if (len(checked.inline_shapes), len(checked.tables), len(checked.sections)) != (28, 19, 9):
        raise RuntimeError("La copia alteró imágenes, tablas o secciones.")
    full_text = "\n".join(paragraph.text for paragraph in checked.paragraphs)
    required = (
        "La comparación entre la corrida histórica y la candidata",
        "ROC-AUC 0,9056",
        "procedimiento metodológicamente más seguro",
        "no como nuevo resultado oficial del MVP",
    )
    missing = [text for text in required if text not in full_text]
    if missing:
        raise RuntimeError(f"Faltan textos de comparación: {missing}")

    print(f"OUTPUT={OUTPUT}")
    print(f"SHA256={digest(OUTPUT).upper()}")
    print(
        f"PARAGRAPHS={len(checked.paragraphs)} "
        f"TABLES={len(checked.tables)} IMAGES={len(checked.inline_shapes)}"
    )


if __name__ == "__main__":
    main()
