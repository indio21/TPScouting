from __future__ import annotations

import sys
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph


sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_CORREGIDA_PAGINADA_21-08-2026.docx"
)
OUTPUT = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_CORREGIDA_PAGINADA_CON_DIAGRAMAS_21-08-2026.docx"
)
DIAGRAM_DIR = ROOT / "docs" / "diagramas" / "export"


DIAGRAMS = [
    # Figuras dentro de los capitulos.
    ("Figura 4-1. Diagrama de componentes de TPScouting.", "03_componentes.png", 6.35),
    ("Figura 4-2. Diagrama de clases del modelo de dominio.", "05_clases.png", 6.00),
    ("Figura 4-3. Diagrama de casos de uso general.", "06_casos_uso_general.png", 5.90),
    ("Figura 4-4. Diagrama de casos de uso de gestión de jugadores.", "07_casos_uso_gestion_jugadores.png", 4.25),
    ("Figura 4-5. Diagrama de casos de uso de análisis y decisión scout.", "08_casos_uso_analisis_decision.png", 6.00),
    (
        "Figura 5-1. Secuencia real de inferencia: score combinado principal y calibración secundaria.",
        "01_secuencia_prediccion.png",
        6.00,
    ),
    ("Figura 5-2. Secuencia de carga del panel general.", "02_secuencia_dashboard.png", 6.00),
    (
        "Figura 5-3. Arquitectura de despliegue en Render y mecanismos operativos del MVP.",
        "04_despliegue.png",
        6.00,
    ),
    # Figuras ampliadas del anexo, renumeradas segun el orden real del Word corregido.
    ("Figura 10-2. Diagrama de componentes de TPScouting ampliado.", "03_componentes.png", 6.00),
    ("Figura 10-3. Diagrama de clases del modelo de dominio ampliado.", "05_clases.png", 6.00),
    ("Figura 10-4. Diagrama de secuencia de inferencia ampliado.", "01_secuencia_prediccion.png", 6.00),
    ("Figura 10-5. Diagrama de secuencia del panel general ampliado.", "02_secuencia_dashboard.png", 6.00),
    ("Figura 10-6. Diagrama de despliegue en Render ampliado.", "04_despliegue.png", 6.00),
    ("Figura 10-7. Diagrama de casos de uso general ampliado.", "06_casos_uso_general.png", 6.00),
    (
        "Figura 10-8. Diagrama de casos de uso de gestión de jugadores ampliado.",
        "07_casos_uso_gestion_jugadores.png",
        4.65,
    ),
    (
        "Figura 10-9. Diagrama de casos de uso de análisis y decisión scout ampliado.",
        "08_casos_uso_analisis_decision.png",
        6.00,
    ),
]


def norm(text: str) -> str:
    return " ".join((text or "").split())


def find_caption(doc: Document, caption: str) -> Paragraph:
    matches = [paragraph for paragraph in doc.paragraphs if norm(paragraph.text) == caption]
    if len(matches) != 1:
        raise ValueError(f"Se esperaba un caption {caption!r}; se encontraron {len(matches)}")
    return matches[0]


def paragraph_has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def find_existing_drawing_before(doc: Document, caption: Paragraph) -> Paragraph | None:
    node = caption._p.getprevious()
    # Las versiones sin imagen pueden conservar uno o mas parrafos vacios.
    # Se revisa todo el bloque hasta llegar al texto introductorio anterior.
    while node is not None and node.tag.endswith("}p"):
        paragraph = Paragraph(node, doc._body)
        if paragraph_has_drawing(paragraph):
            return paragraph
        if norm(paragraph.text):
            return None
        node = node.getprevious()
    return None


def add_image_before(doc: Document, caption: Paragraph, image_path: Path, width: float) -> None:
    if find_existing_drawing_before(doc, caption) is not None:
        return

    previous = caption._p.getprevious()
    if previous is not None and previous.tag.endswith("}p"):
        candidate = Paragraph(previous, doc._body)
        if not norm(candidate.text) and not paragraph_has_drawing(candidate):
            image_paragraph = candidate
        else:
            image_paragraph = None
    else:
        image_paragraph = None

    if image_paragraph is None:
        new_p = OxmlElement("w:p")
        caption._p.addprevious(new_p)
        image_paragraph = Paragraph(new_p, doc._body)

    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(width))


def main() -> int:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    for _, image_name, _ in DIAGRAMS:
        path = DIAGRAM_DIR / image_name
        if not path.exists():
            raise FileNotFoundError(path)

    copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    initial_shapes = len(doc.inline_shapes)
    missing_before = sum(
        find_existing_drawing_before(doc, find_caption(doc, caption_text)) is None
        for caption_text, _, _ in DIAGRAMS
    )

    for caption_text, image_name, width in DIAGRAMS:
        caption = find_caption(doc, caption_text)
        add_image_before(doc, caption, DIAGRAM_DIR / image_name, width)

    inserted = len(doc.inline_shapes) - initial_shapes
    if inserted != missing_before:
        raise RuntimeError(f"Se esperaban {missing_before} diagramas nuevos; se insertaron {inserted}")

    present_after = sum(
        find_existing_drawing_before(doc, find_caption(doc, caption_text)) is not None
        for caption_text, _, _ in DIAGRAMS
    )
    if present_after != len(DIAGRAMS):
        raise RuntimeError(f"Solo {present_after} de {len(DIAGRAMS)} captions quedaron con diagrama")

    doc.save(OUTPUT)
    print(f"Documento con diagramas: {OUTPUT}")
    print(f"Diagramas insertados: {inserted}")
    print(f"Diagramas reutilizados desde el Word paginado: {len(DIAGRAMS) - missing_before}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
