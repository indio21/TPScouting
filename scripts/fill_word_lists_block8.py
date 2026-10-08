"""Replace broken Word figure/table fields with verified static lists."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def pdf_page_map(pdf_path: Path, captions: list[str]) -> dict[str, int]:
    import fitz

    document = fitz.open(pdf_path)
    raw_pages = [page.get_text() for page in document]
    pages = [normalized(text) for text in raw_pages]

    def displayed_page_number(index: int) -> int:
        matches = re.findall(r"(?m)^\s*(\d+)\s*$", raw_pages[index])
        if not matches:
            raise RuntimeError(f"No se halló el número visible en la página física {index + 1}")
        return int(matches[0])
    result = {}
    for caption in captions:
        needle = normalized(caption)
        # Las listas ocupan el frente del documento; se buscan los títulos en el
        # cuerpo para no confundir una entrada de lista con su figura o tabla.
        matches = [index + 1 for index, text in enumerate(pages) if needle in text]
        if len(matches) < 2:
            raise RuntimeError(
                f"El título debe aparecer en la lista y en el cuerpo: {caption!r}; hallado en {matches}"
            )
        # La lista está en el frente y la leyenda real aparece después. Tomar la
        # última coincidencia evita depender de una cantidad fija de páginas
        # preliminares cuando cambia el tamaño tipográfico.
        result[caption] = displayed_page_number(matches[-1] - 1)
    return result


def replace_list(document: Document, heading: str, next_heading: str, captions: list[str], pages: dict[str, int]) -> None:
    paragraphs = list(document.paragraphs)
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == heading)
    stop = next(i for i, p in enumerate(paragraphs) if i > start and p.text.strip() == next_heading)
    anchor = paragraphs[start]._p
    for paragraph in paragraphs[start + 1 : stop]:
        paragraph._p.getparent().remove(paragraph._p)
    for caption in captions:
        element = OxmlElement("w:p")
        anchor.addnext(element)
        paragraph = Paragraph(element, document._body)
        paragraph.style = "toc 1"
        paragraph.add_run(f"{caption}\t{pages.get(caption, 0)}")
        anchor = element


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--pdf", type=Path)
    args = parser.parse_args()

    document = Document(args.source)
    figures = [p.text.strip() for p in document.paragraphs if p.style.name == "Figure Caption Generated" and p.text.strip()]
    tables = [p.text.strip() for p in document.paragraphs if p.style.name == "Table Caption Generated" and p.text.strip()]
    captions = figures + tables
    pages = pdf_page_map(args.pdf, captions) if args.pdf else {}
    replace_list(document, "LISTA DE TABLAS", "RESUMEN", tables, pages)
    replace_list(document, "LISTA DE FIGURAS", "LISTA DE TABLAS", figures, pages)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(f"FIGURES={len(figures)} TABLES={len(tables)} PAGES_MAPPED={len(pages)}")
    print(f"OUTPUT={args.output}")


if __name__ == "__main__":
    main()
