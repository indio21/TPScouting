from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


SOURCE = Path(
    r"C:\Tesis\TPScouting\docs\tesis_final\TRABAJO_FINAL_TPScouting_26-08-2026_v2_audit_source.docx"
)
OUTPUT = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_REVISADA_26-08-2026_v3.docx"
)

SUMMARY = (
    "El seguimiento de futbolistas juveniles en clubes con recursos limitados suele depender de registros "
    "dispersos y criterios difíciles de comparar a lo largo del tiempo. A partir de esta problemática, en "
    "este trabajo final se diseñó y desarrolló TPScouting, un producto mínimo viable orientado a organizar y "
    "analizar información deportiva de jugadores de entre 12 y 18 años. La aplicación permite gestionar "
    "perfiles, atributos técnicos, estadísticas, evaluaciones físicas, reportes, disponibilidad e historiales; "
    "además, incorpora un panel general, herramientas de comparación y una estimación de potencial basada en "
    "un modelo implementado en PyTorch. Para evaluar la solución se utilizaron datos sintéticos, pruebas "
    "automatizadas, métricas de clasificación, comparación con modelos de referencia y evidencia fechada del "
    "funcionamiento de la aplicación y de su integración continua. Los resultados permitieron comprobar la "
    "viabilidad técnica del flujo completo, desde el registro de datos hasta la presentación de una estimación "
    "en la interfaz. La comparación experimental también mostró que PlayerNet no alcanza una superioridad "
    "global frente a la regresión logística en la corrida analizada, por lo que la complejidad del modelo no se "
    "considera un aporte por sí misma. Debido al carácter sintético de los datos y a la ausencia de validación "
    "longitudinal en un club, los resultados no demuestran capacidad predictiva en contextos deportivos reales. "
    "TPScouting se plantea, por tanto, como una herramienta de apoyo para ordenar información, hacer más "
    "trazables los criterios de evaluación y complementar, sin sustituir, el análisis de entrenadores y scouts."
)

CONCLUSIONS = [
    (
        "A partir del desarrollo realizado, se concluye que fue posible construir un producto mínimo viable "
        "funcional que integra en una misma aplicación el registro de jugadores, el seguimiento de atributos e "
        "historiales, la consulta de estadísticas y evaluaciones, la visualización de indicadores, la comparación "
        "de perfiles y un módulo de aprendizaje automático. La solución permitió recorrer el flujo completo, "
        "desde la carga y persistencia de los datos hasta la presentación de una estimación en la interfaz, con "
        "pruebas automatizadas y evidencia técnica de su funcionamiento. Este resultado demuestra la factibilidad "
        "de la integración propuesta dentro del alcance de un MVP, pero no equivale a validar su eficacia en un "
        "proceso deportivo real."
    ),
    (
        "El aporte más concreto del proyecto se encuentra en la organización y trazabilidad de información que, "
        "en clubes con recursos limitados, puede permanecer dispersa entre planillas, observaciones y registros "
        "aislados. La ficha integral, los historiales, el panel general y los comparadores permiten consultar la "
        "evolución de cada jugador y aplicar criterios consistentes al análisis. De esta manera, TPScouting ofrece "
        "una base técnica para apoyar decisiones de seguimiento y formación. Sin embargo, el trabajo no midió una "
        "reducción efectiva de la subjetividad, del tiempo de evaluación o de los costos operativos, por lo que "
        "esos beneficios deben considerarse posibilidades del sistema y no resultados ya demostrados."
    ),
    (
        "En relación con el componente de aprendizaje automático, se logró integrar el preprocesamiento, los "
        "artefactos del modelo, la calibración y la inferencia dentro de la aplicación. Las métricas obtenidas con "
        "datos sintéticos permitieron analizar el comportamiento de PlayerNet y compararlo con modelos de "
        "referencia. La regresión logística obtuvo resultados levemente superiores en varias métricas de la "
        "corrida evaluada, mientras que PlayerNet no presentó una ventaja global. Este hallazgo resulta relevante "
        "porque muestra que una arquitectura más compleja no garantiza por sí sola una mejor solución y que la "
        "elección del modelo debe basarse en evidencia experimental, en el costo de los errores y en las "
        "condiciones reales de uso."
    ),
    (
        "El análisis de los objetivos también exige diferenciar los logros funcionales de los impactos todavía "
        "no comprobados. El objetivo general se alcanzó a nivel de MVP, ya que se desarrolló una aplicación que "
        "integra analítica de datos e inteligencia artificial como apoyo a la evaluación. En cambio, la "
        "identificación de talentos subrepresentados, la promoción de una mayor igualdad de oportunidades y la "
        "optimización efectiva de recursos no pudieron validarse sin datos longitudinales, usuarios reales ni una "
        "comparación con el proceso tradicional de scouting. Esta distinción permite presentar el alcance del "
        "trabajo de manera más precisa y evita atribuir al prototipo efectos que aún no fueron medidos."
    ),
    (
        "En síntesis, TPScouting constituye una base funcional y verificable para continuar investigando el uso de "
        "datos en el seguimiento de futbolistas juveniles. Su valor actual reside en integrar información, "
        "facilitar comparaciones y hacer explícitos los criterios utilizados, siempre como complemento del juicio "
        "de entrenadores y scouts. La siguiente etapa debería priorizar la recolección de datos reales, la "
        "validación longitudinal en clubes y la mejora del diseño experimental antes de aumentar la complejidad "
        "del modelo o plantear un uso operativo de mayor escala."
    ),
]


def find_unique_paragraph(document: Document, exact_text: str):
    matches = [
        (index, paragraph)
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() == exact_text
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Se esperaba un párrafo único para {exact_text!r}; encontrados: {len(matches)}")
    return matches[0]


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is None:
        raise RuntimeError("No se pudo obtener el contenedor del párrafo.")
    parent.remove(element)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if OUTPUT.exists():
        raise FileExistsError(f"El archivo de salida ya existe: {OUTPUT}")

    document = Document(SOURCE)

    summary_index, summary_heading = find_unique_paragraph(document, "RESUMEN")
    summary_paragraph = document.paragraphs[summary_index + 1]
    if not summary_paragraph.text.strip().startswith("En este trabajo final se presenta TPScouting"):
        raise RuntimeError("El resumen encontrado no coincide con la versión esperada.")
    summary_paragraph.text = SUMMARY

    conclusion_index, conclusion_heading = find_unique_paragraph(document, "7.1 Conclusiones")
    objectives_index, objectives_heading = find_unique_paragraph(document, "7.2 Cumplimiento de objetivos")
    paragraphs = document.paragraphs
    start_index = conclusion_index + 1
    end_index = objectives_index
    existing = paragraphs[start_index:end_index]
    if len(existing) != 2:
        raise RuntimeError(f"Se esperaban 2 párrafos de conclusiones; encontrados: {len(existing)}")
    for paragraph in existing:
        remove_paragraph(paragraph)

    for text in CONCLUSIONS:
        paragraph = objectives_heading.insert_paragraph_before(text, style="Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    document.save(OUTPUT)

    validation = Document(OUTPUT)
    texts = [paragraph.text.strip() for paragraph in validation.paragraphs]
    if texts.count(SUMMARY) != 1:
        raise RuntimeError("El resumen nuevo no quedó guardado exactamente una vez.")
    if any(texts.count(text) != 1 for text in CONCLUSIONS):
        raise RuntimeError("Las conclusiones nuevas no quedaron guardadas correctamente.")
    if len(validation.tables) != 22 or len(validation.sections) != 10 or len(validation.inline_shapes) != 27:
        raise RuntimeError("Cambió inesperadamente la estructura del documento.")

    print(f"OK: {OUTPUT}")
    print(
        "Estructura conservada: "
        f"{len(validation.tables)} tablas, {len(validation.sections)} secciones, "
        f"{len(validation.inline_shapes)} imágenes inline."
    )


if __name__ == "__main__":
    main()
