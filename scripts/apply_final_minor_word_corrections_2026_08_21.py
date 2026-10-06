from __future__ import annotations

import sys
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph


sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_COMPLETA_AUDITADA_21-08-2026.docx"
)
OUTPUT = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_REVISADA_21-08-2026.docx"
)
EVIDENCE_DIR = ROOT / "docs" / "evidencia_word_render"

CI_RUN = 73
CI_SHA = "bc5ddd35d0fa3bf6d85faab637772d4e9025fc98"


def norm(text: str) -> str:
    return " ".join((text or "").split())


def find_exact(doc: Document, text: str) -> Paragraph:
    needle = norm(text)
    matches = [paragraph for paragraph in doc.paragraphs if norm(paragraph.text) == needle]
    if len(matches) != 1:
        raise ValueError(f"Se esperaba 1 coincidencia para {text!r}; se encontraron {len(matches)}")
    return matches[0]


def set_paragraph(paragraph: Paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def replace_exact(doc: Document, old: str, new: str) -> None:
    set_paragraph(find_exact(doc, old), new)


def append_tc_field(paragraph: Paragraph, text: str, identifier: str) -> None:
    run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run.append(begin)
    paragraph._p.append(run)

    run = OxmlElement("w:r")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f' TC "{text}" \\f {identifier} '
    run.append(instruction)
    paragraph._p.append(run)

    run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(end)
    paragraph._p.append(run)


def replace_caption(doc: Document, old: str, new: str) -> Paragraph:
    paragraph = find_exact(doc, old)
    set_paragraph(paragraph, new)
    append_tc_field(paragraph, new, "F")
    return paragraph


def remove_section(start: Paragraph, end: Paragraph) -> None:
    node = start._p
    while node is not end._p:
        following = node.getnext()
        if following is None:
            raise RuntimeError("No se encontró el final de la sección a retirar")
        node.getparent().remove(node)
        node = following


def paragraph_has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def remove_preceding_image(doc: Document, caption: Paragraph) -> None:
    node = caption._p.getprevious()
    while node is not None and node.tag.endswith("}p"):
        paragraph = Paragraph(node, doc._body)
        if paragraph_has_drawing(paragraph):
            node.getparent().remove(node)
            return
        if norm(paragraph.text):
            break
        node = node.getprevious()
    raise RuntimeError(f"No se encontró la imagen anterior a: {caption.text}")


def add_image_before(doc: Document, caption: Paragraph, image_path: Path) -> None:
    node = OxmlElement("w:p")
    caption._p.addprevious(node)
    paragraph = Paragraph(node, doc._body)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(5.75))


def replace_ci_images(doc: Document, summary_caption: Paragraph) -> None:
    full_caption = find_exact(
        doc,
        "Figura 10-1. Historial público y ejecución exitosa del workflow CI de TPScouting.",
    )
    replacements = (
        (summary_caption, EVIDENCE_DIR / "08_ci_run_73_resumen.png"),
        (full_caption, EVIDENCE_DIR / "09_ci_run_73_ampliada.png"),
    )
    for caption, image_path in replacements:
        if not image_path.exists():
            raise FileNotFoundError(image_path)
        remove_preceding_image(doc, caption)
        add_image_before(doc, caption, image_path)


def correct_prose(doc: Document) -> None:
    replacements = {
        "Con la expansión de la IA y su acercamiento al ámbito cotidiano, esta dinámica está experimentando un cambio relevante en las formas de trabajo del scouting. Su capacidad para procesar y analizar grandes volúmenes de datos ofrece una nueva perspectiva para la identificación y el desarrollo de jóvenes talentos. Herramientas basadas en IA, como sistemas de análisis de rendimiento y algoritmos predictivos, brindan conclusiones más profundas y precisas sobre el potencial y las habilidades de los jugadores. Por ejemplo, plataformas como Wyscout y Opta Pro, utilizadas profesionalmente, permiten análisis detallados del rendimiento en campo y mejoran la toma de decisiones. También favorecen entrenamientos personalizados, adaptados a las necesidades y fortalezas individuales y colectivas. No obstante, la adopción de estas tecnologías enfrenta desafíos en clubes con recursos limitados, tanto por restricciones presupuestarias como por falta de formación técnica específica.":
            "Con la expansión de la IA y su incorporación al análisis deportivo, los procesos de scouting disponen de nuevas herramientas para organizar y comparar información. Plataformas como Wyscout centralizan video y datos deportivos para apoyar la observación, la comparación y el reclutamiento (Hudl, s. f.). Sin embargo, disponer de más datos no garantiza por sí mismo decisiones más precisas: los resultados dependen de la calidad del registro, del objetivo definido y de la revisión humana.",
        "Actualmente, los métodos de scouting son subjetivos y basados en observación, lo que puede llevar a errores en la identificación del potencial de un jugador.":
            "La observación experta aporta contexto cualitativo indispensable, pero cuando los registros no están estandarizados resulta difícil comparar evaluaciones, conservar trazabilidad y revisar decisiones.",
        "Aplicación en el Deporte: ML puede utilizarse para una variedad de aplicaciones en el deporte, incluyendo la optimización de entrenamientos, análisis táctico, y prevención de lesiones. Permite a entrenadores y analistas deportivos obtener descubrimientos profundos y basados en datos que antes eran imposibles o muy difíciles de obtener. Se aplica para análisis de rendimiento, prevención de lesiones, personalización del entrenamiento, y scouting de jugadores. Incluye ejemplos específicos, como modelos predictivos para identificar futuros talentos o sistemas de análisis de movimiento para optimizar técnicas deportivas.":
            "Aplicación en el deporte: el aprendizaje automático puede apoyar tareas de evaluación de rendimiento, valoración de acciones y detección de perfiles, siempre que el objetivo y los datos estén definidos de manera verificable (Pappalardo et al., 2019; Decroos et al., 2019; Lacan, 2024). En este trabajo su uso se limita a una clasificación experimental sobre datos sintéticos y no demuestra prevención de lesiones ni personalización real del entrenamiento.",
        "Identificación temprana de talentos: La IA está transformando la forma en que se identifican los talentos en el deporte. Al analizar datos de rendimiento desde etapas tempranas, los algoritmos de IA pueden identificar jugadores con alto potencial que quizás no sean evidentes a través de métodos tradicionales.":
            "Identificación temprana de perfiles: los métodos basados en datos pueden apoyar la comparación de jugadores jóvenes y la búsqueda de señales de proyección, pero su desempeño depende del conjunto de datos y no reemplaza la evaluación contextual (Lacan, 2024). En TPScouting esta posibilidad se presenta como contribución potencial, ya que el modelo fue validado únicamente sobre datos sintéticos.",
        "La ejecución local directa del 10/07/2026 obtuvo 83 pruebas aprobadas, 1 omitida y 4 advertencias en 38,88 segundos. La ejecución con pytest-cov obtuvo 79 % de cobertura sobre 5.082 sentencias, con 1.062 no cubiertas. La prueba omitida corresponde al smoke visual de Playwright, configurado como opt-in; las advertencias son RuntimeWarning de scikit-learn por columnas completamente NaN en dos pruebas de preprocesamiento.":
            "La ejecución local directa del 21/08/2026 obtuvo 84 pruebas aprobadas, 1 omitida y 4 advertencias. La ejecución con pytest-cov obtuvo 80 % de cobertura sobre 5.083 sentencias, con 1.032 no cubiertas. La prueba omitida corresponde al smoke visual de Playwright, configurado como opt-in; las advertencias son RuntimeWarning de scikit-learn por columnas completamente NaN en dos pruebas de preprocesamiento.",
        "La ejecución pública CI #71 finalizó con estado success para el commit fa8a50d1173f2760329a45f11fbf93a709721235. Los jobs de Python 3.11 y 3.12 completaron instalación, tests y carga de artefactos de cobertura. Esta evidencia no certifica el HEAD local 7d7680e, que durante la auditoría se encontraba un commit por delante.":
            f"La ejecución pública CI #{CI_RUN} finalizó con estado success para el commit {CI_SHA}. Los jobs de Python 3.11 y 3.12 completaron instalación, tests y carga de artefactos de cobertura. Esta evidencia certifica ese commit de main; no certifica cambios locales posteriores no publicados.",
        "La Figura 10-1 conserva una vista ampliada de la ejecución CI #71, con jobs exitosos para Python 3.11 y 3.12 y dos artefactos de cobertura.":
            f"La Figura 10-1 conserva una vista ampliada de la ejecución CI #{CI_RUN}, con jobs exitosos para Python 3.11 y 3.12 y dos artefactos de cobertura.",
    }
    for old, new in replacements.items():
        replace_exact(doc, old, new)


def remove_technology_subsection(doc: Document) -> None:
    start = find_exact(doc, "2.1.6 Tecnologías consideradas y utilizadas")
    end = find_exact(doc, "2.1.7 Fundamentos de modelado y evaluación")
    remove_section(start, end)
    set_paragraph(end, "2.1.6 Fundamentos de modelado y evaluación")
    replace_exact(
        doc,
        "2.1.7.1 Redes neuronales feed-forward y clasificación binaria",
        "2.1.6.1 Redes neuronales feed-forward y clasificación binaria",
    )
    replace_exact(
        doc,
        "2.1.7.2 Normalización y vector de características",
        "2.1.6.2 Normalización y vector de características",
    )
    replace_exact(
        doc,
        "2.1.7.3 Métricas de evaluación de clasificación",
        "2.1.6.3 Métricas de evaluación de clasificación",
    )


def find_table(doc: Document, first_headers: tuple[str, ...]):
    matches = []
    for table in doc.tables:
        headers = tuple(norm(cell.text) for cell in table.rows[0].cells)
        if headers[: len(first_headers)] == first_headers:
            matches.append(table)
    if len(matches) != 1:
        raise ValueError(f"Se esperaba 1 tabla con cabecera {first_headers}; se encontraron {len(matches)}")
    return matches[0]


def correct_tables(doc: Document) -> None:
    dimensions = find_table(doc, ("Dimensión", "Variables representativas", "Fuente"))
    performance = [row for row in dimensions.rows if norm(row.cells[0].text) == "Rendimiento"]
    if len(performance) != 1:
        raise ValueError("No se encontró una única fila Rendimiento en la tabla de dimensiones")
    performance[0].cells[1].text = (
        "minutes_played, goals, assists, pass_accuracy, shot_accuracy, "
        "duels_won_pct, final_score"
    )

    evidence = find_table(doc, ("Evidencia", "Resultado verificado"))
    values = {
        "pytest local": "84 passed, 1 skipped, 4 warnings",
        "pytest-cov": "80 %; 5.083 sentencias; 1.032 no cubiertas; 55,06 s",
        "CI pública": f"Run #{CI_RUN} completada con éxito en Python 3.11 y 3.12",
        "Commit certificado por CI": CI_SHA,
        "Artefactos CI": "coverage-python-3.11 y coverage-python-3.12",
        "Alcance": f"La run pública certifica el commit {CI_SHA[:7]} de main; no certifica cambios locales posteriores no publicados.",
        "Smoke Render": "Validación histórica del 20/05/2026; los controles posteriores finalizaron por timeout. No se afirma disponibilidad actual.",
    }
    seen = set()
    for row in evidence.rows[1:]:
        key = norm(row.cells[0].text)
        if key in values:
            row.cells[1].text = values[key]
            seen.add(key)
    if seen != set(values):
        raise ValueError(f"Filas de evidencia faltantes: {sorted(set(values) - seen)}")

    glossary = find_table(doc, ("Término", "Definición"))
    scouting_rows = [row for row in glossary.rows[1:] if norm(row.cells[0].text).casefold() == "scouting"]
    definition = (
        "En este trabajo, proceso sistemático de observación, registro y comparación "
        "de jugadores para apoyar decisiones de identificación y seguimiento deportivo."
    )
    if scouting_rows:
        if len(scouting_rows) != 1:
            raise ValueError("El glosario contiene más de una fila Scouting")
        scouting_rows[0].cells[1].text = definition
    else:
        cells = glossary.add_row().cells
        cells[0].text = "Scouting"
        cells[1].text = definition


def main() -> int:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    correct_prose(doc)
    remove_technology_subsection(doc)
    correct_tables(doc)
    summary_caption = replace_caption(
        doc,
        "Figura 6-8. Resumen público de la ejecución CI #71 en GitHub Actions.",
        f"Figura 6-8. Resumen público de la ejecución CI #{CI_RUN} en GitHub Actions.",
    )
    replace_ci_images(doc, summary_caption)

    doc.core_properties.title = "TPScouting - Trabajo Final revisado"
    doc.core_properties.subject = "Revisión técnica final del 21/08/2026"
    doc.save(OUTPUT)
    print(f"Documento revisado guardado en: {OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
