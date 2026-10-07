"""Verify the final Block 8 DOCX/PDF without modifying either artifact."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import fitz
from docx import Document


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--pdf", required=True, type=Path)
    parser.add_argument("--json", required=True, type=Path)
    args = parser.parse_args()

    word = Document(args.docx)
    paragraphs = [p.text.strip() for p in word.paragraphs if p.text.strip()]
    list_entries = [
        text
        for text in paragraphs
        if re.match(r"^(Figura|Tabla) \d+-\d+\..+\t\d+$", text)
    ]

    pdf = fitz.open(args.pdf)
    pdf_texts = [page.get_text("text") for page in pdf]
    normalized_pdf_texts = [re.sub(r"\s+", " ", text).strip().casefold() for text in pdf_texts]
    mismatches = []
    for entry in list_entries:
        caption, listed_page = entry.rsplit("\t", 1)
        needle = re.sub(r"\s+", " ", caption).strip().casefold()
        physical_pages = [
            index
            for index, page_text in enumerate(normalized_pdf_texts)
            if index >= 8 and needle in page_text
        ]
        actual_pages = []
        for index in physical_pages:
            visible_numbers = re.findall(r"(?m)^\s*(\d+)\s*$", pdf_texts[index])
            actual_pages.append(int(visible_numbers[0]) if visible_numbers else None)
        if int(listed_page) not in actual_pages:
            mismatches.append(
                {
                    "caption": caption,
                    "listed_page": int(listed_page),
                    "actual_pages": actual_pages,
                }
            )

    pending = sum(text.count("PENDIENTE DE INFORMAR") for text in paragraphs)
    field_errors = []
    for page_number, page_text in enumerate(pdf_texts, 1):
        for marker in ("Error!", "¡Error!", "Reference source not found", "No se encuentra"):
            if marker.casefold() in page_text.casefold():
                field_errors.append({"page": page_number, "marker": marker})

    page_sizes = sorted(
        {
            (round(page.rect.width, 2), round(page.rect.height, 2))
            for page in pdf
        }
    )
    low_text_pages = []
    for page_number, (page, text) in enumerate(zip(pdf, pdf_texts), 1):
        visible = re.sub(r"\s+", "", text)
        if len(visible) < 40:
            low_text_pages.append(
                {
                    "page": page_number,
                    "characters": len(visible),
                    "images": len(page.get_images(full=True)),
                }
            )

    result = {
        "docx": str(args.docx),
        "docx_sha256": sha256(args.docx),
        "pdf": str(args.pdf),
        "pdf_sha256": sha256(args.pdf),
        "pdf_pages": len(pdf),
        "docx_inline_shapes": len(word.inline_shapes),
        "docx_tables": len(word.tables),
        "docx_sections": len(word.sections),
        "list_entries": len(list_entries),
        "list_mismatches": mismatches,
        "pending_editorial_items": pending,
        "field_errors": field_errors,
        "page_sizes_points": page_sizes,
        "low_text_pages": low_text_pages,
        "passed": (
            len(list_entries) == 45
            and not mismatches
            and pending == 1
            and not field_errors
            and len(page_sizes) == 1
            and all(item["images"] > 0 for item in low_text_pages)
            and len(word.inline_shapes) == 26
            and len(word.tables) == 19
            and len(word.sections) == 9
        ),
    }
    args.json.parent.mkdir(parents=True, exist_ok=True)
    args.json.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(result, indent=2, ensure_ascii=False))
    if not result["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
