"""Create a new thesis copy with the post-audit traceability corrections."""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from docx import Document


EXPECTED_SOURCE_SHA256 = "cbed5344342f4e96e0cd996520093af99a364b851ee24b792a45e171294f474f"

PARAGRAPH_REPLACEMENTS = {
    "La verificación local del 05/10/2026 obtuvo 116 pruebas aprobadas, 1 omitida y 4 advertencias. La cobertura fue 83,74 %. La prueba omitida es el smoke visual Playwright opt-in; las cuatro advertencias son RuntimeWarning de scikit-learn por columnas completamente NaN en dos pruebas de preprocesamiento. Estos resultados corresponden al árbol local y no cuentan todavía con una CI pública del mismo commit.":
        "La verificación local del 10/10/2026 obtuvo 116 pruebas aprobadas, 1 omitida y 4 advertencias, con una cobertura de 83,74 %. El smoke visual Playwright opt-in se ejecutó por separado y aprobó; las cuatro advertencias son RuntimeWarning de scikit-learn por columnas completamente NaN en dos pruebas de preprocesamiento. El repositorio de entrega y el ajuste posterior del repositorio principal cuentan con ejecuciones CI públicas asociadas a sus commits exactos.",
    "La CI #73 es evidencia histórica del commit bc5ddd35d0fa3bf6d85faab637772d4e9025fc98 en el repositorio de desarrollo. El repositorio de entrega revisado se encuentra en ffdefdf8035c994ae285a270de0a4ff4e9f336a8, mientras que las correcciones de octubre se guardaron en commits locales posteriores y todavía no se sincronizaron. Por tanto, esa ejecución histórica no certifica la entrega revisada ni el árbol corregido; la evidencia CI del commit finalmente entregado queda pendiente.":
        "La CI #73 se conserva como evidencia histórica del commit bc5ddd35d0fa3bf6d85faab637772d4e9025fc98 en el repositorio de desarrollo. La entrega corregida se publicó en el commit 7a47bd3f164e865677f2c70075cbedc9fa63427d y fue certificada por la ejecución CI 37707576905, aprobada en Linux con Python 3.11 y 3.12 junto con la auditoría de dependencias. El ajuste posterior de PyTorch 2.14.1 quedó guardado en el repositorio principal mediante el commit 6f83e4066db7ccb9900554568a0389cdf1d12d20 y su CI 37709526994 también finalizó correctamente; ese ajuste aún no fue sincronizado al repositorio de entrega.",
    "Repositorio de entrega revisado: https://github.com/indio21/TPScouting-entrega, commit ffdefdf8035c994ae285a270de0a4ff4e9f336a8. El trabajo de corrección se realizó en el repositorio principal a partir del checkpoint 6e29b45a85396b5cbe82e8b0ece2b0eb394a7bfd, se guardó mediante commits locales posteriores y todavía no fue sincronizado ni publicado en la entrega.":
        "Repositorio de entrega: https://github.com/indio21/TPScouting-entrega, commit publicado 7a47bd3f164e865677f2c70075cbedc9fa63427d, validado por la ejecución CI 37707576905. El ajuste posterior de dependencias se conserva en el repositorio principal de respaldo, commit 6f83e4066db7ccb9900554568a0389cdf1d12d20, con CI 37709526994 aprobada y pendiente de una futura sincronización selectiva a la entrega.",
    "Los comandos siguientes se ejecutan desde la raíz C:\\Tesis\\TPScouting. requirements.txt contiene runtime y requirements-dev.txt agrega pruebas, auditoría y herramientas documentales. La instalación de PyTorch depende de plataforma y acelerador; debe usarse el selector oficial de PyTorch. Se verificó Windows con Python 3.11; no se afirma una ejecución local en Linux o macOS.":
        "Los comandos siguientes se ejecutan desde la raíz del repositorio clonado. En Windows y Linux CPU, el runtime se instala desde requirements-lock.txt y PyTorch desde requirements-torch-cpu.txt; requirements-dev.txt agrega pruebas, auditoría y lint. En macOS se usa requirements.txt para obtener el wheel estándar de PyTorch y luego requirements-dev.txt. requirements-docs.txt se reserva para trabajar con el documento. Se verificó Windows con Python 3.11 y CI en Linux con Python 3.11 y 3.12; no se afirma una ejecución local en macOS.",
    "Windows: .\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt -r requirements-dev.txt\nLinux/macOS: .venv/bin/python -m pip install -r requirements.txt -r requirements-dev.txt":
        "Windows:\n.\\.venv\\Scripts\\python.exe -m pip install -r requirements-lock.txt\n.\\.venv\\Scripts\\python.exe -m pip install -r requirements-torch-cpu.txt\n.\\.venv\\Scripts\\python.exe -m pip install -r requirements-dev.txt\n\nLinux CPU:\n.venv/bin/python -m pip install -r requirements-lock.txt\n.venv/bin/python -m pip install -r requirements-torch-cpu.txt\n.venv/bin/python -m pip install -r requirements-dev.txt\n\nmacOS:\n.venv/bin/python -m pip install -r requirements.txt\n.venv/bin/python -m pip install -r requirements-dev.txt",
}

TABLE_12_REPLACEMENTS = {
    (1, 0): "pytest local (10/10/2026)",
    (3, 1): "Smoke visual Playwright opt-in ejecutado por separado: 1 passed",
    (5, 0): "HEAD del repositorio principal de respaldo",
    (5, 1): "6f83e4066db7ccb9900554568a0389cdf1d12d20; PyTorch 2.14.1 y auditoría directa del manifest",
    (6, 0): "Commit publicado en el repositorio de entrega",
    (6, 1): "7a47bd3f164e865677f2c70075cbedc9fa63427d",
    (7, 1): "Entrega: CI 37707576905 aprobada. Respaldo: CI 37709526994 aprobada.",
    (8, 1): "La CI de entrega certifica 7a47bd3; la CI del respaldo certifica el ajuste posterior 6f83e40, aún no sincronizado a la entrega.",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_text(container, new_text: str) -> None:
    if not container.runs:
        container.add_run(new_text)
        return
    container.runs[0].text = new_text
    for run in container.runs[1:]:
        run.text = ""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    if sha256(args.source) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("El DOCX fuente no coincide con la copia autoritativa esperada.")

    document = Document(args.source)
    counts = {old: 0 for old in PARAGRAPH_REPLACEMENTS}
    for paragraph in document.paragraphs:
        if paragraph.text in PARAGRAPH_REPLACEMENTS:
            old = paragraph.text
            replace_text(paragraph, PARAGRAPH_REPLACEMENTS[old])
            counts[old] += 1

    missing = [old for old, count in counts.items() if count != 1]
    if missing:
        raise RuntimeError(f"Reemplazos no unívocos: {missing}")

    table = document.tables[11]
    for (row_index, cell_index), new_text in TABLE_12_REPLACEMENTS.items():
        cell = table.cell(row_index, cell_index)
        paragraph = cell.paragraphs[0]
        replace_text(paragraph, new_text)
        for extra in cell.paragraphs[1:]:
            replace_text(extra, "")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)

    checked = Document(args.output)
    if len(checked.inline_shapes) != 28 or len(checked.tables) != 19 or len(checked.sections) != 9:
        raise RuntimeError("La corrección alteró imágenes, tablas o secciones.")

    full_text = "\n".join(p.text for p in checked.paragraphs)
    for forbidden in (
        "commit ffdefdf8035c994ae285a270de0a4ff4e9f336a8",
        "evidencia CI del commit finalmente entregado queda pendiente",
        "pip install -r requirements.txt -r requirements-dev.txt",
    ):
        if forbidden in full_text:
            raise RuntimeError(f"Texto obsoleto conservado: {forbidden}")

    print(f"OUTPUT={args.output}")
    print(f"PARAGRAPH_REPLACEMENTS={len(PARAGRAPH_REPLACEMENTS)}")
    print(f"TABLE_CELL_REPLACEMENTS={len(TABLE_12_REPLACEMENTS)}")
    print(f"SHA256={sha256(args.output).upper()}")


if __name__ == "__main__":
    main()
