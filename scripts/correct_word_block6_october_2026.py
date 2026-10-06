"""Create the Block 6 corrected thesis draft without overwriting the approved Word.

The transformation is deliberately tied to the SHA-256 of the approved 26-Aug-2026
document. It only writes an internal backup and a new draft under docs/.
"""

from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_"
    r"ENTREGA_FINAL_REVISADA_26-08-2026_v2.docx"
)
OUT_DIR = ROOT / "docs" / "revision_octubre_2026" / "word"
BACKUP = OUT_DIR / "ORIGINAL_APROBADO_26-08-2026_v2_BACKUP_SHA77261990.docx"
OUTPUT = OUT_DIR / "TRABAJO_FINAL_TPScouting_CORREGIDO_OCTUBRE_2026_BORRADOR.docx"
EXPECTED_SHA256 = "77261990098b6fd0761c7e1d27fcaf8cb91bef8fe13346c39836075e6049955d"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) > 1:
        body_matches = [p for p in matches if not p.style.name.lower().startswith("toc")]
        if len(body_matches) == 1:
            return body_matches[0]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting {prefix!r}; found {len(matches)}")
    return matches[0]


def replace(document: Document, prefix: str, text: str) -> None:
    paragraph = find_paragraph(document, prefix)
    paragraph.text = text


def insert_before(anchor: Paragraph, text: str, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def insert_after(anchor: Paragraph, text: str, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def rewrite_table(table, rows: list[list[str]]) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for col, value in enumerate(rows[0]):
        table.rows[0].cells[col].text = value
    for values in rows[1:]:
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cells[col].text = value


def main() -> None:
    if digest(SOURCE) != EXPECTED_SHA256:
        raise RuntimeError("The approved source differs from the audited SHA-256")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, BACKUP)
    if digest(BACKUP) != EXPECTED_SHA256:
        raise RuntimeError("Backup verification failed")

    document = Document(SOURCE)

    # Cover. The legajo remains explicitly pending because it was not supplied.
    replace(document, "Trabajo Final", "Trabajo Final de Grado")
    replace(document, "Ingeniería en Informática", "Universidad Católica de Santiago del Estero\nDepartamento Académico Rafaela\nIngeniería en Informática")
    replace(document, "Alumno:", "Alumno: Solari, Pablo\nLegajo: [PENDIENTE DE INFORMAR]")
    replace(document, "Prof:", "Directores: Duarte, Jorge – Lovera, Maximiliano")
    replace(document, "Año:", "Rafaela, Santa Fe, Argentina\nOctubre de 2026")

    # Abstract: current authorization of Block 6 supersedes the earlier omission.
    intro = find_paragraph(document, "1. INTRODUCCIÓN")
    insert_before(intro, "ABSTRACT", "Heading 1")
    insert_before(
        intro,
        "Youth player monitoring in resource-constrained football clubs often relies on "
        "scattered records and criteria that are difficult to compare over time. This final "
        "project designed and developed TPScouting, a minimum viable product for organizing "
        "and analyzing sporting information about players aged 12 to 18. The application "
        "manages profiles, technical attributes, statistics, assessments, reports, availability, "
        "and historical records. It includes a dashboard, comparison tools, and a potential "
        "estimate produced by a PyTorch model. The solution was evaluated with synthetic data, "
        "automated tests, classification metrics, baseline comparisons, and dated evidence of "
        "application, deployment, and continuous-integration behavior. The results establish "
        "the technical feasibility of the complete flow from data entry and persistence to the "
        "presentation of an estimate in the interface. They do not establish sporting validity "
        "or social impact because no real players or clubs participated. TPScouting is therefore "
        "presented as a decision-support prototype that complements, rather than replaces, the "
        "judgment of coaches and scouts.",
    )
    insert_before(
        intro,
        "Keywords: youth scouting; youth football; artificial intelligence; machine learning; "
        "data analysis; PyTorch.",
    )

    replace(
        document,
        "Frente a esta realidad",
        "Frente a esta situación, el proyecto evalúa la factibilidad técnica de integrar registros "
        "estructurados, visualizaciones y un modelo predictivo en un MVP de apoyo al scouting. "
        "El trabajo no mide mejoras reales de objetividad, visibilidad ni toma de decisiones, ya "
        "que utiliza datos sintéticos y no fue validado con clubes, scouts o jugadores.",
    )
    replace(
        document,
        "Por lo tanto, esta investigación",
        "La pregunta de investigación se responde dentro del alcance técnico: la ciencia de datos "
        "permite centralizar, historizar y comparar registros, y la inteligencia artificial permite "
        "producir una estimación reproducible a partir de esos registros. El MVP demuestra que ambas "
        "capacidades pueden integrarse en una aplicación accesible. No demuestra que optimicen "
        "recursos, detecten mejor el talento ni mejoren decisiones deportivas reales; esas afirmaciones "
        "requieren validación longitudinal con instituciones y participantes reales.",
    )
    replace(
        document,
        "Objetivo general:",
        "Objetivo general: diseñar, implementar y verificar un MVP que centralice información de "
        "futbolistas juveniles de 12 a 18 años, permita consultar su evolución y comparar perfiles, e "
        "integre una estimación de potencial basada en datos sintéticos, con trazabilidad de datos, "
        "artefactos y pruebas.",
    )
    objective_replacements = {
        "Optimizar la evaluación:": "Implementar registros estructurados e historiales de atributos, estadísticas y evaluaciones con reglas verificables.",
        "Identificar talentos subrepresentados:": "Implementar filtros, comparadores y una estimación de potencial que permitan priorizar perfiles para revisión humana, sin afirmar eficacia deportiva real.",
        "Mejorar la toma de decisiones:": "Presentar en fichas y paneles información trazable que pueda ser utilizada como apoyo por entrenadores, scouts y directivos.",
        "Promover la igualdad de oportunidades:": "Evaluar la factibilidad de una herramienta de bajo costo para clubes con recursos limitados; su efecto sobre la igualdad de oportunidades queda fuera de la validación realizada.",
        "Optimizar la gestión de recursos:": "Verificar que el MVP organiza y recupera información; la reducción de tiempos, costos o recursos requiere un estudio posterior con usuarios reales.",
    }
    for prefix, text in objective_replacements.items():
        replace(document, prefix, text)

    # Theory and method additions grounded in checked literature and the current pipeline.
    related = find_paragraph(document, "2.2 Trabajos relacionados")
    insert_before(related, "2.1.6.4 Calibración, desbalance y fuga de información", "Heading 4")
    insert_before(
        related,
        "La discriminación y la calibración responden a preguntas diferentes: ROC-AUC ordena casos, "
        "mientras que una probabilidad calibrada busca que frecuencias observadas y probabilidades "
        "estimadas sean coherentes. La calibración isotónica debe ajustarse fuera del conjunto de test "
        "para conservar una evaluación final independiente (Niculescu-Mizil & Caruana, 2005). Con una "
        "clase positiva minoritaria, PR-AUC complementa ROC-AUC porque expone la relación entre precisión "
        "y recall para la clase de interés (Saito & Rehmsmeier, 2015). Cualquier estadístico calculado con "
        "todo el conjunto antes del split puede transferir información de validación o test al proceso y "
        "debe declararse como riesgo de fuga.",
    )
    insert_before(related, "2.1.6.5 Datos sintéticos y sesgos de edad y maduración", "Heading 4")
    insert_before(
        related,
        "Los datos sintéticos permiten probar el flujo técnico sin exponer datos personales, pero su "
        "distribución y su etiqueta reflejan reglas de generación y no evidencia observada en futbolistas. "
        "Además, agrupar juveniles por edad cronológica puede favorecer sistemáticamente a quienes nacieron "
        "antes dentro del año de selección o maduraron antes. La literatura identifica este efecto de edad "
        "relativa como un sesgo relevante en el desarrollo deportivo (Cobley et al., 2009). Por ello, edad, "
        "maduración y contexto deben analizarse antes de cualquier uso real del score.",
    )
    replace(
        document,
        "La corrida oficial no utiliza",
        "La corrida oficial no utiliza potential_label como variable objetivo ni incorpora potential_label, "
        "temporal_target_label o predicciones persistidas entre las 64 columnas de entrada (68 tras la "
        "codificación). El target efectivo es temporal_target_label, una etiqueta binaria sintética construida "
        "con señales futuras simuladas. La clase positiva se selecciona mediante un score temporal, controles "
        "de calidad y una cuota por cohorte de posición y edad. Por ello, la prevalencia de 1.597/20.000 "
        "(7,985 %) está fijada por el diseño sintético y no constituye una observación sobre futbolistas reales.",
    )
    replace(
        document,
        "Criterios de éxito:",
        "Criterios de evaluación retrospectiva: la revisión del MVP considera funcionamiento de procesos, "
        "pruebas automatizadas, evidencia operativa y métricas del modelo. No se documentaron criterios de "
        "aceptación cuantitativos establecidos antes del experimento; por ello estos criterios no se presentan "
        "como hipótesis ni como umbrales a priori. El impacto sobre eficiencia, efectividad, igualdad de "
        "oportunidades o decisiones deportivas requiere validación posterior.",
    )

    # Replace generic comparisons with the decisions that were actually made.
    replace(
        document,
        "Comparación: Flask",
        "Decisión tecnológica: se eligió Flask porque el proyecto y el pipeline de datos ya utilizan Python, "
        "y porque su modelo server-side cubre el alcance del MVP con pocas capas. .NET Minimal APIs fue una "
        "alternativa considerada, pero no se realizó un benchmark; por ello no se atribuyen ventajas de "
        "rendimiento o escalabilidad a partir de este trabajo.",
    )
    for prefix in ["Lenguaje: Python", "Enfoque: Es extremadamente", "Flexibilidad: Flask", "Componentes opcionales: Se pueden", "Facilidad de uso: Fácil", "Escalabilidad: Aunque", "Despliegue: Desplegar Flask", "Popularidad: Muy popular", "Lenguaje: C#", "Enfoque: Al igual", "Flexibilidad: Minimal", "Componentes opcionales: .NET", "Facilidad de uso: Aunque C#", "Escalabilidad: Muy adecuado", "Despliegue: .NET", "Popularidad: Cada vez"]:
        try:
            replace(document, prefix, "")
        except RuntimeError:
            pass
    rewrite_table(
        document.tables[1],
        [
            ["Alternativa", "Relación con el MVP", "Decisión"],
            ["Flask (Python)", "Integra aplicación web y pipeline ML en el mismo lenguaje.", "Seleccionada e implementada."],
            [".NET Minimal APIs", "Alternativa viable, pero implicaba mantener C# junto con el pipeline Python.", "No seleccionada; no se realizó benchmark."],
        ],
    )
    rewrite_table(
        document.tables[2],
        [
            ["Alternativa", "Relación con el MVP", "Decisión"],
            ["SQLAlchemy", "Abstrae SQLite y PostgreSQL y ya sostiene las relaciones del dominio.", "Seleccionada e implementada."],
            ["Peewee", "ORM más pequeño, sin una ventaja comprobada para este dominio.", "No seleccionada; no se realizó benchmark."],
        ],
    )
    rewrite_table(
        document.tables[4],
        [
            ["Alternativa", "Relación con el MVP", "Decisión"],
            ["SQLite", "Base local y temporal para desarrollo y pruebas.", "Seleccionada para uso local."],
            ["PostgreSQL", "Persistencia administrada compatible con el despliegue en Render mediante psycopg 3.", "Seleccionada para la demo desplegada."],
            ["MySQL / SQL Server", "Alternativas no necesarias para el alcance.", "No evaluadas mediante benchmark."],
        ],
    )
    rewrite_table(
        document.tables[7],
        [
            ["Entidad", "Propósito"],
            ["Player", "Jugador, datos personales, posición, club, atributos base, birth_date, photo_url y potencial."],
            ["Match", "Partido y contexto competitivo."],
            ["Coach", "Entrenador registrado."],
            ["Director", "Directivo registrado."],
            ["User", "Cuenta, credenciales y rol de acceso."],
            ["PlayerStat", "Estadísticas de rendimiento por fecha."],
            ["PlayerAttributeHistory", "Historial de atributos técnicos, físicos y mentales."],
            ["PlayerMatchParticipation", "Participación del jugador en un partido."],
            ["ScoutReport", "Reporte cualitativo de scouting."],
            ["PhysicalAssessment", "Evaluación física y mediciones."],
            ["PlayerAvailability", "Disponibilidad, fatiga y estado físico."],
        ],
    )

    # Verified model facts and explicit score distinctions.
    document.tables[11].rows[7].cells[1].text = "17.608 (model.parameters()); 386 buffers; 17.994 elementos totales en state_dict"
    replace(
        document,
        "La aplicación también calcula",
        "La salida sigmoid cruda de PlayerNet es una probabilidad sin calibrar. La calibración isotónica "
        "produce otra probabilidad y utiliza un umbral de clasificación seleccionado en validation (0,25; "
        "0,825 para la salida cruda). La interfaz presenta además un score combinado con información histórica "
        "y posicional; sus bandas visuales son bajo <0,60, medio 0,60–0,79 y alto ≥0,80. Esos límites visuales "
        "no son los umbrales usados para medir F1 en test.",
    )
    replace(
        document,
        "El modelo operativo entonces es PlayerNet",
        "PlayerNet contiene 17.608 parámetros entrenables contados con model.parameters(). Su state_dict "
        "contiene 17.994 elementos porque suma 386 buffers de BatchNorm. La arquitectura combina una rama "
        "lineal 68→1 con una rama residual 68→128→64→1, BatchNorm, GELU y dropout 0,15. El entrenamiento "
        "utiliza BCEWithLogitsLoss, AdamW y learning rate inicial 0,0005.",
    )

    # Security wording reflects Blocks 1-2 without claiming production validation.
    security = find_paragraph(document, "5.5 Seguridad")
    next_heading = find_paragraph(document, "5.6 Despliegue")
    for text in [
        "• APP_SECRET_KEY obligatoria en producción y limpieza de sesión al autenticar.",
        "• Protección CSRF con comparación de tiempo constante y formularios POST para operaciones mutantes.",
        "• Redirección next limitada a destinos internos y rechazo de esquemas, barras invertidas y rutas //.",
        "• Revalidación del usuario, estado y rol contra la base en cada solicitud protegida.",
        "• Rate limiting local basado en una clave de cliente definida por una política de proxies confiables; al ser memoria de proceso, no cubre varias instancias.",
        "• /health público sin excepciones internas, CSP compatible con los recursos inventariados, HSTS bajo HTTPS y validación central de photo_url.",
        "• Entrenamiento y generación de datos desactivados en requests de producción; los artefactos joblib y PyTorch deben provenir de fuentes confiables.",
    ]:
        insert_before(next_heading, text)
    # Remove the prior bullets between both headings.
    node = security._p.getnext()
    keep = {p._p for p in document.paragraphs if p.text.startswith("• APP_SECRET_KEY obligatoria en producción y limpieza") or p.text.startswith("• Protección CSRF con comparación") or p.text.startswith("• Redirección next") or p.text.startswith("• Revalidación del usuario") or p.text.startswith("• Rate limiting local") or p.text.startswith("• /health público") or p.text.startswith("• Entrenamiento y generación")}
    while node is not None and node is not next_heading._p:
        nxt = node.getnext()
        if node not in keep:
            node.getparent().remove(node)
        node = nxt

    replace(
        document,
        "La última ejecución local directa",
        "La verificación local del 05/10/2026 obtuvo 116 pruebas aprobadas, 1 omitida y 4 advertencias. "
        "La cobertura fue 83,74 %. La prueba omitida es el smoke visual Playwright opt-in; las cuatro "
        "advertencias son RuntimeWarning de scikit-learn por columnas completamente NaN en dos pruebas de "
        "preprocesamiento. Estos resultados corresponden al árbol local y no cuentan todavía con una CI "
        "pública del mismo commit.",
    )
    rewrite_table(
        document.tables[14],
        [
            ["Evidencia", "Resultado verificado"],
            ["pytest local (05/10/2026)", "116 passed, 1 skipped, 4 warnings"],
            ["pytest-cov local", "83,74 %"],
            ["Prueba omitida", "Smoke visual Playwright opt-in; requiere RUN_PLAYWRIGHT=1"],
            ["Warnings", "4 All-NaN slice encountered en dos pruebas de preprocesamiento"],
            ["HEAD de desarrollo", "49aa51c0167fdb24c5f2f3a6ab6e3f397830b462 más cambios locales no publicados"],
            ["Repo de entrega revisado", "ffdefdf8035c994ae285a270de0a4ff4e9f336a8"],
            ["CI de las correcciones", "Pendiente: no existe evidencia CI pública del árbol local corregido."],
            ["Alcance", "La CI histórica no certifica los cambios locales ni el commit actual del repositorio de entrega."],
            ["Smoke Render", "Verificación histórica del 20/05/2026; no acredita disponibilidad actual."],
        ],
    )
    replace(
        document,
        "La ejecución pública CI #73",
        "La CI #73 es evidencia histórica del commit bc5ddd35d0fa3bf6d85faab637772d4e9025fc98 "
        "en el repositorio de desarrollo. El repositorio de entrega revisado se encuentra en "
        "ffdefdf8035c994ae285a270de0a4ff4e9f336a8, y las correcciones de octubre permanecen como cambios "
        "locales sobre 49aa51c0167fdb24c5f2f3a6ab6e3f397830b462. Por tanto, esa ejecución histórica no "
        "certifica ni la entrega revisada ni el árbol corregido; la evidencia CI del commit finalmente "
        "entregado queda pendiente.",
    )
    baseline = find_paragraph(document, "La comparación incluye")
    insert_after(
        baseline,
        "Un bootstrap pareado no paramétrico de auditoría (2.000 remuestras, semilla 20261005) comparó "
        "PlayerNet crudo con regresión logística sobre el test persistido. La diferencia media crudo menos "
        "logística fue −0,00015 en ROC-AUC, con intervalo percentil del 95 % [−0,00269; 0,00247], y 0,00819 "
        "en PR-AUC, con intervalo [−0,00251; 0,01954]. Ambos intervalos incluyen cero: esta corrida no aporta "
        "evidencia suficiente de superioridad; tampoco demuestra equivalencia. El análisis es retrospectivo, "
        "utiliza un único split y no reemplaza una evaluación con varias semillas y datos reales.",
    )

    # Ethical and legal scope before conclusions.
    conclusion = find_paragraph(document, "7. CONCLUSIÓN")
    insert_before(conclusion, "6.4 Consideraciones éticas y protección de datos", "Heading 2")
    insert_before(
        conclusion,
        "La aplicación fue verificada con datos sintéticos y no se utilizó con menores reales. Un uso real "
        "debería definir una base jurídica y una finalidad específica, informar a los titulares y a sus "
        "representantes, obtener el consentimiento que corresponda, recolectar solo datos necesarios, aplicar "
        "plazos de retención y mecanismos de acceso, rectificación y supresión, y limitar el acceso por rol. "
        "La Ley argentina 25.326 regula los principios de protección de datos, los derechos de los titulares y "
        "las obligaciones de responsables de archivos y bancos de datos. Estas medidas son requisitos "
        "propuestos para una etapa posterior; el MVP no implementa un flujo completo de consentimiento de "
        "tutores ni una política automatizada de retención y baja.",
    )
    insert_before(
        conclusion,
        "Etiquetar a una persona menor de edad como de «potencial bajo» puede reforzar sesgos de edad relativa, "
        "maduración, contexto socioeconómico o calidad del registro, y afectar oportunidades formativas. El "
        "score debe tratarse como señal revisable, acompañado por sus límites y nunca como decisión automática. "
        "Antes de un piloto real se requieren evaluación ética, análisis de sesgos por subgrupos, supervisión "
        "humana, canal de impugnación y seguimiento de consecuencias no deseadas.",
    )
    replace(
        document,
        "Por último, en cuanto a los objetivos",
        "El objetivo general se alcanzó en su dimensión técnica: se implementó y verificó un MVP que integra "
        "registro, análisis, comparación e inferencia. Los objetivos de reducir subjetividad, identificar mejor "
        "talento, optimizar recursos o promover igualdad de oportunidades no fueron validados y permanecen como "
        "hipótesis de trabajo para estudios con clubes, profesionales y datos reales. El valor comprobado reside "
        "en la integración y trazabilidad del flujo, siempre como complemento del juicio humano.",
    )
    rewrite_table(
        document.tables[18],
        [
            ["Objetivo declarado en 1.2", "Estado", "Evidencia y alcance"],
            ["Implementar y verificar el MVP integral", "Cumplido localmente", "Registro, historiales, paneles, comparadores, PlayerNet y 116 pruebas aprobadas; CI del árbol corregido pendiente."],
            ["Estructurar evaluación e historiales", "Cumplido a nivel funcional", "Escalas, validaciones e historiales implementados; no se midió reducción de subjetividad."],
            ["Priorizar perfiles para revisión humana", "Implementado, no validado deportivamente", "Filtros, comparadores y score disponibles; sin contraste con jugadores o procesos reales."],
            ["Apoyar decisiones con información trazable", "Cumplido a nivel funcional", "Ficha, panel y comparadores integran información; no hubo estudio de uso con scouts."],
            ["Evaluar factibilidad para clubes con recursos limitados", "Cumplido técnicamente", "MVP reproducible; igualdad de oportunidades, adopción y efecto social no medidos."],
            ["Organizar y recuperar información", "Cumplido a nivel funcional", "Persistencia y consulta verificadas; no se cuantificaron tiempos, costos ni ahorro de recursos."],
        ],
    )

    # Verified references; retrieval dates are omitted where APA does not require them.
    bibliography = find_paragraph(document, "9. ANEXOS TÉCNICOS")
    for reference in [
        "Niculescu-Mizil, A., & Caruana, R. (2005). Predicting good probabilities with supervised learning. Proceedings of the 22nd International Conference on Machine Learning, 625–632. https://doi.org/10.1145/1102351.1102430",
        "Saito, T., & Rehmsmeier, M. (2015). The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets. PLOS ONE, 10(3), e0118432. https://doi.org/10.1371/journal.pone.0118432",
        "Cobley, S., Baker, J., Wattie, N., & McKenna, J. (2009). Annual age-grouping and athlete development: A meta-analytical review of relative age effects in sport. Sports Medicine, 39(3), 235–256. https://doi.org/10.2165/00007256-200939030-00005",
        "Honorable Congreso de la Nación Argentina. (2000). Ley 25.326 de Protección de los Datos Personales. https://www.argentina.gob.ar/normativa/nacional/ley-25326-2000-64790",
    ]:
        insert_before(bibliography, reference)

    # Reproduction now distinguishes runtime from development and includes the portable demo.
    reproduce = find_paragraph(document, "9.4 Reproducción mínima")
    replace(
        document,
        "Para el pipeline de datos",
        "Los comandos siguientes se ejecutan desde la raíz C:\\Tesis\\TPScouting. requirements.txt "
        "contiene runtime y requirements-dev.txt agrega pruebas, auditoría y herramientas documentales. La "
        "instalación de PyTorch depende de plataforma y acelerador; debe usarse el selector oficial de PyTorch. "
        "Se verificó Windows con Python 3.11; no se afirma una ejecución local en Linux o macOS.",
    )
    code_anchor = find_paragraph(document, "py -3.11 -m venv")
    code_anchor.text = "Windows PowerShell: py -3.11 -m venv .venv\nLinux/macOS: python3.11 -m venv .venv"
    install_anchor = find_paragraph(document, ".\\.venv\\Scripts\\python.exe -m pip install")
    install_anchor.text = "Windows: .\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt -r requirements-dev.txt\nLinux/macOS: .venv/bin/python -m pip install -r requirements.txt -r requirements-dev.txt"
    run_app = find_paragraph(document, ".\\.venv\\Scripts\\python.exe .\\scouting_app\\app.py")
    run_app.text = "Demo portable recomendada: .\\.venv\\Scripts\\python.exe .\\scripts\\iniciar_demo.py\nLinux/macOS: .venv/bin/python scripts/iniciar_demo.py"
    insert_after(
        reproduce,
        "Repositorio de entrega revisado: commit ffdefdf8035c994ae285a270de0a4ff4e9f336a8. "
        "Las correcciones documentadas aquí corresponden al árbol local sobre "
        "49aa51c0167fdb24c5f2f3a6ab6e3f397830b462 y todavía no fueron sincronizadas ni publicadas.",
    )

    document.core_properties.title = "TPScouting — borrador corregido octubre 2026"
    document.core_properties.comments = (
        "Copia de trabajo generada desde el Word aprobado SHA-256 77261990…; "
        "legajo, índices/campos, CI final y revisión visual de PDF pendientes."
    )
    document.save(OUTPUT)
    print(f"backup={BACKUP}")
    print(f"output={OUTPUT}")
    print(f"source_sha256={digest(SOURCE)}")
    print(f"backup_sha256={digest(BACKUP)}")
    print(f"output_sha256={digest(OUTPUT)}")


if __name__ == "__main__":
    main()
