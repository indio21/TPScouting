"""Apply the final, evidence-preserving language pass to the thesis DOCX."""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from docx import Document

EXPECTED_SOURCE_SHA256 = "5eaa742a6cd83bc91c0c406844a74348b3711d25846fa785afB24b9f8c3df36c".lower()

REPLACEMENTS = {
    "4.1.2 Diseño de la Base de datos": "4.1.2 Diseño de la base de datos",
    "4.1.3 Interfaz de Usuario": "4.1.3 Interfaz de usuario",
    "Machine Learning (ML) es una rama de la inteligencia artificial que permite a las máquinas aprender y mejorar a partir de la experiencia, sin ser explícitamente programadas para cada tarea. Esencialmente, ML se centra en el desarrollo de algoritmos que pueden procesar datos y realizar predicciones o tomar decisiones basadas en esos datos (Hastie et al., 2009; Goodfellow et al., 2016).":
        "El aprendizaje automático (machine learning, ML) es una rama de la inteligencia artificial que permite a las máquinas aprender a partir de la experiencia sin ser programadas de manera explícita para cada tarea. Se centra en el desarrollo de algoritmos capaces de procesar datos y producir predicciones o decisiones basadas en ellos (Hastie et al., 2009; Goodfellow et al., 2016).",
    "Bibliotecas para Machine Learning: PyTorch y Scikit-learn.":
        "Bibliotecas de aprendizaje automático: PyTorch y scikit-learn.",
    "La Figura 4-3 resume los casos de uso principales del sistema, agrupando los roles de administrador, scout, director y el sistema de Machine Learning.":
        "La Figura 4-3 resume los casos de uso principales y agrupa los roles de administrador, scout y director, junto con el sistema de aprendizaje automático.",
    "No se detectó una inclusión directa de la etiqueta o de datos futuros dentro de las variables de entrada. Sin embargo, los cuantiles y las cuotas utilizados para construir el target, se calculan sobre el conjunto completo antes de dividir entrenamiento, validación y prueba. Esta decisión introduce un riesgo metodológico: las etiquetas de validación y prueba no son totalmente independientes de la distribución global, por lo que las métricas deben interpretarse como evidencia técnica preliminar.":
        "No se detectó una inclusión directa de la etiqueta ni de datos futuros entre las variables de entrada. Sin embargo, los cuantiles y las cuotas utilizados para construir el target se calculan sobre el conjunto completo antes de dividirlo en entrenamiento, validación y prueba. Esta decisión introduce un riesgo metodológico: las etiquetas de validación y prueba no son totalmente independientes de la distribución global, por lo que las métricas deben interpretarse como evidencia técnica preliminar.",
    "El MVP se desplegó en Render con Gunicorn y PostgreSQL administrado. Por limitación del plan, se utiliza una sola base PostgreSQL para la demostración. La aplicación se valida mediante /health, login, panel general, listado, comparadores y vistas de jugador.":
        "El MVP se desplegó en Render con Gunicorn y PostgreSQL administrado. Debido a las limitaciones del plan, se utiliza una sola base PostgreSQL para la demostración. La aplicación se valida mediante /health, el inicio de sesión, el panel general, el listado, los comparadores y las vistas de jugador.",
    "La Figura 5-3 representa la arquitectura utilizada en el despliegue histórico de Render con Gunicorn, Flask, PostgreSQL y artefactos de Machine Learning. El despliegue automático depende de la rama configurada en Render; la auditoría no pudo verificar cuál es la rama conectada actualmente. El cache y el rate limiting residen en memoria, mientras que el lock del pipeline combina un bloqueo de thread con un archivo creado de forma atómica.":
        "La Figura 5-3 representa la arquitectura utilizada en el despliegue histórico de Render con Gunicorn, Flask, PostgreSQL y artefactos de aprendizaje automático. El despliegue automático depende de la rama configurada en Render; la auditoría no pudo verificar cuál estaba conectada en ese momento. La caché y la limitación de solicitudes residen en memoria, mientras que el bloqueo del pipeline combina un bloqueo de hilo con un archivo creado de forma atómica.",
    "De acuerdo a la última corrida tomada, la misma utiliza semilla 42, 20.000 jugadores sintéticos y una partición de 14.000/3.000/3.000. La clase positiva representa el 8 % del conjunto de prueba. Las métricas se presentan por separado para la salida sigmoid cruda de PlayerNet y para la salida ajustada mediante calibración isotónica; ninguna de ellas corresponde al score combinado mostrado por la aplicación.":
        "La corrida documentada utiliza la semilla 42, 20.000 jugadores sintéticos y una partición de 14.000/3.000/3.000. La clase positiva representa el 8 % del conjunto de prueba. Las métricas se presentan por separado para la salida sigmoid cruda de PlayerNet y para la salida ajustada mediante calibración isotónica; ninguna corresponde al score combinado mostrado por la aplicación.",
    "6.2.3 Evidencia operativa del deploy": "6.2.3 Evidencia operativa del despliegue",
    "Tabla 6-5. Smoke real del deploy en Render.": "Tabla 6-5. Prueba real de disponibilidad del despliegue en Render.",
    "La CI #73 es evidencia histórica del commit bc5ddd35d0fa3bf6d85faab637772d4e9025fc98 en el repositorio de desarrollo. El repositorio de entrega revisado se encuentra en ffdefdf8035c994ae285a270de0a4ff4e9f336a8, y las correcciones de octubre permanecen como cambios locales sobre 49aa51c0167fdb24c5f2f3a6ab6e3f397830b462. Por tanto, esa ejecución histórica no certifica ni la entrega revisada ni el árbol corregido; la evidencia CI del commit finalmente entregado queda pendiente.":
        "La CI #73 es evidencia histórica del commit bc5ddd35d0fa3bf6d85faab637772d4e9025fc98 en el repositorio de desarrollo. El repositorio de entrega revisado se encuentra en ffdefdf8035c994ae285a270de0a4ff4e9f336a8, mientras que las correcciones de octubre se guardaron en commits locales posteriores y todavía no se sincronizaron. Por tanto, esa ejecución histórica no certifica la entrega revisada ni el árbol corregido; la evidencia CI del commit finalmente entregado queda pendiente.",
    "9. ANEXOS TÉCNICOS ": "9. ANEXOS TÉCNICOS",
    "Repositorio de entrega revisado: https://github.com/indio21/TPScouting-entrega, commit ffdefdf8035c994ae285a270de0a4ff4e9f336a8. El trabajo de corrección se realiza en el repositorio principal a partir del checkpoint 6e29b45a85396b5cbe82e8b0ece2b0eb394a7bfd y todavía no fue sincronizado ni publicado en la entrega.":
        "Repositorio de entrega revisado: https://github.com/indio21/TPScouting-entrega, commit ffdefdf8035c994ae285a270de0a4ff4e9f336a8. El trabajo de corrección se realizó en el repositorio principal a partir del checkpoint 6e29b45a85396b5cbe82e8b0ece2b0eb394a7bfd, se guardó mediante commits locales posteriores y todavía no fue sincronizado ni publicado en la entrega.",
    "10.1 Evidencia Render": "10.1 Evidencia de Render",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_paragraph_text(paragraph, old: str, new: str) -> None:
    if paragraph.text != old:
        raise RuntimeError("El texto fuente cambió antes de reemplazarlo.")
    if not paragraph.runs:
        paragraph.add_run(new)
        return
    paragraph.runs[0].text = new
    for run in paragraph.runs[1:]:
        run.text = ""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if sha256(args.source) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("El DOCX fuente no coincide con el cierre verificado del Bloque 8.")

    document = Document(args.source)
    counts = {old: 0 for old in REPLACEMENTS}
    for paragraph in document.paragraphs:
        if paragraph.text in REPLACEMENTS:
            old = paragraph.text
            replace_paragraph_text(paragraph, old, REPLACEMENTS[old])
            counts[old] += 1

    missing = [old for old, count in counts.items() if count != 1]
    if missing:
        raise RuntimeError(f"Reemplazos no unívocos: {missing}")

    removed = 0
    for paragraph in list(document.paragraphs):
        if (
            not paragraph.text.strip()
            and paragraph.style.name in {"Heading 3", "Figure Caption Generated"}
            and not paragraph._p.xpath(".//w:drawing | .//w:pict")
        ):
            paragraph._p.getparent().remove(paragraph._p)
            removed += 1

    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    checked = Document(args.output)
    if len(checked.inline_shapes) != 26 or len(checked.tables) != 19:
        raise RuntimeError("La revisión alteró la estructura gráfica o tabular.")
    print(f"OUTPUT={args.output}")
    print(f"REPLACEMENTS={len(REPLACEMENTS)}")
    print(f"EMPTY_ARTIFACTS_REMOVED={removed}")
    print(f"SHA256={sha256(args.output).upper()}")


if __name__ == "__main__":
    main()
