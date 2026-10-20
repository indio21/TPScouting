"""Check the final Word/PDF layout using their actual text and page geometry."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

import fitz
from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "docs/revision_octubre_2026/final"
DOCX = FOLDER / "TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_2026-10-20.docx"
PDF = DOCX.with_suffix(".pdf")
REPORT = FOLDER / "control_formato_2026-10-20.json"


def normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().casefold()


def footer_numbers(page) -> list[str]:
    return [
        word[4]
        for word in page.get_text("words")
        if word[1] > page.rect.height - 65 and word[4].isdigit()
    ]


def main() -> None:
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    small_sizes = []
    with ZipFile(DOCX) as archive:
        zip_error = archive.testzip()
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                root = etree.fromstring(archive.read(name))
                for node in root.xpath(".//w:sz | .//w:szCs", namespaces=namespace):
                    size = float(node.get("{" + namespace["w"] + "}val"))
                    if size < 20:
                        small_sizes.append({"part": name, "points": size / 2})

    document = Document(DOCX)
    with fitz.open(PDF) as pdf:
        texts = [page.get_text() for page in pdf]
        normalized_texts = [normalized(text) for text in texts]
        numbers = [footer_numbers(page) for page in pdf]
        number_issues = [
            {"physical_page": index + 1, "footer": numbers[index]}
            for index in range(1, len(pdf))
            if numbers[index] != [str(index)]
        ]
        toc_issues = []
        toc_entries = 0
        for paragraph in document.paragraphs:
            if not paragraph.style.name.lower().startswith("toc"):
                continue
            if "\t" not in paragraph.text:
                continue
            heading, listed = paragraph.text.rsplit("\t", 1)
            if not listed.isdigit() or heading.startswith(("Figura ", "Tabla ")):
                continue
            toc_entries += 1
            hits = [
                index
                for index, text in enumerate(normalized_texts)
                if index >= 5 and normalized(heading) in text
            ]
            if not any(listed in numbers[index] for index in hits):
                toc_issues.append({"heading": heading, "listed_page": listed})

        separate_sections = []
        for title in ("RESUMEN", "ABSTRACT", "8. BIBLIOGRAFÍA", "9. ANEXOS TÉCNICOS"):
            for index, text in enumerate(texts):
                lines = [line.strip() for line in text.splitlines() if line.strip()]
                if title in lines:
                    separate_sections.append(
                        {
                            "title": title,
                            "physical_page": index + 1,
                            "starts_on_own_page": lines.index(title) == 1,
                        }
                    )

        out_of_bounds = []
        figure_orphans = []
        for index, page in enumerate(pdf):
            for block in page.get_text("blocks"):
                if (
                    block[0] < -1
                    or block[1] < -1
                    or block[2] > page.rect.width + 1
                    or block[3] > page.rect.height + 1
                ):
                    out_of_bounds.append(index + 1)
            if index >= 10 and re.search(r"(?mi)^\s*figura \d+-\d+\.", texts[index]):
                if not page.get_images(full=True):
                    figure_orphans.append(index + 1)

        result = {
            "docx": str(DOCX),
            "docx_sha256": hashlib.sha256(DOCX.read_bytes()).hexdigest().upper(),
            "pdf_sha256": hashlib.sha256(PDF.read_bytes()).hexdigest().upper(),
            "pdf_pages": len(pdf),
            "zip_error": zip_error,
            "font_sizes_below_10pt": small_sizes,
            "cover_has_visible_page_number": bool(numbers[0]),
            "numbering_issues": number_issues,
            "toc_entries_checked": toc_entries,
            "toc_mismatches": toc_issues,
            "separate_sections": separate_sections,
            "out_of_bounds_pages": out_of_bounds,
            "figure_captions_without_image": figure_orphans,
            "images": len(document.inline_shapes),
            "tables": len(document.tables),
            "sections": len(document.sections),
        }
        result["passed"] = (
            not zip_error
            and not small_sizes
            and not numbers[0]
            and not number_issues
            and not toc_issues
            and toc_entries == 88
            and len(separate_sections) == 4
            and all(item["starts_on_own_page"] for item in separate_sections)
            and not out_of_bounds
            and not figure_orphans
            and (len(document.inline_shapes), len(document.tables), len(document.sections))
            == (28, 19, 9)
        )
    REPORT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=True, indent=2))
    if not result["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
