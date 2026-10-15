"""Create a new thesis copy with verified limitations and future work."""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "docs/revision_octubre_2026/word/"
    "TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_VERIFICADA_2026-10-10.docx"
)
OUTPUT = (
    ROOT
    / "docs/revision_octubre_2026/word/"
    "TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_ACTUALIZADA_2026-10-15.docx"
)
EXPECTED_SOURCE_SHA256 = "d7ffbce018c9fd438012083e48f053618c9e14e67a30ab249412a02522d4c946"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_runs(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_before(reference, text: str, style: str = "Normal"):
    paragraph = reference.insert_paragraph_before(text, style=style)
    paragraph.paragraph_format.keep_together = True
    return paragraph


def main() -> None:
    if digest(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("La fuente no coincide con la copia final verificada.")

    document = Document(SOURCE)

    replacements = {
        457: (
            "• La corrida histórica definió cuantiles y cuotas sobre el conjunto completo antes del split. "
            "Esta limitación no puede corregirse retrospectivamente; una corrida candidata posterior separó "
            "los datos antes de ajustar la política del target y mantuvo los artefactos históricos intactos."
        ),
        458: (
            "• El score combinado mostrado por la interfaz no formó parte de la evaluación histórica original. "
            "Una evaluación posterior sobre test no demostró una mejora defendible frente a la probabilidad "
            "cruda y obtuvo peores resultados de Brier y log loss; por ello se conserva como heurística "
            "secundaria y no como probabilidad ni como mejora demostrada."
        ),
        459: (
            "• La corrida histórica no registra validation loss, duración total ni SHA Git exacto. Esos datos "
            "son irrecuperables y no se reconstruyen. La corrida candidata posterior registra configuración, "
            "dependencias, seed, splits, métricas, tiempos y hashes para mejorar la trazabilidad futura."
        ),
        464: (
            "Los resultados muestran que el flujo técnico es reproducible en el entorno evaluado, pero no "
            "permiten inferir utilidad deportiva. PlayerNet y la regresión logística presentan desempeños "
            "próximos en un único split sintético; el intervalo bootstrap de sus diferencias incluye cero. En "
            "consecuencia, la complejidad de PlayerNet no queda justificada por una superioridad estadística en "
            "esta corrida. La evaluación posterior del score combinado tampoco demostró una mejora defendible "
            "frente a la probabilidad cruda, por lo que se interpreta únicamente como heurística secundaria."
        ),
    }
    for index, text in replacements.items():
        replace_runs(document.paragraphs[index], text)

    limitations_heading = next(p for p in document.paragraphs if p.text.strip() == "7.4 Lecciones aprendidas")
    limitations = [
        (
            "• La calidad estática se controla mediante un conjunto gradual de reglas Ruff. Permanecen 732 "
            "avisos, principalmente vinculados con modernización de anotaciones e importaciones; se consideran "
            "deuda técnica controlada y no defectos funcionales demostrados."
        ),
        (
            "• La adopción de una application factory fue evaluada y diferida. El acoplamiento entre "
            "inicialización, persistencia, artefactos de aprendizaje automático y pruebas convierte ese cambio "
            "en un refactor transversal cuyo riesgo supera el beneficio inmediato para el MVP."
        ),
        (
            "• La compatibilidad fue comprobada en Windows, Linux y macOS sobre Apple Silicon. macOS Intel no "
            "forma parte de las plataformas soportadas porque la versión fijada de PyTorch no dispone de una "
            "distribución compatible para esa arquitectura."
        ),
        (
            "• El sistema continúa restringido a datos sintéticos. Los controles para consentimiento o "
            "representación, ejercicio de derechos, retención, auditoría, seguridad e incidentes fueron "
            "diseñados, pero no están implementados ni habilitan el tratamiento de datos reales de menores."
        ),
        (
            "• Los metadatos ausentes de la corrida histórica son irrecuperables. La trazabilidad se corrige "
            "hacia adelante mediante una corrida nueva y separada, sin atribuirle retrospectivamente duración, "
            "pérdidas o identificadores que no fueron registrados."
        ),
    ]
    for text in limitations:
        insert_before(limitations_heading, text)

    future_heading = next(p for p in document.paragraphs if p.text.strip() == "7.5 Trabajo futuro")
    paragraphs = document.paragraphs
    heading_index = next(i for i, p in enumerate(paragraphs) if p.text.strip() == future_heading.text.strip())
    next_heading_index = next(
        i
        for i, p in enumerate(paragraphs[heading_index + 1 :], start=heading_index + 1)
        if p.style.name.startswith("Heading")
    )
    future_anchor = next(
        p
        for p in paragraphs[heading_index + 1 : next_heading_index]
        if not p.text.strip()
    )
    future_items = [
        (
            "• Reducir gradualmente la deuda estática por regla o módulo, con pruebas específicas y sin aplicar "
            "una modernización masiva que dificulte atribuir regresiones."
        ),
        (
            "• Reconsiderar una application factory solo si el crecimiento del proyecto justifica separar el "
            "ciclo de vida de base, caché, rate limiting y artefactos, mediante etapas reversibles y pruebas de "
            "dos aplicaciones aisladas."
        ),
        (
            "• Evaluar una combinación compatible de Python y PyTorch si se requiere ampliar el soporte a "
            "macOS Intel, repitiendo instalación, carga de artefactos, suite y demo antes de declararlo soportado."
        ),
        (
            "• Implementar y validar los controles institucionales, jurídicos y técnicos diseñados antes de "
            "cualquier piloto con datos reales de menores, junto con validación externa y supervisión humana."
        ),
        (
            "• Mantener para toda nueva corrida el registro de configuración, dependencias, seeds, splits, "
            "métricas, tiempos, commit y hashes, sin modificar ni completar retrospectivamente la evidencia "
            "histórica."
        ),
    ]
    for text in future_items:
        insert_before(future_anchor, text)

    # The figure and table lists in the approved source are static text. The
    # added page shifts must therefore be synchronized explicitly.
    list_pages = {
        "Figura 10-1. Evidencia histórica de la CI #73 del repositorio de desarrollo.": 84,
        "Figura 10-2. Diagrama de componentes de TPScouting ampliado.": 86,
        "Figura 10-3. Diagrama de clases del modelo de dominio ampliado.": 87,
        "Figura 10-4. Diagrama de secuencia de inferencia ampliado.": 88,
        "Figura 10-5. Diagrama de secuencia del panel general ampliado.": 89,
        "Figura 10-6. Diagrama de despliegue en Render ampliado.": 90,
        "Figura 10-7. Diagrama de casos de uso general ampliado.": 92,
        "Figura 10-8. Diagrama de casos de uso de gestión de jugadores ampliado.": 93,
        "Figura 10-9. Diagrama de casos de uso de análisis y decisión scout ampliado.": 94,
        "Tabla 9-1. Rutas principales.": 79,
        "Tabla 9-2. Variables de entorno relevantes.": 79,
        "Tabla 9-3. Glosario técnico.": 80,
    }
    updated_lists = set()
    for paragraph in document.paragraphs:
        if "\t" not in paragraph.text:
            continue
        caption, _, _ = paragraph.text.rpartition("\t")
        if caption not in list_pages:
            continue
        paragraph.runs[-1].text = f"\t{list_pages[caption]}"
        updated_lists.add(caption)
    if updated_lists != set(list_pages):
        missing = sorted(set(list_pages) - updated_lists)
        raise RuntimeError(f"No se pudieron actualizar estas entradas de listas: {missing}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    checked = Document(OUTPUT)
    if (len(checked.inline_shapes), len(checked.tables), len(checked.sections)) != (28, 19, 9):
        raise RuntimeError("La copia alteró imágenes, tablas o secciones.")
    all_text = "\n".join(p.text for p in checked.paragraphs)
    required = (
        "Permanecen 732 avisos",
        "application factory fue evaluada y diferida",
        "macOS Intel no forma parte",
        "no están implementados ni habilitan el tratamiento",
        "metadatos ausentes de la corrida histórica son irrecuperables",
        "evaluación posterior del score combinado",
    )
    missing = [text for text in required if text not in all_text]
    if missing:
        raise RuntimeError(f"Faltan textos requeridos: {missing}")
    if "tampoco cuenta con una evaluación independiente en test" in all_text:
        raise RuntimeError("Se conservó la afirmación obsoleta sobre el score combinado.")

    print(f"OUTPUT={OUTPUT}")
    print(f"SHA256={digest(OUTPUT).upper()}")
    print(f"PARAGRAPHS={len(checked.paragraphs)} TABLES={len(checked.tables)} IMAGES={len(checked.inline_shapes)}")


if __name__ == "__main__":
    main()
