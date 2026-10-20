"""Create clean and review copies of the thesis with six pending points updated.

The source document is immutable and identified by SHA-256. Images and index
fields are intentionally preserved for the final layout pass.
"""

from __future__ import annotations

import hashlib
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    r"C:\Tesis\TPScouting-entrega\documento"
) / "TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_2026-10-15.docx"
OUTPUT_DIR = ROOT / "docs/revision_octubre_2026/word"
CLEAN_OUTPUT = (
    OUTPUT_DIR
    / "TRABAJO_FINAL_TPScouting_ACTUALIZADO_SEIS_PUNTOS_2026-10-19.docx"
)
REVIEW_OUTPUT = (
    OUTPUT_DIR
    / "TRABAJO_FINAL_TPScouting_ACTUALIZADO_SEIS_PUNTOS_CAMBIOS_ROJO_2026-10-19.docx"
)
EXPECTED_SOURCE_SHA256 = (
    "8385153268405e1069496e3e39d774c97db4c47b35f295c4b064324171cc97ae"
)
REVIEW_RED = RGBColor(0xC0, 0x00, 0x00)
MINIMUM_FONT_PT = 10.0


TEXT_REPLACEMENTS = {
    # Neutral academic framing: do not assert institutional prevalence or impact
    # that was not measured in this project.
    "En el dinámico mundo del fútbol contemporáneo, los clubes de comunidades "
    "pequeñas y de recursos limitados enfrentan desafíos significativos, "
    "especialmente en el ámbito del fútbol infantojuvenil. En muchos casos, estas "
    "instituciones constituyen el primer espacio de formación deportiva y social "
    "para niños y adolescentes, pero se ven condicionadas por restricciones "
    "financieras, logísticas y de infraestructura que dificultan el seguimiento "
    "sistemático de sus jóvenes jugadores.":
        "El trabajo toma como contexto a los clubes de fútbol formativo con recursos "
        "limitados que necesitan registrar y consultar de manera sistemática la "
        "evolución de sus jugadores juveniles. El proyecto no midió la frecuencia ni "
        "la magnitud de este problema en instituciones reales; lo utiliza como punto "
        "de partida para evaluar una solución técnica acotada.",
    "Esta situación plantea una paradoja. El talento juvenil puede convertirse en una "
    "fuente de valor deportivo y económico en el mediano plazo. Sin embargo, la falta "
    "de inversión y de procesos de registro y análisis reduce la capacidad de los "
    "clubes para identificar oportunidades, planificar la formación y sostener "
    "decisiones basadas en evidencia.":
        "Dentro de ese contexto, la ausencia de un proceso uniforme de registro y "
        "análisis puede dificultar la recuperación de antecedentes, la comparación "
        "entre jugadores y el seguimiento de su evolución. El trabajo no cuantifica "
        "inversión, desempeño institucional ni valor económico generado.",
    "Además, la recolección y el análisis de datos suelen ser insuficientes. Aunque se "
    "generan registros a través de entrenamientos, partidos y evaluaciones, con "
    "frecuencia dichos datos no se almacenan de forma estructurada o se pierden, lo "
    "que limita la posibilidad de comparar jugadores, detectar patrones de evolución "
    "y evaluar el potencial con criterios consistentes.":
        "Cuando los registros de entrenamientos, partidos y evaluaciones no se "
        "almacenan de forma estructurada, resulta más difícil recuperar la información, "
        "comparar jugadores y observar su evolución con criterios consistentes. Esta "
        "relación constituye la necesidad funcional abordada por el MVP, no un "
        "resultado de campo del estudio.",
    "En consecuencia, muchos jóvenes con potencial de desarrollo pueden pasar "
    "desapercibidos o no recibir seguimiento adecuado. Esto impacta tanto en sus "
    "oportunidades de crecimiento deportivo y profesional como en la sostenibilidad "
    "de los clubes, que dependen de maximizar el rendimiento de recursos formativos "
    "reducidos.":
        "Un seguimiento incompleto puede limitar la detección oportuna de cambios en "
        "el rendimiento y dificultar la revisión de decisiones formativas. El proyecto "
        "no evaluó efectos sobre oportunidades deportivas, trayectorias profesionales "
        "ni sostenibilidad institucional.",
    # Rate limiting reflects the implemented key exactly: normalized username.
    "El login aplica rate limiting local mediante una clave de cliente definida por "
    "la política de proxies confiables. Esta mitigación fue verificada para una "
    "instancia, pero no sustituye un limitador distribuido en un despliegue "
    "multi-instancia.":
        "El login aplica rate limiting local por nombre de usuario normalizado y no "
        "utiliza encabezados de IP controlables por el cliente para construir la clave. "
        "Esta mitigación fue verificada para una instancia, pero no sustituye un "
        "limitador distribuido en un despliegue multi-instancia.",
    "• Rate limiting local basado en una clave de cliente definida por una política de "
    "proxies confiables; al ser memoria de proceso, no cubre varias instancias.":
        "• Rate limiting local por nombre de usuario normalizado, sin confiar en "
        "encabezados de IP suministrados por el cliente; al utilizar memoria de proceso, "
        "no cubre varias instancias.",
    "• /health público sin excepciones internas, CSP compatible con los recursos "
    "inventariados, HSTS bajo HTTPS y validación central de photo_url.":
        "• /health público sin excepciones internas, CSP compatible con los recursos "
        "inventariados y validación central de photo_url. HSTS se emite cuando la "
        "aplicación reconoce una solicitud HTTPS; su efectividad en producción depende "
        "de la terminación TLS y de la configuración confiable del proxy.",
    # Distinguish the UI's primary displayed value from predictive evidence.
    "Flujo técnico de inferencia: Flask recupera el jugador y sus historiales mediante "
    "SQLAlchemy; el preprocesador construye y transforma un vector de 68 variables; "
    "PlayerNet produce logits y sigmoid obtiene base_prob, la probabilidad cruda. Si "
    "está disponible, el calibrador isotónico genera una referencia secundaria. La "
    "interfaz calcula combined_prob mediante combine_probability(base_prob, "
    "stats_summary, fit_score), combinando la señal cruda, el promedio histórico y el "
    "ajuste posicional, con renormalización cuando falta algún componente. "
    "combined_prob es el resultado operativo principal y la salida calibrada no lo "
    "sustituye.":
        "Flujo técnico de inferencia: Flask recupera el jugador y sus historiales "
        "mediante SQLAlchemy; el preprocesador transforma las 64 columnas de entrada "
        "en un vector de 68 variables; PlayerNet produce logits y sigmoid obtiene "
        "base_prob, la probabilidad cruda. Si está disponible, el calibrador isotónico "
        "genera una referencia secundaria. La interfaz calcula combined_prob mediante "
        "combine_probability(base_prob, stats_summary, fit_score), combinando la señal "
        "cruda, el promedio histórico y el ajuste posicional, con renormalización "
        "cuando falta algún componente. combined_prob es el valor principal mostrado "
        "por la interfaz, pero se interpreta como una heurística operativa y no como "
        "una probabilidad calibrada ni como una mejora predictiva demostrada.",
    "La verificación local del 15/10/2026 obtuvo 123 pruebas aprobadas, 1 omitida y 4 "
    "advertencias conocidas, con una cobertura de 83,29 %. El smoke visual Playwright "
    "opt-in se ejecutó por separado. Las advertencias corresponden a RuntimeWarning de "
    "scikit-learn por columnas completamente NaN en dos pruebas de preprocesamiento. "
    "Las ejecuciones de CI y del contenedor se vinculan con commits exactos de los "
    "repositorios principal y de entrega.":
        "La verificación local del 19/10/2026 obtuvo 123 pruebas aprobadas, 1 omitida y "
        "4 advertencias conocidas, con una cobertura de 83,29 %. El smoke visual "
        "Playwright opt-in se ejecutó por separado. Las advertencias corresponden a "
        "RuntimeWarning de scikit-learn por columnas completamente NaN en dos pruebas "
        "de preprocesamiento. La cobertura global supera el umbral configurado, aunque "
        "no es uniforme: compare registra 70 % y settings 65 %, por lo que ambos "
        "dominios permanecen como prioridades de ampliación. Las ejecuciones de CI y "
        "del contenedor se vinculan con commits exactos de los repositorios principal y "
        "de entrega.",
    "La CI #73 se conserva como evidencia histórica del commit "
    "bc5ddd35d0fa3bf6d85faab637772d4e9025fc98 en el repositorio de desarrollo. La "
    "entrega ejecutable más reciente se publicó en el commit "
    "8093fe616c9c50b72078ff4f241d9ab3427e763d y fue certificada por la ejecución CI "
    "38656130065 y el smoke de contenedor 38656129892. La figura siguiente documenta "
    "la ejecución histórica y no sustituye esa evidencia posterior.":
        "La CI #73 se conserva como evidencia histórica del commit "
        "bc5ddd35d0fa3bf6d85faab637772d4e9025fc98 en el repositorio de desarrollo. El "
        "código funcional de la entrega corresponde al commit "
        "8093fe616c9c50b72078ff4f241d9ab3427e763d, certificado por la ejecución CI "
        "38656130065 y el smoke de contenedor 38656129892. Ese código se integró en el "
        "paquete de entrega publicado en "
        "4c5ef2c88494e257a3bd8232bccba7f33f79b66a, verificado por la ejecución CI "
        "38815440730 y el smoke de contenedor 38815440745. La figura siguiente documenta "
        "la ejecución histórica y no sustituye esas evidencias posteriores.",
    "• El score combinado mostrado por la interfaz no formó parte de la evaluación "
    "histórica original. Una evaluación posterior sobre test no demostró una mejora "
    "defendible frente a la probabilidad cruda y obtuvo peores resultados de Brier y "
    "log loss; por ello se conserva como heurística secundaria y no como probabilidad "
    "ni como mejora demostrada.":
        "• El score combinado mostrado por la interfaz no formó parte de la evaluación "
        "histórica original. Una evaluación posterior sobre test no demostró una mejora "
        "defendible frente a la probabilidad cruda y obtuvo peores resultados de Brier "
        "y log loss. Por ello se conserva como valor operativo principal de la interfaz, "
        "pero como heurística secundaria desde el punto de vista de la evidencia "
        "predictiva; no se presenta como probabilidad ni como mejora demostrada.",
    "Los resultados muestran que el flujo técnico es reproducible en el entorno "
    "evaluado, pero no permiten inferir utilidad deportiva. PlayerNet y la regresión "
    "logística presentan desempeños próximos en un único split sintético; el intervalo "
    "bootstrap de sus diferencias incluye cero. En consecuencia, la complejidad de "
    "PlayerNet no queda justificada por una superioridad estadística en esta corrida. "
    "La evaluación posterior del score combinado tampoco demostró una mejora defendible "
    "frente a la probabilidad cruda, por lo que se interpreta únicamente como "
    "heurística secundaria.":
        "Los resultados muestran que el flujo técnico es reproducible en el entorno "
        "evaluado, pero no permiten inferir utilidad deportiva. PlayerNet y la regresión "
        "logística presentan desempeños próximos en un único split sintético; el "
        "intervalo bootstrap de sus diferencias incluye cero. En consecuencia, la "
        "complejidad de PlayerNet no queda justificada por una superioridad estadística "
        "en esta corrida. La evaluación posterior del score combinado tampoco demostró "
        "una mejora defendible frente a la probabilidad cruda. El score se mantiene "
        "como valor operativo de la interfaz, pero su interpretación metodológica es la "
        "de una heurística secundaria sin superioridad predictiva demostrada.",
    "Repositorio de entrega: https://github.com/indio21/TPScouting-entrega, commit "
    "publicado 8093fe616c9c50b72078ff4f241d9ab3427e763d, validado por la ejecución CI "
    "38656130065 y el smoke de contenedor 38656129892. El repositorio principal conserva "
    "el historial de trabajo y las evidencias internas que no forman parte de la "
    "entrega académica.":
        "Repositorio de entrega: https://github.com/indio21/TPScouting-entrega. El código "
        "funcional corresponde al commit "
        "8093fe616c9c50b72078ff4f241d9ab3427e763d, validado por la ejecución CI "
        "38656130065 y el smoke de contenedor 38656129892. La revisión documental y el "
        "paquete previo se publicaron en el commit "
        "4c5ef2c88494e257a3bd8232bccba7f33f79b66a, validado por la ejecución CI "
        "38815440730 y el smoke de contenedor 38815440745. El repositorio principal "
        "conserva el historial de trabajo y las evidencias internas que no forman parte "
        "de la entrega académica.",
    # Cross-platform commands. These are documentation changes only; they do not
    # claim a platform execution beyond the evidence already stated in the thesis.
    ".\\.venv\\Scripts\\python.exe -m pytest -q --cov=scouting_app "
    "--cov-report=term-missing":
        "Windows PowerShell:\n"
        ".\\.venv\\Scripts\\python.exe -m pytest -q --cov=scouting_app "
        "--cov-report=term-missing\n"
        "Linux/macOS:\n"
        ".venv/bin/python -m pytest -q --cov=scouting_app --cov-report=term-missing",
    "Pipeline reproducible de datos, entrenamiento, evaluación y sincronización de "
    "demo:":
        "Pipeline reproducible opcional de datos, entrenamiento, evaluación y "
        "sincronización de demo. Debe ejecutarse en una copia de trabajo y con artefactos "
        "de salida separados para no sobrescribir la corrida histórica:",
    "Set-Location .\\scouting_app":
        "Windows PowerShell: Set-Location .\\scouting_app\n"
        "Linux/macOS: cd scouting_app",
    "..\\.venv\\Scripts\\python.exe generate_data.py --num-players 20000 --db-url "
    "sqlite:///players_training.db --seed 42 --min-age 12 --max-age 18 --reset":
        "Windows PowerShell:\n"
        "..\\.venv\\Scripts\\python.exe generate_data.py --num-players 20000 "
        "--db-url sqlite:///players_training.db --seed 42 --min-age 12 --max-age 18 --reset\n"
        "Linux/macOS:\n"
        "../.venv/bin/python generate_data.py --num-players 20000 --db-url "
        "sqlite:///players_training.db --seed 42 --min-age 12 --max-age 18 --reset",
    "..\\.venv\\Scripts\\python.exe train_model.py --db-url "
    "sqlite:///players_training.db --model-out model.pt --preprocessor-out "
    "preprocessor.joblib --calibrator-out probability_calibrator.joblib "
    "--metadata-out training_metadata.json --splits-out training_splits.json "
    "--epochs 45 --lr 5e-4 --patience 10":
        "Windows PowerShell:\n"
        "..\\.venv\\Scripts\\python.exe train_model.py --db-url "
        "sqlite:///players_training.db --model-out model_candidate.pt --preprocessor-out "
        "preprocessor_candidate.joblib --calibrator-out "
        "probability_calibrator_candidate.joblib --metadata-out "
        "training_metadata_candidate.json --splits-out training_splits_candidate.json "
        "--epochs 45 --lr 5e-4 --patience 10\n"
        "Linux/macOS:\n"
        "../.venv/bin/python train_model.py --db-url sqlite:///players_training.db "
        "--model-out model_candidate.pt --preprocessor-out "
        "preprocessor_candidate.joblib --calibrator-out "
        "probability_calibrator_candidate.joblib --metadata-out "
        "training_metadata_candidate.json --splits-out training_splits_candidate.json "
        "--epochs 45 --lr 5e-4 --patience 10",
    "..\\.venv\\Scripts\\python.exe evaluate_saved_model.py --db-url "
    "sqlite:///players_training.db --metadata-path training_metadata.json":
        "Windows PowerShell:\n"
        "..\\.venv\\Scripts\\python.exe evaluate_saved_model.py --db-url "
        "sqlite:///players_training.db --metadata-path training_metadata_candidate.json\n"
        "Linux/macOS:\n"
        "../.venv/bin/python evaluate_saved_model.py --db-url "
        "sqlite:///players_training.db --metadata-path training_metadata_candidate.json",
    "..\\.venv\\Scripts\\python.exe sync_shortlist.py --src-db "
    "sqlite:///players_training.db --dst-db sqlite:///players_updated_v2.db --limit 100 "
    "--min-age 12 --max-age 18 --replace":
        "Windows PowerShell:\n"
        "..\\.venv\\Scripts\\python.exe sync_shortlist.py --src-db "
        "sqlite:///players_training.db --dst-db sqlite:///players_updated_v2.db "
        "--limit 100 --min-age 12 --max-age 18 --replace\n"
        "Linux/macOS:\n"
        "../.venv/bin/python sync_shortlist.py --src-db "
        "sqlite:///players_training.db --dst-db sqlite:///players_updated_v2.db "
        "--limit 100 --min-age 12 --max-age 18 --replace",
    "Set-Location ..":
        "Windows PowerShell: Set-Location ..\n"
        "Linux/macOS: cd ..",
    ".\\.venv\\Scripts\\python.exe scripts\\smoke_render.py --base-url "
    "https://tpscouting-mvp.onrender.com":
        "Windows PowerShell:\n"
        ".\\.venv\\Scripts\\python.exe scripts\\smoke_render.py --base-url "
        "https://tpscouting-mvp.onrender.com\n"
        "Linux/macOS:\n"
        ".venv/bin/python scripts/smoke_render.py --base-url "
        "https://tpscouting-mvp.onrender.com",
}


TABLE_UPDATES = {
    ("Evidencia", 1, 0): "pytest local (19/10/2026)",
    ("Evidencia", 5, 0): "Commit funcional del repositorio principal de respaldo",
    ("Evidencia", 5, 1): (
        "754ccb3601a489ed1249a078ad700d1eb546a494; distribución reproducible verificada"
    ),
    ("Evidencia", 6, 0): "Commit funcional del código de entrega",
    ("Evidencia", 6, 1): (
        "8093fe616c9c50b72078ff4f241d9ab3427e763d; CI 38656130065 y contenedor "
        "38656129892 aprobados"
    ),
    ("Evidencia", 7, 0): "Paquete de entrega publicado antes de esta actualización",
    ("Evidencia", 7, 1): "4c5ef2c88494e257a3bd8232bccba7f33f79b66a",
    ("Evidencia", 8, 0): "Validación del paquete publicado",
    ("Evidencia", 8, 1): (
        "CI 38815440730 y smoke de contenedor 38815440745 aprobados"
    ),
    ("Objetivo declarado en 1.2", 1, 2): (
        "Registro, historiales, paneles, comparadores y PlayerNet; 123 pruebas "
        "aprobadas, 1 omitida y 83,29 % de cobertura. Código funcional: CI "
        "38656130065. Paquete publicado: CI 38815440730."
    ),
}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def all_paragraphs(document: Document):
    """Yield body, table, header and footer paragraphs without duplicates."""
    seen = set()

    def emit(paragraph):
        key = paragraph._p
        if key not in seen:
            seen.add(key)
            return paragraph
        return None

    for paragraph in document.paragraphs:
        result = emit(paragraph)
        if result is not None:
            yield result
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    result = emit(paragraph)
                    if result is not None:
                        yield result
    for section in document.sections:
        parts = (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        )
        for part in parts:
            for paragraph in part.paragraphs:
                result = emit(paragraph)
                if result is not None:
                    yield result
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            result = emit(paragraph)
                            if result is not None:
                                yield result


def set_paragraph_text(paragraph, text: str, review: bool) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
    else:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    if review:
        for run in paragraph.runs:
            if run.text:
                run.font.color.rgb = REVIEW_RED


def set_cell_text(cell, text: str, review: bool) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text, review)
    for extra in cell.paragraphs[1:]:
        set_paragraph_text(extra, "", False)


def apply_content_updates(document: Document, review: bool) -> tuple[int, int]:
    by_text = {paragraph.text: paragraph for paragraph in all_paragraphs(document)}
    missing = [text for text in TEXT_REPLACEMENTS if text not in by_text]
    if missing:
        raise RuntimeError(
            "No se encontraron todos los textos fuente; se detiene para no aplicar "
            f"cambios ambiguos. Faltan {len(missing)}: {missing[:2]}"
        )

    for old_text, new_text in TEXT_REPLACEMENTS.items():
        set_paragraph_text(by_text[old_text], new_text, review)

    tables = {table.cell(0, 0).text: table for table in document.tables}
    for (header, row, column), text in TABLE_UPDATES.items():
        if header not in tables:
            raise RuntimeError(f"No se encontró la tabla con encabezado {header!r}.")
        set_cell_text(tables[header].cell(row, column), text, review)

    return len(TEXT_REPLACEMENTS), len(TABLE_UPDATES)


def normalize_exposed_font_sizes(document: Document) -> int:
    changes = 0
    for style in document.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            size = style.font.size
            if size is not None and size.pt < MINIMUM_FONT_PT:
                style.font.size = Pt(MINIMUM_FONT_PT)
                changes += 1
    for paragraph in all_paragraphs(document):
        for run in paragraph.runs:
            size = run.font.size
            if run.text and size is not None and size.pt < MINIMUM_FONT_PT:
                run.font.size = Pt(MINIMUM_FONT_PT)
                changes += 1
    return changes


_FONT_SIZE_RE = re.compile(
    rb'(<w:sz(?:Cs)?\b[^>]*\bw:val=")([0-9]+(?:\.[0-9]+)?)(")'
)


def normalize_package_font_sizes(path: Path) -> int:
    """Raise WordprocessingML font sizes below 10 pt, including text boxes."""
    replacements = 0

    def update_xml(data: bytes) -> bytes:
        nonlocal replacements

        def replace(match: re.Match[bytes]) -> bytes:
            nonlocal replacements
            value = float(match.group(2))
            if value < MINIMUM_FONT_PT * 2:
                replacements += 1
                return match.group(1) + b"20" + match.group(3)
            return match.group(0)

        return _FONT_SIZE_RE.sub(replace, data)

    with tempfile.TemporaryDirectory(prefix="tpscouting_docx_") as temp_dir:
        temp_path = Path(temp_dir) / path.name
        with zipfile.ZipFile(path, "r") as source_zip, zipfile.ZipFile(
            temp_path, "w", zipfile.ZIP_DEFLATED
        ) as output_zip:
            for item in source_zip.infolist():
                data = source_zip.read(item.filename)
                if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                    data = update_xml(data)
                output_zip.writestr(item, data)
        shutil.copy2(temp_path, path)
    return replacements


def count_package_small_sizes(path: Path) -> int:
    count = 0
    with zipfile.ZipFile(path, "r") as archive:
        for name in archive.namelist():
            if not (name.startswith("word/") and name.endswith(".xml")):
                continue
            data = archive.read(name)
            count += sum(
                1
                for match in _FONT_SIZE_RE.finditer(data)
                if float(match.group(2)) < MINIMUM_FONT_PT * 2
            )
    return count


def validate_output(path: Path, expected_red: bool) -> dict[str, int | str]:
    if not zipfile.is_zipfile(path):
        raise RuntimeError(f"El DOCX generado no es un paquete ZIP válido: {path}")
    if count_package_small_sizes(path):
        raise RuntimeError(f"Persisten tamaños menores a 10 pt en {path.name}.")

    document = Document(path)
    if (len(document.inline_shapes), len(document.tables), len(document.sections)) != (
        28,
        19,
        9,
    ):
        raise RuntimeError("La copia alteró imágenes, tablas o secciones.")

    full_text = "\n".join(paragraph.text for paragraph in all_paragraphs(document))
    required = (
        "4c5ef2c88494e257a3bd8232bccba7f33f79b66a",
        "rate limiting local por nombre de usuario normalizado",
        "heurística secundaria desde el punto de vista de la evidencia predictiva",
        "Linux/macOS: cd scouting_app",
        "compare registra 70 % y settings 65 %",
        "Los metadatos ausentes de la corrida histórica son irrecuperables",
    )
    missing = [text for text in required if text not in full_text]
    if missing:
        raise RuntimeError(f"Faltan actualizaciones requeridas: {missing}")

    red_runs = 0
    for paragraph in all_paragraphs(document):
        for run in paragraph.runs:
            if run.text and run.font.color.rgb == REVIEW_RED:
                red_runs += 1
    if expected_red and red_runs == 0:
        raise RuntimeError("La copia de revisión no contiene cambios en rojo.")
    if not expected_red and red_runs:
        raise RuntimeError("La copia limpia contiene texto rojo de revisión.")

    return {
        "sha256": digest(path).upper(),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "images": len(document.inline_shapes),
        "sections": len(document.sections),
        "red_runs": red_runs,
    }


def build(output: Path, review: bool) -> tuple[int, int, int, int]:
    document = Document(SOURCE)
    content_changes, table_changes = apply_content_updates(document, review)
    exposed_font_changes = normalize_exposed_font_sizes(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    package_font_changes = normalize_package_font_sizes(output)
    return content_changes, table_changes, exposed_font_changes, package_font_changes


def main() -> None:
    if digest(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("La fuente no coincide con el Word final verificado.")

    print(f"SOURCE={SOURCE}")
    print(f"SOURCE_SHA256={digest(SOURCE).upper()}")
    for output, review in ((CLEAN_OUTPUT, False), (REVIEW_OUTPUT, True)):
        counts = build(output, review)
        result = validate_output(output, expected_red=review)
        print(f"OUTPUT={output}")
        print(f"CHANGES=paragraphs:{counts[0]} tables:{counts[1]}")
        print(f"FONT_NORMALIZATION=exposed:{counts[2]} package:{counts[3]}")
        for key, value in result.items():
            print(f"{key.upper()}={value}")


if __name__ == "__main__":
    main()
