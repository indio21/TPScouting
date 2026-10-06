from __future__ import annotations

import sys
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


sys.stdout.reconfigure(encoding="utf-8")

SOURCE = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_AUDITADA_21-08-2026.docx"
)
OUTPUT = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_CORREGIDA_21-08-2026.docx"
)


def norm(text: str) -> str:
    return " ".join((text or "").split())


def find_exact(doc: Document, text: str) -> Paragraph:
    needle = norm(text)
    matches = [p for p in doc.paragraphs if norm(p.text) == needle]
    if len(matches) != 1:
        raise ValueError(f"Se esperaba 1 coincidencia para {text!r}; se encontraron {len(matches)}")
    return matches[0]


def set_paragraph(paragraph: Paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def replace_exact(doc: Document, old: str, new: str) -> None:
    set_paragraph(find_exact(doc, old), new)


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    node = OxmlElement("w:p")
    paragraph._p.addnext(node)
    inserted = Paragraph(node, paragraph._parent)
    if style:
        inserted.style = style
    inserted.add_run(text)
    return inserted


def remove_between(start: Paragraph, end: Paragraph) -> None:
    node = start._p.getnext()
    while node is not None and node is not end._p:
        following = node.getnext()
        node.getparent().remove(node)
        node = following


def append_word_field(paragraph: Paragraph, instruction: str) -> None:
    run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run.append(begin)
    paragraph._p.append(run)

    run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run.append(instr)
    paragraph._p.append(run)

    run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run.append(separate)
    paragraph._p.append(run)

    result_run = OxmlElement("w:r")
    result_text = OxmlElement("w:t")
    result_text.text = "Actualizar campo en Microsoft Word"
    result_run.append(result_text)
    paragraph._p.append(result_run)

    run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(end)
    paragraph._p.append(run)


def append_tc_field(paragraph: Paragraph, text: str, identifier: str) -> None:
    safe_text = text.replace('"', "'")
    run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run.append(begin)
    paragraph._p.append(run)

    run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' TC "{safe_text}" \\f {identifier} '
    run.append(instr)
    paragraph._p.append(run)

    run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(end)
    paragraph._p.append(run)


def add_field_after(heading: Paragraph, instruction: str) -> Paragraph:
    field_paragraph = insert_after(heading, "")
    append_word_field(field_paragraph, instruction)
    return field_paragraph


def ensure_style(doc: Document, name: str, base_name: str) -> None:
    if name in doc.styles:
        return
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles[base_name]


def rebuild_automatic_lists(doc: Document) -> None:
    ensure_style(doc, "Figure Caption Generated", "Caption")
    ensure_style(doc, "Table Caption Generated", "Caption")

    index_heading = find_exact(doc, "ÍNDICE")
    figure_heading = find_exact(doc, "LISTA DE FIGURAS")
    table_heading = find_exact(doc, "LISTA DE TABLAS")
    summary_heading = find_exact(doc, "RESUMEN")

    remove_between(index_heading, figure_heading)
    remove_between(figure_heading, table_heading)
    remove_between(table_heading, summary_heading)

    for paragraph in doc.paragraphs:
        text = norm(paragraph.text)
        if paragraph.style.name == "Caption" and text.startswith("Figura "):
            paragraph.style = "Figure Caption Generated"
            append_tc_field(paragraph, text, "F")
        elif paragraph.style.name == "Caption" and text.startswith("Tabla "):
            paragraph.style = "Table Caption Generated"
            append_tc_field(paragraph, text, "T")

    add_field_after(index_heading, 'TOC \\o "1-4" \\h \\z')
    add_field_after(figure_heading, 'TOC \\f F \\h \\z')
    add_field_after(table_heading, 'TOC \\f T \\h \\z')

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def correct_text(doc: Document) -> None:
    replacements = {
        "Visualización y comunicación de datos: Mediante la presentación gráfica y la generación de informes basados en los análisis realizados, el sistema facilita la interpretación de los datos y también, mejora la comunicación y comprensión de la información por parte de entrenadores, directivos y otras partes interesadas.":
            "Visualización y comunicación de datos: mediante paneles, gráficos, comparadores, vistas imprimibles y el registro de reportes scout, el sistema facilita la interpretación y comunicación de la información entre entrenadores, directivos y otras partes interesadas. La exportación formal de informes y archivos PDF se mantiene como trabajo futuro.",
        "Migraciones con Flask-Migrate: A menudo se usa junto con Flask-Migrate, que permite gestionar los cambios en el esquema de la base de datos (migraciones) de manera sencilla utilizando el sistema de migraciones de Alembic. Esto es útil para mantener sincronizada la base de datos con el código a medida que evoluciona la aplicación.":
            "Migraciones del MVP: SQLAlchemy puede integrarse con Flask-Migrate y Alembic, pero TPScouting no incorpora esas herramientas. La versión auditada crea el esquema con Base.metadata.create_all y aplica ampliaciones compatibles mediante funciones de migración manual en db_utils.py. La adopción de migraciones versionadas con Alembic queda como mejora futura.",
        "La base de datos se diseñó de forma relacional. En el MVP se implementa en SQLite para manejar grandes cantidades de datos históricos relacionados con jugadores juveniles.":
            "La base de datos se diseñó de forma relacional mediante SQLAlchemy. En desarrollo local y pruebas, el MVP utiliza SQLite por su simplicidad y portabilidad; el despliegue público en Render utiliza PostgreSQL administrado para conservar los datos fuera del sistema de archivos efímero del servicio. El MVP no implementa infraestructura Big Data ni demuestra escalabilidad para grandes volúmenes.",
        "Gestión de Bases de Datos: SQL se utiliza para gestionar bases de datos donde se almacenan grandes volúmenes de datos deportivos. Es esencial para realizar consultas eficientes, almacenamiento y recuperación de datos relacionados con jugadores, estadísticas de partidos, etc.":
            "Gestión de bases de datos: SQL permite consultar, almacenar y recuperar datos deportivos estructurados, como jugadores, atributos, evaluaciones y estadísticas. En TPScouting se utiliza mediante SQLAlchemy sobre SQLite local y PostgreSQL en Render; el MVP no demuestra escalabilidad para grandes volúmenes.",
        "La aplicación también calcula combined_prob, un score operativo que combina la probabilidad cruda de PlayerNet, el promedio histórico del jugador y el ajuste por posición con pesos 0,35, 0,35 y 0,30. Cuando falta algún componente, los pesos disponibles se renormalizan. Este score es el valor integrado utilizado por la interfaz, pero no fue evaluado de manera independiente en el conjunto de prueba.":
            "La aplicación también calcula combined_prob, un score operativo que combina la probabilidad cruda de PlayerNet, el promedio histórico del jugador y el ajuste por posición. Los pesos por defecto son 0,35 para el modelo, 0,35 para el promedio histórico y 0,30 para el ajuste posicional; pueden configurarse mediante POT_W_MODEL, POT_W_RATING y POT_W_FIT. Cuando falta algún componente, los pesos disponibles se renormalizan. Este score es el valor integrado utilizado por la interfaz, pero no fue evaluado de manera independiente en el conjunto de prueba.",
        "El modelo de datos se implementa con SQLAlchemy. La entidad central es Player, vinculada a historiales de atributos, estadísticas, partidos, disponibilidad, evaluaciones físicas y reportes scout. La fecha de nacimiento es la fuente para calcular edad actual y categoría juvenil.":
            "El modelo de datos se implementa con SQLAlchemy. La entidad central es Player, vinculada a historiales de atributos, estadísticas, partidos, disponibilidad, evaluaciones físicas y reportes scout. La fecha de nacimiento es la fuente persistida para calcular la edad actual y la categoría juvenil; current_age y category_year son valores derivados, no columnas almacenadas.",
        "La Figura 4-2 muestra las clases principales del modelo de dominio persistido con SQLAlchemy y sus relaciones mas relevantes.":
            "La Figura 4-2 muestra las clases principales del modelo de dominio persistido con SQLAlchemy y sus relaciones más relevantes. La edad actual y la categoría juvenil se identifican como propiedades derivadas de birth_date.",
        "Flujo técnico de inferencia: Flask recupera el jugador mediante SQLAlchemy; el preprocesador transforma edad, posición, atributos e historial en un vector de 68 variables; PlayerNet produce logits y la función sigmoid obtiene la probabilidad cruda. La calibración isotónica se utiliza en la evaluación experimental. Para la interfaz, la aplicación calcula combined_prob mediante un promedio ponderado de la probabilidad cruda del modelo (0,35), el promedio histórico (0,35) y el ajuste posicional (0,30), con renormalización cuando falta algún componente.":
            "Flujo técnico de inferencia: Flask recupera el jugador y sus historiales mediante SQLAlchemy; el preprocesador construye y transforma un vector de 68 variables; PlayerNet produce logits y sigmoid obtiene base_prob, la probabilidad cruda. Si está disponible, el calibrador isotónico genera una referencia secundaria. La interfaz calcula combined_prob mediante combine_probability(base_prob, stats_summary, fit_score), combinando la señal cruda, el promedio histórico y el ajuste posicional, con renormalización cuando falta algún componente. combined_prob es el resultado operativo principal y la salida calibrada no lo sustituye.",
        "La Figura 5-1 detalla el flujo de inferencia cuando un usuario consulta la proyección de potencial de un jugador.":
            "La Figura 5-1 detalla el flujo real de inferencia y separa las ramas de error: si el modelo no está disponible, la ruta responde HTTP 500; si el modelo está cargado pero no existen datos suficientes, se renderiza una vista controlada sin proyección.",
        "Figura 5-1. Secuencia de predicción de potencial de jugador.":
            "Figura 5-1. Secuencia real de inferencia: score combinado principal y calibración secundaria.",
        "La Figura 5-3 representa el despliegue verificado del MVP en Render con Gunicorn, Flask, PostgreSQL y artefactos de Machine Learning.":
            "La Figura 5-3 representa la arquitectura utilizada en el despliegue histórico de Render con Gunicorn, Flask, PostgreSQL y artefactos de Machine Learning. El despliegue automático depende de la rama configurada en Render; la auditoría no pudo verificar cuál es la rama conectada actualmente. El cache y el rate limiting residen en memoria, mientras que el lock del pipeline combina un bloqueo de thread con un archivo creado de forma atómica.",
        "Figura 5-3. Diagrama de despliegue en Render.":
            "Figura 5-3. Arquitectura de despliegue en Render y mecanismos operativos del MVP.",
        "Como evidencia operativa adicional, la Tabla 6-5 presenta el smoke HTTP ejecutado contra Render el 20/05/2026. El primer request refleja el arranque frío del plan Free; las rutas principales respondieron correctamente en esa verificación. Un control posterior del 10/07/2026 terminó por timeout, por lo que no se afirma disponibilidad continua actual.":
            "Como evidencia operativa adicional, la Tabla 6-5 presenta el smoke HTTP ejecutado contra Render el 20/05/2026. El primer request refleja el arranque frío del plan Free; las rutas principales respondieron correctamente en esa verificación histórica. Los controles del 10/07/2026, 13/07/2026 y 21/08/2026 finalizaron por timeout. Por ello no se afirma disponibilidad continua ni estado vigente del servicio; un timeout aislado tampoco demuestra una caída permanente.",
        "La Figura 6-6 presenta la vista de predicción, donde se exponen señales de decisión y evolución histórica.":
            "La Figura 6-6 presenta la vista de predicción con la edad visible del jugador y el ajuste combinado respecto de PlayerNet crudo. Ese ajuste incorpora tanto el historial disponible como el ajuste posicional.",
        "Figura 6-6. Vista de predicción de potencial del jugador.":
            "Figura 6-6. Vista de predicción con edad y ajuste combinado de historial y posición.",
        "• Healthcheck histórico documentado con 100 jugadores demo y faltantes críticos en 0; un control del 10/07/2026 finalizó por timeout.":
            "• Healthcheck histórico documentado con 100 jugadores demo y faltantes críticos en 0; los controles del 10/07/2026, 13/07/2026 y 21/08/2026 finalizaron por timeout, por lo que no se afirma disponibilidad vigente.",
        "• Cache, rate limiting y lock de pipeline son locales al proceso.":
            "• Cache y rate limiting son locales al proceso. El lock del pipeline combina un bloqueo de thread con un archivo atómico; reduce ejecuciones simultáneas en el MVP, pero no reemplaza una coordinación distribuida multi-instancia.",
    }
    for old, new in replacements.items():
        replace_exact(doc, old, new)


def correct_tables(doc: Document) -> None:
    migration_table = doc.tables[2]
    if norm(migration_table.rows[4].cells[0].text) != "Sistema de migraciones":
        raise ValueError("No se encontro la fila Sistema de migraciones en Tabla 4-2")
    migration_table.rows[4].cells[1].text = (
        "El MVP usa creación de esquema y migraciones manuales en db_utils.py; "
        "Flask-Migrate/Alembic constituye una mejora futura para versionar cambios."
    )

    dimensions = doc.tables[5]
    expected = ["Técnica", "Física", "Defensiva", "Mental", "Rendimiento"]
    actual = [norm(row.cells[0].text) for row in dimensions.rows[1:]]
    if actual != expected:
        raise ValueError(f"Filas inesperadas en Tabla 4-5: {actual}")
    values = {
        "Técnica": ("pace, shooting, passing, dribbling, technique, vision", "Player e historial de atributos"),
        "Física": ("physical, estimated_speed, endurance, height_cm, weight_kg", "Player y PhysicalAssessment"),
        "Defensiva": ("defending, tackling", "Player e historial de atributos"),
        "Mental": ("determination, decision_making, tactical_reading, mental_profile, adaptability", "Player y ScoutReport"),
        "Rendimiento": ("minutes, goals, assists, pass_accuracy, shot_accuracy, duels_won_pct, final_score", "PlayerStat y participaciones"),
    }
    for row in dimensions.rows[1:]:
        variables, source = values[norm(row.cells[0].text)]
        row.cells[1].text = variables
        row.cells[2].text = source

    deployment = doc.tables[12]
    if not any(norm(row.cells[0].text) == "Rama de deploy" for row in deployment.rows):
        cells = deployment.add_row().cells
        cells[0].text = "Rama de deploy"
        cells[1].text = (
            "Evidencia histórica: render-free-deploy. La rama conectada actualmente en Render no fue verificada."
        )


def correct_numbering(doc: Document) -> None:
    replace_exact(
        doc,
        "Tabla 4-4. Capas y responsabilidades del MVP.",
        "Tabla 4-3. Capas y responsabilidades del MVP.",
    )
    replace_exact(
        doc,
        "Tabla 4-3. Comparación entre motores relacionales.",
        "Tabla 4-4. Comparación entre motores relacionales.",
    )

    figure_replacements = {
        "La Figura 10-9 conserva una vista ampliada de la ejecución CI #71, con jobs exitosos para Python 3.11 y 3.12 y dos artefactos de cobertura.":
            "La Figura 10-1 conserva una vista ampliada de la ejecución CI #71, con jobs exitosos para Python 3.11 y 3.12 y dos artefactos de cobertura.",
        "Figura 10-9. Historial público y ejecución exitosa del workflow CI de TPScouting.":
            "Figura 10-1. Historial público y ejecución exitosa del workflow CI de TPScouting.",
        "La Figura 10-1 presenta en tamaño ampliado los componentes del cliente, la aplicación Flask, servicios, persistencia y módulo de Machine Learning.":
            "La Figura 10-2 presenta en tamaño ampliado los componentes del cliente, la aplicación Flask, servicios, persistencia y módulo de Machine Learning.",
        "Figura 10-1. Diagrama de componentes de TPScouting ampliado.":
            "Figura 10-2. Diagrama de componentes de TPScouting ampliado.",
        "La Figura 10-2 presenta en tamaño ampliado las clases principales definidas en el modelo SQLAlchemy y sus relaciones.":
            "La Figura 10-3 presenta en tamaño ampliado las clases principales definidas en el modelo SQLAlchemy y sus relaciones; current_age y category_year son propiedades derivadas de birth_date.",
        "Figura 10-2. Diagrama de clases del modelo de dominio ampliado.":
            "Figura 10-3. Diagrama de clases del modelo de dominio ampliado.",
        "La Figura 10-3 presenta en tamaño ampliado el flujo de consulta de potencial de un jugador.":
            "La Figura 10-4 presenta en tamaño ampliado el flujo real de inferencia: base_prob cruda, calibración secundaria, historial, ajuste posicional y combined_prob como resultado operativo principal.",
        "Figura 10-3. Diagrama de secuencia de predicción ampliado.":
            "Figura 10-4. Diagrama de secuencia de inferencia ampliado.",
        "La Figura 10-4 presenta en tamaño ampliado el flujo de carga del panel general, incluyendo cache y consultas a base de datos.":
            "La Figura 10-5 presenta en tamaño ampliado el flujo de carga del panel general, incluyendo cache y consultas a base de datos.",
        "Figura 10-4. Diagrama de secuencia del panel general ampliado.":
            "Figura 10-5. Diagrama de secuencia del panel general ampliado.",
        "La Figura 10-5 presenta en tamaño ampliado la relación entre navegador, Render, Gunicorn, Flask, PostgreSQL y artefactos ML.":
            "La Figura 10-6 presenta en tamaño ampliado la relación entre navegador, Render, Gunicorn, Flask, PostgreSQL y artefactos ML. El deploy depende de la rama configurada en Render; cache y rate limiting son de memoria, y el lock del pipeline usa thread y archivo atómico.",
        "Figura 10-5. Diagrama de despliegue en Render ampliado.":
            "Figura 10-6. Diagrama de despliegue en Render ampliado.",
        "La Figura 10-6 presenta en tamaño ampliado la vista general de casos de uso del sistema.":
            "La Figura 10-7 presenta en tamaño ampliado la vista general de casos de uso del sistema.",
        "Figura 10-6. Diagrama de casos de uso general ampliado.":
            "Figura 10-7. Diagrama de casos de uso general ampliado.",
        "La Figura 10-7 presenta en tamaño ampliado los casos de uso del módulo de jugadores.":
            "La Figura 10-8 presenta en tamaño ampliado los casos de uso del módulo de jugadores.",
        "Figura 10-7. Diagrama de casos de uso de gestión de jugadores ampliado.":
            "Figura 10-8. Diagrama de casos de uso de gestión de jugadores ampliado.",
        "La Figura 10-8 presenta en tamaño ampliado los casos de uso de análisis, comparación y predicción.":
            "La Figura 10-9 presenta en tamaño ampliado los casos de uso de análisis, comparación y predicción.",
        "Figura 10-8. Diagrama de casos de uso de análisis y decisión scout ampliado.":
            "Figura 10-9. Diagrama de casos de uso de análisis y decisión scout ampliado.",
    }
    for old, new in figure_replacements.items():
        replace_exact(doc, old, new)


def correct_reproduction_commands(doc: Document) -> None:
    replace_exact(
        doc,
        "Los siguientes comandos corresponden al flujo documentado del repositorio. Deben ejecutarse desde la raíz del proyecto en Windows y con las variables de entorno necesarias. No se conserva el transcript exacto de la corrida original de entrenamiento.",
        "Los comandos generales se ejecutan desde la raíz del proyecto en Windows. Para el pipeline de datos y entrenamiento se cambia al directorio scouting_app, de modo que las rutas SQLite y los artefactos coincidan con el RUNBOOK. Los comandos se reproducen completos; no se conserva el transcript exacto de la corrida original y esta corrección documental no reentrena el modelo.",
    )

    label_generate = find_exact(doc, "Generar datos de entrenamiento:")
    cmd_generate = find_exact(doc, "generate_data.py --num-players 20000 --db-url sqlite:///players_training.db --seed 42 --min-age 12 --max-age 18 --reset")
    label_train = find_exact(doc, "Entrenar:")
    cmd_train = find_exact(doc, "train_model.py ... --epochs 45 --lr 5e-4 --patience 10")
    label_eval = find_exact(doc, "Evaluar el modelo:")
    cmd_eval = find_exact(doc, "evaluate_saved_model.py --db-url sqlite:///players_training.db --metadata-path training_metadata.json")

    set_paragraph(label_generate, "Pipeline reproducible de datos, entrenamiento, evaluación y sincronización de demo:")
    set_paragraph(cmd_generate, "Set-Location .\\scouting_app")
    set_paragraph(label_train, "Generar el dataset sintético oficial:")
    set_paragraph(cmd_train, "..\\.venv\\Scripts\\python.exe generate_data.py --num-players 20000 --db-url sqlite:///players_training.db --seed 42 --min-age 12 --max-age 18 --reset")
    set_paragraph(label_eval, "Entrenar PlayerNet con la configuración documentada:")
    set_paragraph(cmd_eval, "..\\.venv\\Scripts\\python.exe train_model.py --db-url sqlite:///players_training.db --model-out model.pt --preprocessor-out preprocessor.joblib --calibrator-out probability_calibrator.joblib --metadata-out training_metadata.json --splits-out training_splits.json --epochs 45 --lr 5e-4 --patience 10")

    cursor = cmd_eval
    cursor = insert_after(cursor, "Evaluar el modelo guardado:")
    cursor = insert_after(cursor, "..\\.venv\\Scripts\\python.exe evaluate_saved_model.py --db-url sqlite:///players_training.db --metadata-path training_metadata.json", "Código técnico")
    cursor = insert_after(cursor, "Sincronizar la base operativa de demostración:")
    cursor = insert_after(cursor, "..\\.venv\\Scripts\\python.exe sync_shortlist.py --src-db sqlite:///players_training.db --dst-db sqlite:///players_updated_v2.db --limit 100 --min-age 12 --max-age 18 --replace", "Código técnico")
    cursor = insert_after(cursor, "Volver a la raíz del proyecto:")
    insert_after(cursor, "Set-Location ..", "Código técnico")


def main() -> int:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    correct_text(doc)
    correct_tables(doc)
    correct_numbering(doc)
    correct_reproduction_commands(doc)
    rebuild_automatic_lists(doc)

    doc.core_properties.title = "TPScouting - Trabajo Final corregido"
    doc.core_properties.subject = "Versión corregida a partir de la auditoría técnica del 21/08/2026"
    doc.save(OUTPUT)
    print(f"Documento corregido guardado en: {OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
