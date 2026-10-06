from __future__ import annotations

import sys
from pathlib import Path
from shutil import copy2
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph


sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_CORREGIDA_PAGINADA_CON_DIAGRAMAS_21-08-2026.docx"
)
OLD_EVIDENCE_DOCX = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_AUDITADA.docx"
)
OUTPUT = Path(
    r"C:\Users\Usuario\Desktop\TRABAJO_FINAL_TPScouting_ENTREGA_FINAL_COMPLETA_AUDITADA_21-08-2026.docx"
)
EVIDENCE_DIR = ROOT / "docs" / "evidencia_word_render"


EXTRACTED_MEDIA = {
    "word/media/image17.png": "04a_player_detail_perfil.png",
    "word/media/image18.png": "04b_player_detail_historiales.png",
    "word/media/image19.png": "04c_player_detail_reportes.png",
    "word/media/image20.png": "08_ci_run_71_resumen.png",
    "word/media/image21.png": "09_ci_run_71_ampliada.png",
}

FIGURES = [
    ("Figura 6-1. Curva real de entrenamiento generada desde training_metadata.json.", "07_training_curve_real.png", 5.75),
    ("Figura 6-2. Login del MVP desplegado en Render.", "01_login_render.png", 5.75),
    ("Figura 6-3. Panel general o mesa de scouting con métricas accionables.", "02_dashboard_render.png", 4.848),
    ("Figura 6-4. Listado paginado de jugadores con edad, categoría y potencial.", "03_players_render.png", 4.467),
    ("a) Perfil, atributos y resumen de rendimiento.", "04a_player_detail_perfil.png", 5.75),
    ("b) Historiales, gráficos y evaluaciones periódicas.", "04b_player_detail_historiales.png", 5.75),
    ("c) Reportes, disponibilidad y acciones de mantenimiento.", "04c_player_detail_reportes.png", 5.75),
    (
        "Figura 6-6. Vista de predicción con edad y ajuste combinado de historial y posición.",
        "05_prediction_render_corregida_2026-08-21.png",
        5.75,
    ),
    ("Figura 6-7. Comparador múltiple de jugadores.", "06_compare_multi_render.png", 4.764),
    ("Figura 6-8. Resumen público de la ejecución CI #71 en GitHub Actions.", "08_ci_run_71_resumen.png", 5.75),
    (
        "Figura 10-1. Historial público y ejecución exitosa del workflow CI de TPScouting.",
        "09_ci_run_71_ampliada.png",
        3.831,
    ),
]


def norm(text: str) -> str:
    return " ".join((text or "").split())


def find_exact(doc: Document, text: str) -> Paragraph:
    needle = norm(text)
    matches = [paragraph for paragraph in doc.paragraphs if norm(paragraph.text) == needle]
    if len(matches) != 1:
        raise ValueError(f"Se esperaba una coincidencia para {text!r}; se encontraron {len(matches)}")
    return matches[0]


def replace_exact(doc: Document, old: str, new: str) -> None:
    paragraph = find_exact(doc, old)
    paragraph.clear()
    paragraph.add_run(new)


def paragraph_has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def drawing_immediately_before(doc: Document, anchor: Paragraph) -> bool:
    node = anchor._p.getprevious()
    while node is not None and node.tag.endswith("}p"):
        paragraph = Paragraph(node, doc._body)
        if paragraph_has_drawing(paragraph):
            return True
        if norm(paragraph.text):
            return False
        node = node.getprevious()
    return False


def add_image_before(doc: Document, anchor: Paragraph, image_path: Path, width: float) -> None:
    if drawing_immediately_before(doc, anchor):
        raise RuntimeError(f"El ancla ya tiene una imagen: {anchor.text}")

    previous = anchor._p.getprevious()
    image_paragraph = None
    if previous is not None and previous.tag.endswith("}p"):
        candidate = Paragraph(previous, doc._body)
        if not norm(candidate.text) and not paragraph_has_drawing(candidate):
            image_paragraph = candidate

    if image_paragraph is None:
        new_p = OxmlElement("w:p")
        anchor._p.addprevious(new_p)
        image_paragraph = Paragraph(new_p, doc._body)

    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(width))


def extract_historical_evidence() -> None:
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    with ZipFile(OLD_EVIDENCE_DOCX) as archive:
        for member, output_name in EXTRACTED_MEDIA.items():
            payload = archive.read(member)
            if len(payload) < 1_000:
                raise RuntimeError(f"El recurso {member} no parece una captura valida")
            (EVIDENCE_DIR / output_name).write_bytes(payload)


def main() -> int:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not OLD_EVIDENCE_DOCX.exists():
        raise FileNotFoundError(OLD_EVIDENCE_DOCX)

    extract_historical_evidence()
    for _, image_name, _ in FIGURES:
        image_path = EVIDENCE_DIR / image_name
        if not image_path.exists():
            raise FileNotFoundError(image_path)

    copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    initial_shapes = len(doc.inline_shapes)

    replace_exact(
        doc,
        "Las siguientes figuras fueron capturadas desde el despliegue público en Render mediante una sesión autenticada. Cada captura documenta un flujo funcional del MVP y se interpreta como evidencia fechada, no como garantía de disponibilidad permanente.",
        "Las Figuras 6-2 a 6-5 y 6-7 fueron capturadas desde el despliegue público en Render mediante una sesión autenticada. La Figura 6-6 fue reproducida el 21 de agosto de 2026 en un navegador real contra una base SQLite temporal aislada, después de corregir el mapeo de current_age en la vista de predicción. Cada captura se interpreta como evidencia fechada, no como garantía de disponibilidad permanente.",
    )
    replace_exact(
        doc,
        "La Figura 6-6 presenta la vista de predicción con la edad visible del jugador y el ajuste combinado respecto de PlayerNet crudo. Ese ajuste incorpora tanto el historial disponible como el ajuste posicional.",
        "La Figura 6-6 presenta una reproducción local auditada de la vista de predicción, con la edad derivada de birth_date y el ajuste combinado respecto de PlayerNet crudo. Ese ajuste incorpora tanto el historial disponible como el ajuste posicional; la base temporal utilizada para la captura fue eliminada al finalizar.",
    )
    replace_exact(
        doc,
        ".\\.venv\\Scripts\\python.exe scripts\\smoke_render.py --base-url <URL>",
        ".\\.venv\\Scripts\\python.exe scripts\\smoke_render.py --base-url https://tpscouting-mvp.onrender.com",
    )

    for anchor_text, image_name, width in FIGURES:
        anchor = find_exact(doc, anchor_text)
        add_image_before(doc, anchor, EVIDENCE_DIR / image_name, width)

    inserted = len(doc.inline_shapes) - initial_shapes
    if inserted != len(FIGURES):
        raise RuntimeError(f"Se esperaban {len(FIGURES)} imágenes nuevas; se insertaron {inserted}")

    doc.save(OUTPUT)
    print(f"Documento completo generado: {OUTPUT}")
    print(f"Imágenes insertadas: {inserted} (9 figuras; Figura 6-5 contiene 3 capturas)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
